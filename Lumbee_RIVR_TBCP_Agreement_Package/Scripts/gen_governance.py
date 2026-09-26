"""
gen_governance.py — Generator for the GOVERNANCE & AUTHORIZATION documents
(package folder 03_Governance_Authorizations) of the coordinated TBCP fiber
package for the Lumbee Tribe of North Carolina and RIVR Tech.

SINGLE SOURCE OF TRUTH is Scripts/common.py — imported as C. NO deal facts are
invented here: party names, defined terms, term, payment, citations, flags, and
formatting all come from C. Where a fact was not supplied (award number,
approvals, consents, vote tallies, dollar amounts), a highlighted placeholder
C.PH(...) or a flag marker is emitted instead.

Builds 7 files into 03_Governance_Authorizations/:
  1. 01_Lumbee_Tribal_Council_Resolution.docx
  2. 02_RIVR_Tech_Member_Authorization.docx
  3. 03_LREMC_Consent_and_Joinder.docx
  4. 04_Limited_Waiver_of_Sovereign_Immunity.docx
  5. 05_Tax_TERO_Permitting_and_Regulatory_Schedule.xlsx   (table-suited -> Excel)
  6. 06_Insurance_Schedule.xlsx                             (table-suited -> Excel)
  7. 07_Service_Level_Schedule.docx
"""

import sys
sys.path.insert(0, '/home/user/claude-skills/Lumbee_RIVR_TBCP_Agreement_Package/Scripts')
import common as C

from docx.shared import Pt, Inches
from openpyxl import Workbook

OUTDIR = "03_Governance_Authorizations"

# The ten Definitive Agreements (verbatim from C.AGREEMENTS).
MASTER = C.AGREEMENTS['master']
OTHER_NINE = ['depc', 'iru', 'interconnect', 'om', 'retail',
              'finance', 'land', 'privacy', 'transition']


# ---------------------------------------------------------------------------
# Local formatting helpers (docx). The engine (C) supplies the primitives;
# these compose the governance-specific shapes the assignment calls for.
# ---------------------------------------------------------------------------
def head(doc, text, size=12):
    """Left-aligned bold navy section heading."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = C.NAVY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def whereas(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("WHEREAS, ")
    r.bold = True
    C._emit_runs(p, text)
    return p


def resolved(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run("RESOLVED, that ")
    r.bold = True
    C._emit_runs(p, text)
    return p


def sig(doc, entity, roles):
    """Custom signature block: entity heading + one By/Name/Title/Date stack
    per role. roles = list of dicts {name, title}."""
    p = doc.add_paragraph()
    r = p.add_run(entity)
    r.bold = True
    r.font.color.rgb = C.NAVY
    for role in roles:
        lines = [
            ("By:", "____________________________________"),
            ("Name:", C.PH(role["name"])),
            ("Title:", role["title"]),
            ("Date:", "____________________________________"),
        ]
        for label, val in lines:
            lp = doc.add_paragraph()
            lp.paragraph_format.left_indent = Inches(0.25)
            lr = lp.add_run(f"{label}  ")
            lr.bold = True
            C._emit_runs(lp, val)
        C.spacer(doc, 1)


def disclaimer(doc):
    """Standard closing disclaimer that ends every Word document body."""
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    r = p.add_run("DISCLAIMER.  ")
    r.bold = True
    r.font.color.rgb = C.RED
    C._emit_runs(
        p,
        "This document is a preliminary working draft prepared without a "
        "controlling TBCP award or source documents. The provisional baseline "
        f"is the {C.NOFO_NAME}. {C.NOFO_LIVE_NOTE} Nothing in this document "
        "constitutes final legal advice, a binding commitment, or a "
        "representation that any award, approval, consent, resolution, or vote "
        "has been obtained; all such items remain to be confirmed. All "
        "highlighted placeholders and every "
        f"{C.FLAG_BUSINESS}, {C.FLAG_ATTORNEY}, {C.FLAG_GRANT}, and "
        f"{C.FLAG_TECH} item must be resolved by qualified Tribal counsel, "
        "grant advisors, and financial/insurance professionals before "
        "execution. " + C.DRAFT_STATUS_SHORT + ".",
    )
    return p


def intro_recital_parties(doc):
    """Common opening paragraph that fixes the three distinct entities."""
    C.para(
        doc,
        f"The parties and entities referenced in this instrument are: "
        f"(i) the {C.TRIBE_FULL} (the “{C.TRIBE_SHORT}”), the "
        f"prospective TBCP award recipient, steward, and owner of the "
        f"{C.GRANT_FUNDED} and the {C.NETWORK}; (ii) {C.OPERATOR_FULL} "
        f"(“{C.OPERATOR_SHORT}”), the prospective operator/partner; "
        f"and (iii) {C.LREMC_FULL} (“{C.LREMC_SHORT}”), the separate "
        f"owner of the {C.LREMC_ASSETS} (poles, conduit, fiber, easements, "
        f"huts, power, and land), whose participation is by limited consent "
        f"and joinder only. {C.OPERATOR_SHORT} and {C.LREMC_SHORT} are "
        f"distinct entities and this instrument does not treat the "
        f"{C.LREMC_ASSETS} as owned or controlled by {C.OPERATOR_SHORT}.",
    )


# ===========================================================================
# 1.  TRIBAL COUNCIL RESOLUTION
# ===========================================================================
def build_resolution():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "GOVERNANCE & AUTHORIZATION 03-01",
        "Lumbee Tribal Council Resolution",
        "Authorizing Participation in the Tribal Broadband Connectivity Program "
        "and Execution of the Definitive Agreements",
    )
    C.setup_header_footer(doc, "Tribal Council Resolution")
    C.status_banner(doc)
    C.spacer(doc, 1)

    C.title_line(
        doc,
        f"RESOLUTION NO. {C.PH('resolution number')} OF THE TRIBAL COUNCIL OF "
        f"THE {C.TRIBE_FULL.upper()}",
        size=14,
    )
    C.para(
        doc,
        f"A RESOLUTION authorizing the {C.TRIBE_SHORT} to apply for and accept "
        f"an award under the {C.PROGRAM_FULL} ({C.PROGRAM_SHORT}) of the "
        f"{C.AGENCY_FULL} ({C.AGENCY_SHORT}); to own and steward the "
        f"{C.GRANT_FUNDED}; to accept the {C.FEDERAL_INTEREST} and related "
        f"federal obligations; and to execute the Definitive Agreements with "
        f"{C.OPERATOR_FULL}.",
        italic=True,
    )

    head(doc, "RECITALS")
    intro_recital_parties(doc)
    whereas(
        doc,
        f"the {C.TRIBE_SHORT} has the opportunity to participate in the "
        f"{C.PROGRAM_FULL} ({C.PROGRAM_SHORT}) administered by {C.AGENCY_SHORT} "
        f"within the {C.DOC_FULL}, to expand affordable broadband service in "
        f"{C.GEOGRAPHY}, at service levels of not less than {C.SPEED_FLOOR};",
    )
    whereas(
        doc,
        f"the provisional program baseline for this preliminary package is the "
        f"{C.NOFO_NAME}, and {C.NOFO_LIVE_NOTE}",
    )
    whereas(
        doc,
        f"the {C.TRIBE_SHORT} intends to be the TBCP award recipient, and to "
        f"own, hold, and steward the {C.GRANT_FUNDED} comprising the "
        f"{C.NETWORK}, as the party responsible to {C.AGENCY_SHORT} for the "
        f"award;",
    )
    whereas(
        doc,
        f"the {C.TRIBE_SHORT} intends to partner with {C.OPERATOR_FULL} "
        f"(“{C.OPERATOR_SHORT}”) to design, build, operate, and "
        f"maintain the {C.NETWORK} and to provide retail broadband service, on "
        f"the terms of a suite of Definitive Agreements;",
    )
    whereas(
        doc,
        f"{C.LREMC_FULL} (“{C.LREMC_SHORT}”) separately owns the "
        f"{C.LREMC_ASSETS} (poles, conduit, fiber, easements, huts, power, and "
        f"land) that may be required for the project, and its participation is "
        f"to be limited to a separate owner consent and joinder and does not "
        f"constitute any guarantee, transfer of assets, assumption of debt, or "
        f"duty to finance the project;",
    )
    whereas(
        doc,
        f"the {C.TRIBE_SHORT}'s eligibility to receive a TBCP award turns on "
        f"its federal-recognition status, which presents a threshold question "
        f"under the {C.CITES['lumbee_act']} that must be confirmed. "
        f"{C.FLAG_GRANT} {C.FLAG_ATTORNEY}",
    )
    whereas(
        doc,
        f"acceptance of a TBCP award will subject the {C.GRANT_FUNDED} to the "
        f"{C.FEDERAL_INTEREST} and to the Uniform Guidance at "
        f"{C.CITES['ug_part']}, including {C.CITES['real_property']}, "
        f"{C.CITES['equipment']}, {C.CITES['intangible']}, and "
        f"{C.CITES['prog_income']};",
    )

    head(doc, "RESOLUTIONS")
    C.para(
        doc,
        f"NOW, THEREFORE, BE IT RESOLVED by the Tribal Council of the "
        f"{C.TRIBE_FULL} as follows:",
        bold=True,
    )
    resolved(
        doc,
        f"the {C.TRIBE_SHORT} is authorized to prepare, submit, and (if "
        f"selected) accept an application and award under the {C.PROGRAM_SHORT}, "
        f"consistent with the controlling NOFO and award once confirmed; "
        f"{C.FLAG_GRANT}",
    )
    resolved(
        doc,
        f"the {C.TRIBE_SHORT} is authorized to take, own, hold, and steward the "
        f"{C.GRANT_FUNDED}, and to accept the {C.FEDERAL_INTEREST} and the "
        f"obligations of {C.CITES['ug_part']} applicable to the award, "
        f"including property, procurement, program-income, record-retention, "
        f"and audit obligations; {C.FLAG_GRANT}",
    )
    resolved(
        doc,
        f"the {C.TRIBE_SHORT} is authorized to negotiate and execute the "
        f"{MASTER}, and each of the following nine (9) Definitive Agreements, in "
        f"forms approved by Tribal counsel:",
    )
    for k in OTHER_NINE:
        C.bullet(doc, C.AGREEMENTS[k])
    resolved(
        doc,
        f"the {C.PH('authorized signatory office(s) — e.g., Chairman and/or Tribal Administrator')} "
        f"is/are designated as the authorized signatory(ies) empowered to "
        f"execute and deliver the Definitive Agreements, together with related "
        f"certificates, schedules, and federal submissions, on behalf of the "
        f"{C.TRIBE_SHORT}; {C.FLAG_ATTORNEY}",
    )
    resolved(
        doc,
        f"all approvals granted by this Resolution are LIMITED and expressly "
        f"subject to review and final form approval by Tribal counsel; and any "
        f"limited waiver of sovereign immunity is RESERVED and is NOT granted by "
        f"this Resolution — no such waiver is effective unless and until the "
        f"Tribal Council expressly acts by a separate duly adopted resolution "
        f"(see Governance Document 04, {C.AGREEMENTS['master']} cross-reference "
        f"and the Limited Waiver of Sovereign Immunity alternatives). "
        f"{C.FLAG_ATTORNEY}",
    )
    resolved(
        doc,
        f"the {C.TRIBE_SHORT} is authorized to perform its grant-stewardship "
        f"role and to make all required federal submissions, reports, "
        f"disclosures, and closeout filings to {C.AGENCY_SHORT}/{C.DOC_FULL} in "
        f"accordance with the award and {C.CITES['ug_part']}; {C.FLAG_GRANT}",
    )
    C.flag_para(
        doc,
        C.FLAG_ATTORNEY,
        "Confirm quorum, notice, meeting, and adoption formalities under the "
        "Tribal Constitution and governing documents before reliance.",
    )

    head(doc, "CERTIFICATION")
    C.para(
        doc,
        f"I, the undersigned Secretary of the Tribal Council of the "
        f"{C.TRIBE_FULL}, certify that the foregoing Resolution was duly adopted "
        f"at a {C.PH('regular / special')} meeting of the Tribal Council held on "
        f"{C.PH('meeting date')}, at which a quorum was present and acting "
        f"throughout, by a vote of {C.PH('number in favor')} in favor, "
        f"{C.PH('number opposed')} opposed, and {C.PH('number abstaining')} "
        f"abstaining, and that this Resolution has not been rescinded or "
        f"modified and remains in full force and effect.",
    )
    C.spacer(doc, 1)
    sig(
        doc,
        f"{C.TRIBE_FULL} — Tribal Council",
        [
            {"name": "Secretary name", "title": "Secretary, Lumbee Tribal Council"},
            {"name": "Chairman name", "title": "Chairman, Lumbee Tribal Council (attest)"},
        ],
    )

    disclaimer(doc)
    return C.save(doc, OUTDIR, "01_Lumbee_Tribal_Council_Resolution.docx")


# ===========================================================================
# 2.  RIVR TECH MEMBER AUTHORIZATION
# ===========================================================================
def build_rivr_authorization():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "GOVERNANCE & AUTHORIZATION 03-02",
        "RIVR Tech Member Authorization",
        f"Written Consent of the Members and Managers of {C.OPERATOR_FULL}",
    )
    C.setup_header_footer(doc, "RIVR Tech Member Authorization")
    C.status_banner(doc)
    C.spacer(doc, 1)

    C.title_line(
        doc,
        f"WRITTEN CONSENT OF THE MEMBERS AND MANAGERS OF "
        f"{C.OPERATOR_FULL.upper()}",
        size=13,
    )
    C.para(
        doc,
        f"The undersigned, constituting the members and managers of "
        f"{C.OPERATOR_FULL} (“{C.OPERATOR_SHORT}”), a limited "
        f"liability company, acting by written consent in lieu of a meeting to "
        f"the extent permitted by the company's operating agreement and "
        f"applicable law, hereby adopt the following authorizations. "
        f"{C.FLAG_ATTORNEY}",
    )

    head(doc, "RECITALS")
    intro_recital_parties(doc)
    whereas(
        doc,
        f"{C.OPERATOR_SHORT} intends to serve as the operator/partner of the "
        f"{C.TRIBE_SHORT} for the design, construction, operation, maintenance, "
        f"and retail service of the {C.NETWORK} under the {C.PROGRAM_SHORT};",
    )
    whereas(
        doc,
        f"{C.OPERATOR_SHORT} is a separate entity from {C.LREMC_FULL} "
        f"(“{C.LREMC_SHORT}”), and the {C.LREMC_ASSETS} are the "
        f"subject of a separate consent and joinder by {C.LREMC_SHORT};",
    )
    whereas(
        doc,
        f"the compensation model contemplated by the Definitive Agreements is "
        f"{C.DEAL['per_subscriber_amount']}, over an initial term of "
        f"{C.DEAL['iru_term_recommended']} years from {C.DEAL['term_trigger']} "
        f"plus {C.DEAL['iru_renewal']};",
    )

    head(doc, "AUTHORIZATIONS")
    C.para(doc, "NOW, THEREFORE, the members and managers RESOLVE as follows:",
           bold=True)
    resolved(
        doc,
        f"{C.OPERATOR_SHORT} is authorized to negotiate, execute, deliver, and "
        f"perform the {MASTER} and each of the following nine (9) Definitive "
        f"Agreements:",
    )
    for k in OTHER_NINE:
        C.bullet(doc, C.AGREEMENTS[k])
    resolved(
        doc,
        f"the {C.PH('authorized RIVR Tech signatory office — e.g., Chief Executive Officer / Manager')} "
        f"is designated as the authorized signatory of {C.OPERATOR_SHORT} "
        f"empowered to execute and deliver the Definitive Agreements and related "
        f"certificates and schedules; {C.FLAG_BUSINESS}",
    )
    resolved(
        doc,
        f"{C.OPERATOR_SHORT} confirms it is duly organized, validly existing, "
        f"and in good standing under the laws of {C.STATE} "
        f"({C.PH('state of organization — confirm')}), holds or will obtain the "
        f"authority, registrations, and licenses required to perform, and that "
        f"execution has been duly authorized; {C.FLAG_ATTORNEY}",
    )
    resolved(
        doc,
        f"{C.OPERATOR_SHORT} commits to procure and maintain the insurance and "
        f"any performance/payment bonding required by the Definitive Agreements "
        f"and the applicable insurance schedule (additional-insured, waiver of "
        f"subrogation, and primary/non-contributory endorsements in favor of "
        f"the {C.TRIBE_SHORT}); {C.FLAG_BUSINESS}",
    )
    resolved(
        doc,
        f"{C.OPERATOR_SHORT} accepts and will comply with the applicable "
        f"grant flow-down requirements passed through by the {C.TRIBE_SHORT} "
        f"under {C.CITES['pass_through']} and {C.CITES['subrecipient']}, "
        f"including procurement, domestic-preference/BABA, "
        f"prohibited-telecommunications ({C.CITES['telecom_ban']}), "
        f"record-retention, and audit-access requirements; {C.FLAG_GRANT}",
    )

    head(doc, "EXECUTION")
    C.para(
        doc,
        "The undersigned execute this written consent, which may be signed in "
        "counterparts, as of the date(s) set forth below.",
    )
    C.spacer(doc, 1)
    sig(
        doc,
        C.OPERATOR_FULL,
        [
            {"name": "member / manager name", "title": "Member / Manager"},
            {"name": "additional member / manager (as required by the operating agreement)",
             "title": "Member / Manager"},
        ],
    )

    disclaimer(doc)
    return C.save(doc, OUTDIR, "02_RIVR_Tech_Member_Authorization.docx")


# ===========================================================================
# 3.  LREMC CONSENT AND LIMITED JOINDER
# ===========================================================================
def build_lremc_consent():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "GOVERNANCE & AUTHORIZATION 03-03",
        "LREMC Consent and Limited Joinder",
        f"Limited Consent and Joinder of {C.LREMC_FULL}",
    )
    C.setup_header_footer(doc, "LREMC Consent and Limited Joinder")
    C.status_banner(doc)
    C.spacer(doc, 1)

    C.title_line(
        doc,
        f"CONSENT AND LIMITED JOINDER OF {C.LREMC_FULL.upper()}",
        size=13,
    )
    intro_recital_parties(doc)

    head(doc, "RECITALS")
    whereas(
        doc,
        f"{C.LREMC_FULL} (“{C.LREMC_SHORT}”) is the owner of the "
        f"{C.LREMC_ASSETS}, consisting of the poles, conduit, fiber, easements, "
        f"huts/shelters, electric power facilities, and land more particularly "
        f"identified in {C.PH('LREMC asset/facility schedule reference (Schedule S9)')};",
    )
    whereas(
        doc,
        f"{C.LREMC_SHORT} is a separate entity from {C.OPERATOR_FULL} "
        f"(“{C.OPERATOR_SHORT}”), and nothing herein makes the "
        f"{C.LREMC_ASSETS} the property of, or under the control of, "
        f"{C.OPERATOR_SHORT} or the {C.TRIBE_SHORT};",
    )
    whereas(
        doc,
        f"the {C.TRIBE_SHORT} and {C.OPERATOR_SHORT} require certain access and "
        f"use rights to the {C.LREMC_ASSETS} to deploy and operate the "
        f"{C.NETWORK}, and {C.LREMC_SHORT} is willing to consent and join on a "
        f"limited basis on the terms below;",
    )

    head(doc, "CONSENT AND LIMITED JOINDER")
    C.numbered(
        doc,
        f"Consent. {C.LREMC_SHORT} consents to the use of the {C.LREMC_ASSETS} "
        f"for the {C.NETWORK}, and joins the Definitive Agreements SOLELY as to "
        f"the specifically identified {C.LREMC_ASSETS} and the obligations "
        f"expressly assumed in this instrument, and for no other purpose.",
    )
    C.numbered(
        doc,
        f"Grant of access/use rights. {C.LREMC_SHORT} grants the access, "
        f"attachment, occupancy, and use rights over the {C.LREMC_ASSETS} "
        f"necessary for the deployment, interconnection, operation, and "
        f"maintenance of the {C.NETWORK}, on the specific terms, fees, and "
        f"conditions set out in the {C.AGREEMENTS['land']} and the "
        f"{C.AGREEMENTS['interconnect']}. {C.PH('specific poles / conduit / fiber strands / sites / power terms — confirm')}",
    )
    C.numbered(
        doc,
        f"No guarantee. {C.LREMC_SHORT} does not guarantee, and shall not be "
        f"deemed to guarantee, any obligation, debt, performance, or funding of "
        f"the {C.TRIBE_SHORT} or {C.OPERATOR_SHORT}.",
    )
    C.numbered(
        doc,
        f"No transfer of assets. This instrument conveys no ownership of the "
        f"{C.LREMC_ASSETS}; {C.LREMC_SHORT} retains title, and grants only the "
        f"limited access/use rights expressly stated.",
    )
    C.numbered(
        doc,
        f"No debt assumption. {C.LREMC_SHORT} assumes no indebtedness or "
        f"liability of the {C.TRIBE_SHORT} or {C.OPERATOR_SHORT}, and no "
        f"{C.FEDERAL_INTEREST} attaches to the {C.LREMC_ASSETS} by reason of "
        f"this consent. {C.FLAG_GRANT}",
    )
    C.numbered(
        doc,
        f"No duty to finance. {C.LREMC_SHORT} has no obligation to fund, "
        f"finance, advance costs for, or provide credit support to the project.",
    )
    C.numbered(
        doc,
        f"Scope limitation and reservation. All rights not expressly granted are "
        f"reserved to {C.LREMC_SHORT}; this consent is limited to the "
        f"{C.LREMC_ASSETS} and does not extend to any other {C.LREMC_SHORT} "
        f"assets, systems, or operations. {C.FLAG_ATTORNEY}",
    )
    C.flag_para(
        doc,
        C.FLAG_ATTORNEY,
        f"Confirm {C.LREMC_SHORT} corporate authority, any RUS/lender or "
        f"regulatory consents affecting the {C.LREMC_ASSETS}, and the pole-"
        f"attachment framework, including {C.CITES['nc_pole']}.",
    )

    head(doc, "EXECUTION")
    C.para(
        doc,
        f"IN WITNESS WHEREOF, {C.LREMC_SHORT} has executed this Consent and "
        f"Limited Joinder by its duly authorized officer as of the date below.",
    )
    C.spacer(doc, 1)
    sig(
        doc,
        C.LREMC_FULL,
        [
            {"name": "authorized LREMC officer",
             "title": "e.g., Chief Executive Officer / General Manager"},
            {"name": "attesting officer", "title": "Secretary (attest)"},
        ],
    )

    disclaimer(doc)
    return C.save(doc, OUTDIR, "03_LREMC_Consent_and_Joinder.docx")


# ===========================================================================
# 4.  LIMITED WAIVER OF SOVEREIGN IMMUNITY — ALTERNATIVES FOR COUNSEL
# ===========================================================================
def build_sovereign_immunity():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "GOVERNANCE & AUTHORIZATION 03-04",
        "Limited Waiver of Sovereign Immunity",
        "Alternative Forms for Tribal Counsel Review — Nothing Waived Herein",
    )
    C.setup_header_footer(doc, "Limited Waiver of Sovereign Immunity")
    C.status_banner(doc)
    C.spacer(doc, 1)

    C.title_line(
        doc,
        "LIMITED WAIVER OF SOVEREIGN IMMUNITY — ALTERNATIVE FORMS FOR TRIBAL "
        "COUNSEL REVIEW",
        size=13,
    )
    C.flag_para(
        doc,
        C.FLAG_ATTORNEY,
        "FOR COUNSEL REVIEW ONLY. This document presents drafting alternatives; "
        "it is not itself a waiver and grants nothing.",
    )
    C.para(
        doc,
        f"NOTHING IN THIS DOCUMENT WAIVES, OR SHALL BE DEEMED TO WAIVE, THE "
        f"SOVEREIGN IMMUNITY OF THE {C.TRIBE_FULL}. Tribal sovereign immunity is "
        f"not waived absent a clear and express act; see {C.CITES['sovereign_immunity']}. "
        f"No waiver is effective unless and until the Tribal Council expressly "
        f"acts by a duly adopted resolution adopting a specific form below. "
        f"{C.FLAG_ATTORNEY}",
        bold=True,
    )

    head(doc, "COMMON PARAMETERS (ALL ALTERNATIVES)")
    C.para(
        doc,
        "Any waiver alternative, if adopted, is to be drafted narrowly and tied "
        "to the following parameters, and to nothing beyond them:",
    )
    C.subsection(doc, "a", f"Agreements covered: only the specifically identified "
                 f"Definitive Agreement(s) — "
                 f"{C.PH('list the exact agreement(s), e.g., the Master Agreement only')} "
                 f"— and no others.")
    C.subsection(doc, "b", "Claims covered: only claims arising out of the covered "
                 "agreement(s); consequential, punitive, and non-contractual "
                 "claims excluded unless expressly stated.")
    C.subsection(doc, "c", "Remedies: limited to the specific remedies stated (e.g., "
                 f"specific contract remedies and a monetary cap of "
                 f"{C.PH('remedy/damages cap')}); no execution against Tribal assets "
                 f"beyond those expressly designated.")
    C.subsection(doc, "d", f"EXCLUDED assets: the waiver excludes, and no remedy may "
                 f"reach, any trust or restricted assets/land and any assets "
                 f"subject to the {C.FEDERAL_INTEREST}; recovery is limited to "
                 f"{C.PH('specifically designated non-trust, non-Federal-Interest assets or funds')}. "
                 f"{C.FLAG_GRANT} {C.FLAG_ATTORNEY}")
    C.subsection(doc, "e", "Enforcement source / consideration: the waiver, if any, is "
                 "granted expressly as consideration for the covered agreement(s) "
                 "and is enforceable only per its stated terms.")
    C.subsection(doc, "f", f"Exhaustion of tribal remedies: identify whether exhaustion "
                 f"of {C.TRIBE_SHORT} tribal-court or administrative remedies is a "
                 f"condition precedent to any external forum. "
                 f"{C.PH('required / not required — counsel to specify')} {C.FLAG_ATTORNEY}")

    head(doc, "ALTERNATIVE 1 — NO WAIVER (BASELINE)")
    C.flag_para(doc, C.FLAG_ATTORNEY, "Recommended default absent a compelling, "
                "bounded business reason to waive.")
    C.para(
        doc,
        f"[ALTERNATIVE 1: The {C.TRIBE_SHORT} grants NO waiver of sovereign "
        f"immunity. Disputes are addressed through non-binding escalation, "
        f"and any binding process requires a separate, later express Tribal "
        f"Council act. All defenses, including sovereign immunity and "
        f"exhaustion of tribal remedies, are fully preserved.]",
    )

    head(doc, "ALTERNATIVE 2 — NARROW LIMITED WAIVER (COURT-BASED)")
    C.para(
        doc,
        f"[ALTERNATIVE 2: Solely for claims arising under the covered "
        f"agreement(s), the {C.TRIBE_SHORT} grants a limited, express waiver of "
        f"sovereign immunity to suit in {C.PH('designated forum — e.g., the U.S. District Court, or the Tribal Court, as counsel specifies')}, "
        f"consenting to that forum's jurisdiction for the limited purpose of "
        f"enforcing the covered agreement(s). Remedies are limited to those in "
        f"the Common Parameters; recovery may not reach trust/restricted assets "
        f"or {C.FEDERAL_INTEREST} assets. Exhaustion of tribal remedies is "
        f"{C.PH('a condition precedent / not required')} per counsel. Governing "
        f"law: {C.PH('governing law — counsel to specify')}.] {C.FLAG_ATTORNEY}",
    )

    head(doc, "ALTERNATIVE 3 — ARBITRATION-ONLY LIMITED WAIVER")
    C.para(
        doc,
        f"[ALTERNATIVE 3: Solely for claims arising under the covered "
        f"agreement(s), the {C.TRIBE_SHORT} grants a limited, express waiver "
        f"consenting to binding arbitration administered by "
        f"{C.PH('arbitral body / rules — e.g., AAA Commercial Rules')} seated in "
        f"{C.PH('seat/venue')}, with a limited waiver extending only to "
        f"confirmation and enforcement of the resulting award in "
        f"{C.PH('court for confirmation of award')}. Remedies and asset "
        f"exclusions (trust/restricted and {C.FEDERAL_INTEREST} assets) are per "
        f"the Common Parameters. Exhaustion of tribal remedies is "
        f"{C.PH('a condition precedent / not required')} per counsel.] "
        f"{C.FLAG_ATTORNEY}",
    )

    head(doc, "RESERVATION")
    C.para(
        doc,
        f"Until the Tribal Council expressly adopts a specific alternative by "
        f"duly adopted resolution, the {C.TRIBE_SHORT} reserves all rights and "
        f"defenses and waives nothing. Cross-reference: the Tribal Council "
        f"Resolution (Governance Document 01) RESERVES this matter. "
        f"{C.FLAG_ATTORNEY}",
        bold=True,
    )

    disclaimer(doc)
    return C.save(doc, OUTDIR, "04_Limited_Waiver_of_Sovereign_Immunity.docx")


# ===========================================================================
# 5.  TAX / TERO / PERMITTING / REGULATORY SCHEDULE  (xlsx — table-suited)
# ===========================================================================
def build_tax_tero_schedule():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Tax, TERO, Permitting, and Regulatory Obligations Schedule",
        "Governance & Authorization 03-05 — preliminary register of known/potential "
        "Tribal, federal, state, and local obligations for the Lumbee Tribe / RIVR Tech "
        "TBCP fiber project.",
        "",
        "## How to use",
        "• The Obligations sheet lists each obligation with its level, authority, and "
        "applicability status.",
        "• Every applicability determination is marked [COUNSEL TO CONFIRM] and must be "
        "confirmed by qualified Tribal, tax, and regulatory counsel before reliance.",
        "• [BUSINESS DECISION REQUIRED] marks items with financial exposure to be sized "
        "with advisors. [GRANT COMPLIANCE REVIEW REQUIRED] marks award/flow-down items.",
        "• The Change-in-Law sheet states the economic-adjustment mechanism.",
        "",
        "## Status",
        C.DRAFT_STATUS,
        C.NOFO_LIVE_NOTE,
    ])

    ws = wb.create_sheet("Obligations")
    C.xl_title(ws, "Tax, TERO, Permitting & Regulatory Obligations",
               subtitle="Preliminary register — all applicability items [COUNSEL TO CONFIRM]",
               span=7)
    headers = ["#", "Category", "Obligation / Item", "Level",
               "Authority / Citation", "Applicability / Status", "Notes"]
    C.xl_header_row(ws, 5, headers)
    widths = [4, 20, 34, 12, 46, 30, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[C.get_column_letter(i + 1)].width = w

    rows = [
        ("Tax", "Possessory-interest tax on use of exempt/public property",
         "State/Local",
         "N.C. property-tax law on possessory interests in exempt property",
         "[COUNSEL TO CONFIRM]",
         "Applies to a private operator's beneficial use of tax-exempt/grant assets; confirm NC treatment."),
        ("Tax", "Ad valorem / real & personal property tax",
         "State/Local",
         "N.C. Gen. Stat. Ch. 105 (Machinery Act) [COUNSEL TO CONFIRM]",
         "[COUNSEL TO CONFIRM]",
         "Confirm exemptions for Tribally/grant-owned Grant-Funded Assets vs. operator-owned equipment."),
        ("Tax", "Sales & use tax on equipment/construction",
         "State/Local",
         "N.C. Gen. Stat. Ch. 105, Art. 5 (sales & use) [COUNSEL TO CONFIRM]",
         "[COUNSEL TO CONFIRM]",
         "Confirm any exemption/refund for grant-funded telecom infrastructure; [GRANT COMPLIANCE REVIEW REQUIRED]."),
        ("TERO / Preference", "Tribal Employment Rights / Tribal preference",
         "Tribal",
         C.CITES['tero'],
         "[COUNSEL TO CONFIRM]",
         "Confirm whether a TERO/tribal-preference ordinance applies to the Lumbee Tribe and this project."),
        ("Permits", "Building, electrical, and construction permits",
         "State/Local",
         "N.C. State Building Code and local permitting [COUNSEL TO CONFIRM]",
         "[COUNSEL TO CONFIRM]",
         "Per-jurisdiction; coordinate with make-ready and construction schedule."),
        ("Permits", "Underground utility locates (811)",
         "State",
         C.CITES['nc_dig'],
         "[COUNSEL TO CONFIRM]",
         "Damage-prevention / 811 locate compliance during construction."),
        ("Environmental / Historic", "NEPA and NHPA Section 106 review",
         "Federal",
         f"{C.CITES['nepa']}; {C.CITES['nhpa']}",
         "[COUNSEL TO CONFIRM]",
         "Award-condition environmental/historic review; [GRANT COMPLIANCE REVIEW REQUIRED]."),
        ("Telecom / FCC", "FCC authorizations / USAC / ETC designation",
         "Federal",
         C.CITES['usac_lifeline'],
         "[COUNSEL TO CONFIRM]",
         "Confirm ETC designation, Lifeline, and USAC obligations for the retail provider of record."),
        ("Telecom / Privacy", "CPNI obligations",
         "Federal",
         C.CITES['cpni'],
         "[COUNSEL TO CONFIRM]",
         "Customer proprietary network information handling for the retail operator."),
        ("Franchise / ROW", "Right-of-way, franchise, and pole attachment",
         "State/Local",
         f"{C.CITES['nc_pole']}; {C.CITES['nc_ncdot']}",
         "[COUNSEL TO CONFIRM]",
         f"Coordinate with LREMC ({C.LREMC_SHORT}) pole/facility consent and NCDOT encroachment."),
        ("Land (if trust/restricted)", "Leasing / ROW / BIA approval on Indian land",
         "Federal",
         f"{C.CITES['indian_leasing']}; {C.CITES['indian_row']}; {C.CITES['bia_approval']}",
         "[COUNSEL TO CONFIRM]",
         "Applies only if trust/restricted land is used or crossed; confirm status."),
        ("Eligibility", "Federal-recognition status affecting TBCP eligibility",
         "Federal/Tribal",
         C.CITES['lumbee_act'],
         "[COUNSEL TO CONFIRM]",
         "Threshold eligibility question; [GRANT COMPLIANCE REVIEW REQUIRED]."),
    ]
    r = 6
    for i, row in enumerate(rows, start=1):
        C.xl_cell(ws, r, 1, i, kind="text", align="center")
        C.xl_cell(ws, r, 2, row[0], kind="section" if False else "text", wrap=True)
        C.xl_cell(ws, r, 3, row[1], wrap=True)
        C.xl_cell(ws, r, 4, row[2], align="center", wrap=True)
        C.xl_cell(ws, r, 5, row[3], wrap=True)
        C.xl_cell(ws, r, 6, row[4], kind="input", wrap=True)
        C.xl_cell(ws, r, 7, row[5], wrap=True)
        r += 1
    C.xl_legend(ws, r + 1)

    # Change-in-law / economic-adjustment mechanism sheet
    ws2 = wb.create_sheet("Change-in-Law")
    C.xl_title(ws2, "Change-in-Law / Economic-Adjustment Mechanism", span=2)
    ws2.column_dimensions["A"].width = 4
    ws2.column_dimensions["B"].width = 120
    mech = [
        "## Purpose",
        "Allocate the risk of new, changed, or newly-applied taxes, fees, TERO/preference "
        "charges, permits, or regulatory obligations arising after the Effective Date.",
        "",
        "## Mechanism",
        "• Trigger: enactment, amendment, repeal, or change in interpretation/enforcement of any "
        "law, tax, fee, or regulatory requirement after the Effective Date that materially changes "
        "a Party's cost or obligations (a “Change in Law”).",
        "• Notice: the affected Party gives written notice with supporting detail within "
        "[BUSINESS DECISION REQUIRED: e.g., 30] days of becoming aware.",
        "• Adjustment: the Parties negotiate in good faith an equitable adjustment to the "
        "per-Active-Subscriber payment (" + C.DEAL['per_subscriber_amount'] + ") or other affected "
        "terms to restore the pre-change economic balance. [BUSINESS DECISION REQUIRED]",
        "• Grant limits: no adjustment may authorize any act inconsistent with the Award or "
        "applicable federal law; program-income and cost treatment follow the controlling Award. "
        "[GRANT COMPLIANCE REVIEW REQUIRED]",
        "• Dispute: unresolved adjustments follow the dispute-resolution provisions of the "
        "Master Agreement. [ATTORNEY REVIEW REQUIRED]",
        "",
        "## Status",
        C.DRAFT_STATUS_SHORT,
    ]
    rr = 5
    for ln in mech:
        c = ws2.cell(row=rr, column=2, value=ln[2:].strip() if ln.startswith("##") else ln)
        if ln.startswith("##"):
            c.font = C.Font(bold=True, size=11, color="1F3864")
        else:
            c.font = C.Font(size=10)
        c.alignment = C.Alignment(wrap_text=True, vertical="top")
        rr += 1

    return C.xl_save(wb, OUTDIR, "05_Tax_TERO_Permitting_and_Regulatory_Schedule.xlsx")


# ===========================================================================
# 6.  INSURANCE SCHEDULE  (xlsx — table-suited)
# ===========================================================================
def build_insurance_schedule():
    ins = C.DEAL["insurance"]
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Insurance Schedule",
        "Governance & Authorization 03-06 — required coverages and minimum limits for the "
        "Lumbee Tribe / RIVR Tech TBCP fiber project.",
        "",
        "## How to use",
        "• The Coverages sheet lists each required coverage and a PLACEHOLDER minimum limit.",
        "• All limits are placeholders pending broker/risk-advisor confirmation "
        "[BUSINESS DECISION REQUIRED]; they are canonical defaults from the deal engine, not final.",
        "• The Endorsements sheet lists required endorsements and certificate/notice terms.",
        "",
        "## Status",
        C.DRAFT_STATUS,
    ])

    ws = wb.create_sheet("Coverages")
    C.xl_title(ws, "Required Insurance Coverages & Minimum Limits",
               subtitle="All limits are PLACEHOLDERS pending broker advice [BUSINESS DECISION REQUIRED]",
               span=4)
    C.xl_header_row(ws, 5, ["#", "Coverage", "Minimum Limit (placeholder)", "Basis / Notes"])
    for i, w in enumerate([4, 40, 34, 60]):
        ws.column_dimensions[C.get_column_letter(i + 1)].width = w

    cov = [
        ("Commercial General Liability (CGL)",
         f"{ins['cgl_occurrence']} per occurrence / {ins['cgl_aggregate']} aggregate",
         "Bodily injury, property damage, products/completed operations."),
        ("Automobile Liability", ins["auto"],
         "Owned, hired, and non-owned autos; combined single limit."),
        ("Umbrella / Excess Liability", ins["umbrella"],
         "Excess over CGL, auto, and employers' liability."),
        ("Workers' Compensation", ins["workers_comp"],
         "Statutory limits per N.C. law."),
        ("Employers' Liability", ins["employers_liability"],
         "Each accident / disease."),
        ("Professional / Technology E&O", ins["professional_tech_eo"],
         "Errors & omissions for design, engineering, and technology services."),
        ("Cyber Liability", ins["cyber"],
         "Network security, privacy, breach response; coordinate with the privacy addendum."),
        ("Property / Builders' Risk", ins["property_builders_risk"],
         "Builders' risk during construction of the Grant-Funded Assets; then property coverage."),
        ("Pollution / Environmental", ins["pollution"],
         "Contractor's pollution liability for construction operations."),
    ]
    r = 6
    for i, (name, limit, note) in enumerate(cov, start=1):
        C.xl_cell(ws, r, 1, i, align="center")
        C.xl_cell(ws, r, 2, name, wrap=True, bold=True)
        C.xl_cell(ws, r, 3, limit, kind="input", wrap=True)
        C.xl_cell(ws, r, 4, note, wrap=True)
        r += 1
    C.xl_legend(ws, r + 1)

    ws2 = wb.create_sheet("Endorsements")
    C.xl_title(ws2, "Required Endorsements & Certificate Terms", span=2)
    ws2.column_dimensions["A"].width = 4
    ws2.column_dimensions["B"].width = 120
    endts = [
        "## Required endorsements (in favor of the Lumbee Tribe)",
        f"• Additional Insured: the {C.TRIBE_FULL} (and, as applicable, {C.LREMC_SHORT} "
        "and the United States/NTIA) named as additional insured on CGL, auto, umbrella, "
        "and (where available) pollution and property.",
        "• Waiver of Subrogation: in favor of the Tribe on all policies to the extent permitted.",
        "• Primary and Non-Contributory: the operator's coverage is primary and "
        "non-contributory to any coverage carried by the Tribe.",
        "• Builders' Risk: covering the full replacement cost of the Grant-Funded Assets "
        "during construction, with the Tribe as insured/loss payee as its interest appears.",
        "",
        "## Certificates and notice",
        "• Certificates of insurance delivered before commencement and on each renewal.",
        "• Notice of cancellation or material change: [BUSINESS DECISION REQUIRED: e.g., 30] "
        "days' written notice to the Tribe (10 days for non-payment).",
        "• Insurers rated [BUSINESS DECISION REQUIRED: e.g., A- VII or better] by A.M. Best.",
        "",
        "## Grant coordination",
        f"• Insurance on the Grant-Funded Assets must be consistent with {C.CITES['insurance']}. "
        "[GRANT COMPLIANCE REVIEW REQUIRED]",
        "",
        "## Status",
        C.DRAFT_STATUS_SHORT,
    ]
    rr = 5
    for ln in endts:
        c = ws2.cell(row=rr, column=2, value=ln[2:].strip() if ln.startswith("##") else ln)
        if ln.startswith("##"):
            c.font = C.Font(bold=True, size=11, color="1F3864")
        else:
            c.font = C.Font(size=10)
        c.alignment = C.Alignment(wrap_text=True, vertical="top")
        rr += 1

    return C.xl_save(wb, OUTDIR, "06_Insurance_Schedule.xlsx")


# ===========================================================================
# 7.  SERVICE LEVEL SCHEDULE  (docx)
# ===========================================================================
def build_service_level_schedule():
    s = C.DEAL["sla"]
    doc = C.new_doc()
    C.add_cover(
        doc,
        "GOVERNANCE & AUTHORIZATION 03-07",
        "Service Level Schedule",
        f"Service Levels, Measurement, Reporting, and Remedies — "
        f"cross-referenced to the {C.AGREEMENTS['om']}",
    )
    C.setup_header_footer(doc, "Service Level Schedule")
    C.status_banner(doc)
    C.spacer(doc, 1)

    C.title_line(doc, "SERVICE LEVEL SCHEDULE", size=14)
    C.para(
        doc,
        f"This Schedule sets the service levels for operation and maintenance "
        f"of the {C.NETWORK}. It is cross-referenced to, and governed by, the "
        f"{C.AGREEMENTS['om']} and Schedule {C.PH('S6')} "
        f"({C.SCHEDULES['S6']}). In any conflict, the {C.AGREEMENTS['om']} and "
        f"the controlling Award govern.",
    )

    head(doc, "1.  DEFINITIONS")
    C.subsection(doc, "a", "“Incident” means any unplanned interruption or "
                 "degradation of service or a network fault reported or detected.")
    C.subsection(doc, "b", "“Response Time” means the interval from Incident "
                 "detection/report to acknowledgment by the operator's NOC.")
    C.subsection(doc, "c", "“Dispatch Time” means the interval from acknowledgment "
                 "to dispatch of qualified field personnel where a truck roll is required.")
    C.subsection(doc, "d", "“Restore/Resolve Time” means the interval from detection "
                 "to restoration of service or resolution of the Incident.")
    C.subsection(doc, "e", f"“Availability” means the percentage of time the "
                 f"{C.NETWORK} is available in a calendar month, excluding Excused Outages.")
    C.subsection(doc, "f", "“NOC” means the network operations center, staffed "
                 f"{s['noc']}.")

    head(doc, "2.  SEVERITY MATRIX")
    C.add_table(
        doc,
        ["Priority", "Definition", "Response", "Dispatch", "Restore / Resolve"],
        [
            ["P1 — Critical",
             "Total outage / multiple subscribers or critical facility down; public-safety impact",
             f"{s['P1_response_min']} min", f"{s['P1_dispatch_hr']} hr", f"{s['P1_restore_hr']} hr"],
            ["P2 — Major",
             "Significant degradation or single-site outage; partial service loss",
             f"{s['P2_response_min']} min", f"{s['P2_dispatch_hr']} hr", f"{s['P2_restore_hr']} hr"],
            ["P3 — Minor",
             "Isolated/limited-impact fault; service degraded but usable",
             f"{s['P3_response_hr']} hr", f"{s['P3_dispatch_hr']} hr", f"{s['P3_restore_days']} days"],
            ["P4 — Low",
             "Cosmetic, informational, or scheduled work; no service impact",
             f"{s['P4_response_hr']} hr", "as scheduled", f"{s['P4_restore_days']} days"],
        ],
        widths=[1.1, 2.7, 0.9, 0.9, 1.1],
        col_align=[None, None, "c", "c", "c"],
    )

    head(doc, "3.  PERFORMANCE TARGETS")
    C.add_table(
        doc,
        ["Metric", "Target", "Notes"],
        [
            ["Network Availability", s["availability_target"],
             "Monthly, excluding Excused Outages."],
            ["Latency (round-trip)", f"≤ {s['latency_ms']} ms",
             "Measured across the defined core path."],
            ["Packet Loss", f"≤ {s['packet_loss']}", "Monthly average."],
            ["Jitter", f"≤ {s['jitter_ms']} ms", "Monthly average."],
            ["NOC Coverage", s["noc"], "Staffed monitoring and incident intake."],
        ],
        widths=[1.8, 1.6, 3.0],
        col_align=[None, "c", None],
    )
    C.flag_para(doc, C.FLAG_TECH, "Confirm measurement reference points, test paths, "
                "and instrumentation with the engineering exhibits.")

    head(doc, "4.  MEASUREMENT METHODS")
    C.bullet(doc, "Availability, latency, packet loss, and jitter are measured by the "
             "operator's monitoring systems at agreed reference points, on a calendar-month basis.")
    C.bullet(doc, "Incident timestamps are taken from the NOC ticketing system; clocks "
             "synchronized to a common time source.")
    C.bullet(doc, f"Methods and reference points are defined in the engineering exhibits and "
             f"Schedule {C.PH('S6')}. {C.FLAG_TECH}")

    head(doc, "5.  EXCLUSIONS (EXCUSED OUTAGES)")
    C.bullet(doc, "Scheduled/emergency maintenance within an agreed maintenance window and notice.")
    C.bullet(doc, "Force majeure events.")
    C.bullet(doc, f"Outages caused by the {C.TRIBE_SHORT}, subscribers, or third parties not "
             f"under the operator's control, including third-party power or "
             f"{C.LREMC_SHORT}-owned facility failures outside the operator's responsibility.")
    C.bullet(doc, "Subscriber premises equipment or inside-wiring faults.")

    head(doc, "6.  REPORTING")
    C.bullet(doc, "Monthly service-level report within [BUSINESS DECISION REQUIRED: e.g., 10] "
             "business days of month-end, covering each metric, Incident log, and credits due.")
    C.bullet(doc, "Prompt notification of P1/P2 Incidents to the Tribe's designated contact.")
    C.bullet(doc, "Quarterly review meeting and trend analysis.")

    head(doc, "7.  CHRONIC-FAILURE RULES")
    C.para(
        doc,
        "A chronic failure occurs when the operator misses the same service level "
        "for [BUSINESS DECISION REQUIRED: e.g., three (3)] consecutive months, or for "
        "[BUSINESS DECISION REQUIRED: e.g., any four (4)] months in a rolling twelve (12)-month "
        "period.",
    )
    C.bullet(doc, "Chronic failure triggers a root-cause analysis and a corrective-action plan.")
    C.bullet(doc, f"Uncured chronic failure is an escalation/step-in and default trigger under the "
             f"{C.AGREEMENTS['om']} and the {C.AGREEMENTS['transition']}. {C.FLAG_ATTORNEY}")

    head(doc, "8.  REMEDIES / SERVICE CREDITS")
    C.add_table(
        doc,
        ["Missed Level", "Service Credit (placeholder)"],
        [
            ["Availability below target",
             f"{C.PH('X')}% of the monthly per-Active-Subscriber payment per affected subscriber"],
            ["P1 restore-time miss", f"{C.PH('X')}% credit per missed Incident"],
            ["P2 restore-time miss", f"{C.PH('X')}% credit per missed Incident"],
            ["Chronic failure", f"escalated credit and/or step-in per {C.AGREEMENTS['transition']}"],
        ],
        widths=[2.6, 3.8],
    )
    C.flag_para(doc, C.FLAG_BUSINESS, "Credit percentages, caps, and whether credits are the "
                "sole monetary remedy are business decisions to be finalized.")
    C.para(
        doc,
        f"Service credits are calculated against the {C.DEAL['per_subscriber_amount']} model and "
        f"reconciled in the monthly report. Nothing in this Schedule authorizes any act "
        f"inconsistent with the Award or applicable federal law. {C.FLAG_GRANT}",
    )

    disclaimer(doc)
    return C.save(doc, OUTDIR, "07_Service_Level_Schedule.docx")


# ===========================================================================
def main():
    builders = [
        build_resolution,
        build_rivr_authorization,
        build_lremc_consent,
        build_sovereign_immunity,
        build_tax_tero_schedule,
        build_insurance_schedule,
        build_service_level_schedule,
    ]
    paths = []
    for b in builders:
        p = b()
        paths.append(p)
        print("WROTE:", p)
    print("\n%d files generated." % len(paths))
    return paths


if __name__ == "__main__":
    main()
