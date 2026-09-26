"""
gen_phase1.py — Generator for the PHASE 1 DILIGENCE set (9 deliverables) of the
Lumbee Tribe of North Carolina / RIVR Tech TBCP fiber-project transaction package.

SINGLE GENERATOR. Uses the canonical engine Scripts/common.py for every party
name, defined term, citation, deal-mechanic, placeholder and flag so the whole
Phase 1 set stays internally consistent.

No source documents were supplied. This is a PRELIMINARY draft keyed to the
provisional March 2024 amended Round 2 NOFO baseline (Round 3 is the live round
as of the version date). Nothing is invented: unknowns are C.PH() placeholders
and open questions carry the appropriate C.FLAG_* marker.
"""

import sys
sys.path.insert(0, '/home/user/claude-skills/Lumbee_RIVR_TBCP_Agreement_Package/Scripts')
import common as C

from openpyxl import Workbook

SUBDIR = "01_Phase1_Diligence"

# ---------------------------------------------------------------------------
# Shared boilerplate
# ---------------------------------------------------------------------------
DISCLAIMER = (
    "These drafts are provided for business planning and attorney review only. They are "
    "preliminary, non-binding, incomplete, and prepared without any controlling TBCP award "
    "or source documents. The provisional baseline is the March 2024 amended Round 2 NOFO; "
    "Round 3 is the live TBCP round as of the version date, and the controlling round and "
    "executed award must be confirmed before anything here is finalized. Every bracketed item, "
    "highlighted placeholder, and flagged provision requires confirmation by qualified counsel, "
    "grant, financial, and technical advisors. Nothing in these drafts constitutes legal, "
    "financial, tax, or grant-compliance advice, creates any binding obligation, or waives any "
    "right, remedy, or sovereign immunity. In the event of any conflict, the executed TBCP award, "
    "applicable federal and state law, and any required governmental consent control; nothing in "
    "these drafts may be relied upon to supersede the TBCP award, applicable law, or required "
    "governmental consent."
)

DOC_INTRO_NOTE = (
    "No executed TBCP award, CD-450 Financial Assistance Award, Notice of Funding Opportunity "
    "PDF, Specific Award Conditions, Tribal resolution, title/easement record, or existing "
    "RIVR Tech / LREMC contract was supplied for this work product. All conclusions are "
    "provisional and keyed to the assumptions register in Document 01."
)


def doc_start(number, title, subtitle, short):
    doc = C.new_doc()
    C.add_cover(doc, number, title, subtitle)
    C.setup_header_footer(doc, short)
    C.add_toc(doc)
    C.status_banner(doc)
    C.spacer(doc, 1)
    return doc


def doc_close(doc, n):
    C.article(doc, n, "Closing Disclaimer and Reservation")
    C.para(doc, DISCLAIMER, italic=True)
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "This document is attorney work product prepared for the common interest of the "
                f"{C.TRIBE_SHORT} and {C.OPERATOR_SHORT}; it is not final legal advice and must be "
                "reviewed and completed by qualified counsel before use.")


# ===========================================================================
# 1. 01_Source_and_Assumptions_Register.xlsx
# ===========================================================================
def build_01():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Source and Assumptions Register — Phase 1 Diligence",
        "Document 01 of the Phase 1 Diligence set.",
        DOC_INTRO_NOTE,
        "",
        "## How to use this workbook",
        "• Sheet 'Source Register' lists every document that WOULD be reviewed to draft and finalize the transaction. "
        "Because no source documents were supplied, nearly every item is MISSING or CONFIRM.",
        "• Sheet 'Assumptions' lists every working assumption used in the Phase 1 drafts. Each is tagged "
        "ASSUMPTION - NOT FACT and must be confirmed against the controlling award before finalizing.",
        "• Status legend: CONTROLLING = governs and must be obtained; INFORMATIVE = useful context; "
        "MISSING = not supplied and required; CONFIRM = referenced but version/applicability unverified.",
        "• Placeholders shown as [# ...] mark data that was not supplied and must not be invented.",
        "",
        "## Precedence reminder",
        C.PRECEDENCE_RULE,
    ], title="Instructions")

    # ---- Source Register sheet ----
    ws = wb.create_sheet("Source Register")
    C.xl_title(ws, "Source Document Register — Documents That Would Be Reviewed",
               "Status is MISSING/CONFIRM for nearly all items because no source documents were supplied.", span=6)
    widths = [42, 26, 20, 16, 34, 60]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    hdr_row = 5
    C.xl_header_row(ws, hdr_row, ["Item", "Type", "Date/Version", "Status", "Where Used", "Notes"])

    rows = [
        # (Item, Type, Date/Version, Status, Where Used, Notes)
        ("Executed Federal Award / Form CD-450 (Financial Assistance Award)", "Grant / award", C.PH("award date - not supplied"),
         "MISSING", "All documents; order of precedence", "Controlling instrument. Award number, amount, period NOT supplied - do not invent. " + C.CITES["sac"]),
        ("Specific Award Conditions (SAC)", "Grant / award terms", C.PH("SAC version"),
         "MISSING", "Docs 03, 04, 05, 07, 08, 09", "May restrict IRU, disposition, program income, and operator role. " + C.CITES["sac"]),
        ("NTIA Standard Terms and Conditions", "Grant / award terms", C.PH("effective version"),
         "CONFIRM", "Docs 03, 04, 05, 09", C.CITES["sac"]),
        ("DOC Financial Assistance Standard Terms and Conditions", "Grant / award terms", C.PH("effective version"),
         "CONFIRM", "Docs 03, 04, 05", C.CITES["sac"]),
        ("Controlling Notice of Funding Opportunity (NOFO)", "Program rules", "Round 2 amended Mar 2024 (PROVISIONAL); Round 3 live Sept 2026",
         "CONFIRM", "All documents", C.NOFO_LIVE_NOTE + " " + C.CITES["nofo"]),
        ("NOFO amendments / addenda", "Program rules", C.PH("amendment list"),
         "MISSING", "Docs 03, 05, 08", "Confirm all amendments to the controlling round."),
        ("TBCP Infrastructure Deployment Application Guidance", "Program guidance", C.PH("version matching award"),
         "CONFIRM", "Docs 02, 05, 06, 07", C.CITES["id_guidance"]),
        ("Approved application (technical & programmatic narrative)", "Application record", C.PH("as approved"),
         "MISSING", "Docs 02, 06, 07, 08", "Defines approved scope; NOT supplied."),
        ("Approved budget and budget narrative", "Application record", C.PH("as approved"),
         "MISSING", "Docs 03, 05, 07", "Basis for reimbursement and cost allocation. " + C.CITES["allowable"]),
        ("Approved network design, routes, and GIS/KMZ", "Engineering record", C.PH("as approved"),
         "MISSING", "Docs 02, 06, 07", "Routes/prices/useful life NOT supplied - do not invent."),
        ("Approved eligible-locations / homes-passed list", "Application record", C.PH("as approved"),
         "MISSING", "Docs 02, 06, 07, 08", "Approved Project Area definition. " + C.DEAL["ph_homes_passed"]),
        ("Approved service commitments (speed, affordability)", "Application record", C.PH("as approved"),
         "MISSING", "Docs 02, 08", f"Service floor assumed {C.SPEED_FLOOR}; confirm against award."),
        ("Tribal Council authorizing resolution(s)", "Governance / authority", C.PH("resolution no./date"),
         "MISSING", "Docs 02, 06, 08; Schedule S1", "Authority of Tribal signatory NOT supplied."),
        ("Tribal designation of entity/instrumentality (if any)", "Governance / authority", C.PH("if applicable"),
         "MISSING", "Docs 06, 08", C.TRIBE_ENTITY_ALT),
        ("RIVR Tech (LREMC Technologies, LLC) formation & governance docs", "Corporate", C.PH("charter/operating agreement"),
         "MISSING", "Docs 04, 05, 06", "Confirm authority, ownership, affiliation with LREMC."),
        ("LREMC corporate authority / board consent", "Corporate / affiliate", C.PH("board consent"),
         "MISSING", "Docs 06, 07; Schedule S9", "Separate owner of poles/fiber/easements/huts/power/land."),
        ("Procurement record for RIVR Tech selection", "Procurement", C.PH("solicitation/award file"),
         "MISSING", "Docs 04, 05", C.CITES["procurement"]),
        ("Existing RIVR Tech network as-builts, inventory & valuation", "Asset / valuation", C.PH("inventory & valuation"),
         "MISSING", "Docs 05, 07", "Pre-existing assets stay RIVR Tech's; value/cost-allocate - never free. " + C.CITES["allowable"]),
        ("Existing RIVR Tech customer & service contracts", "Contracts", C.PH("contract list"),
         "MISSING", "Docs 06, 07, 08", "Provider-of-record and non-project use analysis."),
        ("LREMC pole attachment agreement(s)", "Land / access", C.PH("agreement date"),
         "MISSING", "Docs 06, 07; Schedule S9", C.CITES["nc_pole"]),
        ("LREMC easement / right-of-way records", "Land / access", C.PH("record set"),
         "MISSING", "Docs 07; Schedule S9", C.CITES["indian_row"] + " (if Indian land)."),
        ("LREMC land / hut / site leases", "Land / access", C.PH("lease set"),
         "MISSING", "Docs 07; Schedule S9", C.CITES["indian_leasing"] + " (if trust/restricted land)."),
        ("Title / real-property records", "Land / title", C.PH("title set"),
         "MISSING", "Docs 07; Schedule S9", C.CITES["real_property"]),
        ("Power / electrical service agreements (huts/cabinets)", "Utility", C.PH("service agreements"),
         "MISSING", "Docs 06, 07; Schedule S9", "LREMC/utility-owned; confirm consideration."),
        ("NEPA / environmental review records", "Environmental", C.PH("EA/CATEX record"),
         "MISSING", "Docs 03, 06", C.CITES["nepa"]),
        ("NHPA Section 106 consultation records", "Historic preservation", C.PH("Sec.106 record"),
         "MISSING", "Docs 03, 06", C.CITES["nhpa"]),
        ("Insurance policies / certificates (both parties)", "Insurance", C.PH("policies/COIs"),
         "MISSING", "Docs 02, 03; Schedule S12", C.CITES["insurance"]),
        ("BABA compliance documentation / waivers", "Supply chain", C.PH("BABA record"),
         "MISSING", "Docs 03, 05", C.CITES["baba"]),
        ("Section 889 covered-equipment attestations", "Supply chain", C.PH("attestations"),
         "MISSING", "Docs 03, 05", C.CITES["telecom_ban"]),
        ("SAM.gov registration / UEI / active registration", "Eligibility", C.PH("UEI / status"),
         "CONFIRM", "Docs 03, 04", "Confirm active registration and exclusions check."),
        ("Debarment / suspension (SAM exclusions) checks", "Eligibility", C.PH("check date"),
         "MISSING", "Docs 04, 05", C.CITES["debarment"]),
        ("Interconnection / transport / transit agreements", "Network", C.PH("agreement set"),
         "MISSING", "Docs 06, 07", "Middle-mile / IP transit; RIVR Tech Existing Network."),
        ("Billing / OSS / BSS software licenses", "Systems", C.PH("license set"),
         "MISSING", "Docs 06, 07", "Confirm license scope, data ownership, portability."),
        ("Customer data & privacy policies (CPNI)", "Privacy", C.PH("policies"),
         "MISSING", "Docs 06, 08", C.CITES["cpni"]),
        ("Cybersecurity program / System Security Plan", "Cybersecurity", C.PH("SSP"),
         "MISSING", "Docs 02, 06", "Supply-chain risk and incident response."),
        ("Lumbee federal-recognition status documentation", "Eligibility (threshold)", C.PH("status confirmation"),
         "CONFIRM", "Docs 02, 03, 09", C.CITES["lumbee_act"] + " - threshold TBCP eligibility issue."),
        ("TERO ordinance / Tribal preference requirements", "Tribal / labor", C.PH("ordinance"),
         "CONFIRM", "Docs 02, 05", C.CITES["tero"]),
        ("NC state regulatory (pole attach / NCDOT encroachment)", "State regulatory", C.PH("filings/permits"),
         "CONFIRM", "Docs 06, 07", C.CITES["nc_pole"] + "; " + C.CITES["nc_ncdot"] + "; " + C.CITES["nc_dig"]),
        ("Prior / related federal awards (program-income history)", "Grant history", C.PH("award history"),
         "CONFIRM", "Docs 03, 08", C.CITES["prog_income"]),
        ("Single-audit reports (recipient)", "Audit", C.PH("most recent"),
         "CONFIRM", "Docs 03, 04", C.CITES["single_audit"]),
    ]
    r = hdr_row + 1
    for item, typ, dv, status, used, notes in rows:
        C.xl_cell(ws, r, 1, item, "text", wrap=True)
        C.xl_cell(ws, r, 2, typ, "text", wrap=True)
        C.xl_cell(ws, r, 3, dv, "input", wrap=True)
        kind = "output" if status in ("MISSING", "CONFIRM") else "formula"
        C.xl_cell(ws, r, 4, status, kind, bold=True, align="center")
        C.xl_cell(ws, r, 5, used, "text", wrap=True)
        C.xl_cell(ws, r, 6, notes, "text", wrap=True)
        r += 1
    C.xl_legend(ws, r + 1)
    C.xl_cell(ws, r + 3, 1, "Counts (of listed items):", "text", bold=True)
    C.xl_cell(ws, r + 3, 2, f"=COUNTIF(D{hdr_row+1}:D{r-1},\"MISSING\")", "formula", bold=True)
    C.xl_cell(ws, r + 3, 3, "MISSING", "text")
    C.xl_cell(ws, r + 3, 4, f"=COUNTIF(D{hdr_row+1}:D{r-1},\"CONFIRM\")", "formula", bold=True)
    C.xl_cell(ws, r + 3, 5, "CONFIRM", "text")

    # ---- Assumptions sheet ----
    wa = wb.create_sheet("Assumptions")
    C.xl_title(wa, "Working Assumptions Register — ASSUMPTION, NOT FACT",
               "Every row is a working assumption used in the Phase 1 drafts; each must be confirmed against the controlling award.", span=6)
    awidths = [8, 46, 40, 22, 30, 44]
    for i, w in enumerate(awidths):
        wa.column_dimensions[chr(65 + i)].width = w
    ahdr = 5
    C.xl_header_row(wa, ahdr, ["ID", "Working Assumption", "Rationale / Basis", "Tag", "Confirm With", "Impact if Wrong"])
    assumptions = [
        ("A1", f"Parties are the {C.TRIBE_FULL} (recipient/steward) and {C.OPERATOR_FULL}.",
         "Deal brief; party constants in common.py.", "ASSUMPTION - NOT FACT", "Award; Tribal resolution",
         "Wrong party voids authority and precedence."),
        ("A2", f"The {C.TRIBE_SHORT} owns the {C.TRIBAL_ASSETS} unless the Award + written federal approval say otherwise.",
         C.CITES["equipment"] + "; " + C.CITES["real_property"], "ASSUMPTION - NOT FACT", "Award; SAC; grants officer",
         "Alters IRU structure, disposition, federal interest."),
        ("A3", f"{C.OPERATOR_SHORT} is a contractor/vendor, not a subrecipient (provisional).",
         C.CITES["subrecipient"], "ASSUMPTION - NOT FACT", "Grants officer written determination",
         "If subrecipient, triggers pass-through duties (200.332)."),
        ("A4", f"{C.LREMC_SHORT} is a SEPARATE owner of poles/fiber/easements/huts/power/land ({C.LREMC_ASSETS}); separate consent/joinder required.",
         "Deal brief; " + C.CITES["nc_pole"], "ASSUMPTION - NOT FACT", "LREMC board; title/easement records",
         "No consent = no lawful access; construction blocked."),
        ("A5", f"Primary economics: {C.OPERATOR_SHORT} pays the {C.TRIBE_SHORT} {C.DEAL['per_subscriber_amount']}; characterization Option D (provisional).",
         C.DEAL["program_income_note"], "ASSUMPTION - NOT FACT", "Grants officer; finance advisor",
         "Program-income treatment turns on the Award, not the label."),
        ("A6", f"Term = {C.IRU_TERM_RECOMMENDED} years from {C.TERM_TRIGGER} + {C.DEAL['iru_renewal']}.",
         "Deal brief; ceiling: " + C.DEAL["term_ceiling_note"], "ASSUMPTION - NOT FACT", "Counsel; grants officer",
         "Term cannot exceed useful life / federal-interest period."),
        ("A7", f"Provisional baseline = {C.NOFO_NAME}.",
         C.NOFO_LIVE_NOTE, "ASSUMPTION - NOT FACT", "Grants officer; NOFO PDFs",
         "Wrong round changes eligibility, match, deadlines."),
        ("A8", f"Pre-existing {C.OPERATOR_EXISTING} stays {C.OPERATOR_SHORT}'s and must be valued / cost-allocated - never a free contribution.",
         C.CITES["allowable"], "ASSUMPTION - NOT FACT", "Valuation; finance advisor",
         "Free use of grant capacity or unpriced contribution risks disallowance."),
        ("A9", "Lumbee federal-recognition status is adequate for TBCP eligibility (THRESHOLD).",
         C.CITES["lumbee_act"], "ASSUMPTION - NOT FACT", "Counsel; NTIA (threshold)",
         "Eligibility failure is dispositive - flag GRANT + ATTORNEY."),
        ("A10", f"Approved service floor is {C.SPEED_FLOOR}.",
         "TBCP guidance; confirm against approved commitments.", "ASSUMPTION - NOT FACT", "Award; application",
         "Affects design, SLA, and affordability commitments."),
        ("A11", "Grant-funded capacity is never free commercial capacity; commercial/non-project use requires valuation + approval.",
         C.CITES["prog_income"] + "; " + C.CITES["equipment"], "ASSUMPTION - NOT FACT", "Grants officer",
         "Unapproved commercial use risks program income / disposition findings."),
        ("A12", "Sovereign-immunity approach, governing law, forum, and dispute resolution are UNRESOLVED (bracketed alternatives).",
         C.CITES["sovereign_immunity"], "ASSUMPTION - NOT FACT", "Tribal counsel; Council",
         "Silent choice is impermissible; express waiver (if any) must be Council-approved."),
    ]
    r = ahdr + 1
    for aid, txt, basis, tag, confirm, impact in assumptions:
        C.xl_cell(wa, r, 1, aid, "text", bold=True, align="center")
        C.xl_cell(wa, r, 2, txt, "text", wrap=True)
        C.xl_cell(wa, r, 3, basis, "text", wrap=True)
        C.xl_cell(wa, r, 4, tag, "output", bold=True, wrap=True)
        C.xl_cell(wa, r, 5, confirm, "input", wrap=True)
        C.xl_cell(wa, r, 6, impact, "text", wrap=True)
        r += 1
    C.xl_legend(wa, r + 1)

    return C.xl_save(wb, SUBDIR, "01_Source_and_Assumptions_Register.xlsx")


# ===========================================================================
# 2. 02_Missing_Information_Request.docx
# ===========================================================================
def build_02():
    doc = doc_start("DOCUMENT 02 — MISSING INFORMATION REQUEST",
                    "Missing Information Request",
                    "Phase 1 Diligence — Information and Document Requests by Topic",
                    "Missing-Info Request")
    C.article(doc, 1, "Purpose and Instructions")
    C.para(doc, "This request lists the information and documents needed to complete Phase 1 diligence and "
                f"to draft definitive agreements between the {C.TRIBE_SHORT} and {C.OPERATOR_SHORT} for the "
                f"{C.PROGRAM_FULL} ({C.PROGRAM_SHORT}) fiber project. " + DOC_INTRO_NOTE)
    C.para(doc, "Items are organized by topic. For each item: what is needed, why it matters, the owner best "
                "positioned to provide it, and a priority (P1 = blocks drafting/execution; P2 = needed before "
                "finalizing; P3 = needed before closeout/operations).")
    C.flag_para(doc, C.FLAG_GRANT, "Provide the executed Award, Specific Award Conditions, controlling NOFO, and "
                "approved application/budget/routes first; these control the entire package.")

    headers = ["#", "Information Needed", "Why It Matters", "Owner", "Priority"]
    widths = [0.35, 1.9, 2.35, 1.0, 0.9]

    topics = [
        ("Legal", [
            ("Executed Federal Award / CD-450 and all incorporated terms", "Controlling instrument; sets precedence, federal interest, and permitted structure.", f"{C.TRIBE_SHORT} / NTIA", "P1"),
            ("Tribal Council authorizing resolution(s) and delegated signatory authority", "Confirms who may bind the Tribe and on what terms.", f"{C.TRIBE_SHORT}", "P1"),
            (f"{C.OPERATOR_SHORT} formation, governance, and signatory authority; affiliation with {C.LREMC_SHORT}", "Confirms capacity to contract and reveals conflict/OCI issues.", f"{C.OPERATOR_SHORT}", "P1"),
            ("Sovereign-immunity, governing-law, forum, and dispute-resolution positions", "Cannot be chosen silently; drives enforceability. " + C.FLAG_ATTORNEY, f"{C.TRIBE_SHORT} counsel", "P1"),
            ("Lumbee federal-recognition status documentation", "Threshold TBCP eligibility question. " + C.CITES["lumbee_act"], f"{C.TRIBE_SHORT} counsel", "P1"),
        ]),
        ("Grant", [
            ("Specific Award Conditions and NTIA/DOC Standard Terms", "May restrict IRU, operator role, disposition, and program income. " + C.CITES["sac"], f"{C.TRIBE_SHORT}", "P1"),
            ("Controlling NOFO and all amendments (confirm Round 2 vs Round 3)", C.NOFO_LIVE_NOTE, f"{C.TRIBE_SHORT} / NTIA", "P1"),
            ("Approved application, budget, budget narrative, and cost categories", "Basis for reimbursement, allowability, and cost allocation. " + C.CITES["allowable"], f"{C.TRIBE_SHORT}", "P1"),
            ("Written determinations already issued by the grants officer, if any", "Avoids re-litigating settled points; sets baseline for Doc 09 requests.", f"{C.TRIBE_SHORT} / NTIA", "P2"),
            ("BABA and Section 889 compliance records/waivers", "Non-compliance risks disallowed costs. " + C.CITES["baba"] + "; " + C.CITES["telecom_ban"], f"{C.OPERATOR_SHORT}", "P2"),
        ]),
        ("Engineering", [
            ("Approved network design, routes, GIS/KMZ, and useful-life data", "Defines scope, acceptance, and IRU capacity. Routes/useful life NOT supplied.", f"{C.OPERATOR_SHORT}", "P1"),
            ("Bill of materials and equipment list (with §889/BABA status)", "Supply-chain compliance and cost allocation.", f"{C.OPERATOR_SHORT}", "P2"),
            ("Interconnection points, transport, and demarcation design", "Determines interconnection approvals and asset boundaries.", f"{C.OPERATOR_SHORT}", "P2"),
            ("Acceptance-testing standards and segment-acceptance procedure", f"Trigger for the term start ({C.TERM_TRIGGER}).", f"{C.OPERATOR_SHORT}", "P2"),
        ]),
        ("Finance", [
            (f"Valuation and cost allocation of the {C.OPERATOR_EXISTING}", "Existing assets are not a free contribution; must be priced. " + C.CITES["allowable"], f"{C.OPERATOR_SHORT}", "P1"),
            ("Proposed per-Active-Subscriber amount and supporting model", f"Sets primary economics; {C.DEAL['per_subscriber_amount']}. " + C.CITES["prog_income"], f"{C.OPERATOR_SHORT} / {C.TRIBE_SHORT}", "P1"),
            ("Program-income history and treatment for prior awards", "Informs program-income determination. " + C.CITES["prog_income"], f"{C.TRIBE_SHORT}", "P2"),
            ("Contractor pricing / margin basis for reasonableness review", "Required cost/price analysis. " + C.CITES["procurement"], f"{C.OPERATOR_SHORT}", "P2"),
            ("Most recent single-audit report(s)", "Risk assessment and audit posture. " + C.CITES["single_audit"], f"{C.TRIBE_SHORT}", "P3"),
        ]),
        ("Customer Operations", [
            ("Provider-of-record structure and existing customer contracts", "Confirms retail role and non-project customers.", f"{C.OPERATOR_SHORT}", "P1"),
            ("Billing/OSS/BSS systems, data ownership, and portability", "Affects step-in, transition, and program-income tracking.", f"{C.OPERATOR_SHORT}", "P2"),
            ("Affordability program design (e.g., low-cost service option)", "Confirm against approved service commitments.", f"{C.OPERATOR_SHORT} / {C.TRIBE_SHORT}", "P2"),
            ("Customer-receipt and revenue-recording procedures", "Program-income tracking and record retention. " + C.CITES["records"], f"{C.OPERATOR_SHORT}", "P2"),
        ]),
        ("Land", [
            (f"{C.LREMC_SHORT} pole, conduit, easement, ROW, land, hut, and power records", f"{C.LREMC_SHORT} is a separate owner; access needs consent/joinder. " + C.CITES["nc_pole"], f"{C.LREMC_SHORT}", "P1"),
            ("Title and real-property records for all sites", "Confirms rights and federal-interest placement. " + C.CITES["real_property"], f"{C.LREMC_SHORT} / {C.TRIBE_SHORT}", "P1"),
            ("Whether any trust/restricted Indian land is used or crossed", "Triggers BIA leasing/ROW approval. " + C.CITES["bia_approval"], f"{C.TRIBE_SHORT} counsel", "P1"),
            ("NCDOT encroachment and 811/locate compliance records", C.CITES["nc_ncdot"] + "; " + C.CITES["nc_dig"], f"{C.OPERATOR_SHORT}", "P2"),
        ]),
        ("Cybersecurity", [
            ("System Security Plan, supply-chain risk plan, and incident-response plan", "Required for operations and data protection.", f"{C.OPERATOR_SHORT}", "P2"),
            ("CPNI and customer-data handling policies", C.CITES["cpni"], f"{C.OPERATOR_SHORT}", "P2"),
            ("Disaster-recovery and business-continuity plans", "Supports step-in and continuity obligations.", f"{C.OPERATOR_SHORT}", "P3"),
        ]),
        ("Insurance", [
            ("Certificates of insurance and full policies (both parties)", "Confirm coverage adequacy and additional-insured status. " + C.CITES["insurance"], "Both Parties", "P2"),
            ("Builder's risk / property coverage for grant-funded assets", f"Protects the {C.TRIBAL_ASSETS} during construction.", f"{C.OPERATOR_SHORT}", "P2"),
        ]),
        ("Tax / TERO", [
            ("Applicable Tribal tax, TERO, and preference requirements", C.CITES["tero"], f"{C.TRIBE_SHORT}", "P2"),
            ("State/local tax treatment of assets, services, and payments", "Affects economics and payment characterization.", "Tax advisor", "P3"),
        ]),
    ]

    n = 2
    for topic, items in topics:
        C.article(doc, n, topic)
        rows = [[str(i + 1), it[0], it[1], it[2], it[3]] for i, it in enumerate(items)]
        C.add_table(doc, headers, rows, widths=widths, col_align=["c", "l", "l", "l", "c"])
        C.spacer(doc, 1)
        n += 1

    doc_close(doc, n)
    return C.save(doc, SUBDIR, "02_Missing_Information_Request.docx")


# ===========================================================================
# 3. 03_Red_Flag_and_Approval_Memorandum.docx
# ===========================================================================
def build_03():
    doc = doc_start("DOCUMENT 03 — RED-FLAG AND APPROVAL MEMORANDUM",
                    "Red-Flag and Approval Memorandum",
                    "Phase 1 Diligence — Issues That Could Prevent or Delay the Transaction",
                    "Red-Flag Memo")
    C.article(doc, 1, "Purpose")
    C.para(doc, "This memorandum identifies issues that could prevent or delay execution, reimbursement, "
                "construction, interconnection, service launch, receipt retention, the per-Active-Subscriber "
                f"payment, or enforceability of the {C.TRIBE_SHORT} / {C.OPERATOR_SHORT} transaction. " + DOC_INTRO_NOTE)
    C.para(doc, "Each red flag states the risk, the consequence, a recommended mitigation, and the governing "
                "citation. Section 3 lists the specific written determinations that should be requested from "
                "NTIA / the grants officer (see Document 09).")

    C.article(doc, 2, "Red-Flag Register")
    rf_headers = ["Risk", "Consequence", "Recommended Mitigation", "Citation"]
    rf_widths = [1.5, 1.5, 2.0, 1.5]
    rf = [
        ("No controlling award or source documents supplied", "Cannot finalize any definitive clause; drafts are provisional.",
         "Obtain executed Award, SAC, controlling NOFO, approved application/budget/routes before finalizing. " + C.FLAG_GRANT, C.CITES["sac"]),
        ("Contractor vs. subrecipient status of RIVR Tech undetermined", "Wrong classification imposes or omits pass-through duties; risks findings.",
         "Complete the 200.331 substance analysis (Doc 04) and obtain a written determination. " + C.FLAG_GRANT, C.CITES["subrecipient"]),
        ("RIVR Tech as retail provider of record not approved", "Operator/retail role over grant-funded assets may be impermissible without approval.",
         "Request written operating-appointment determination; structure via IRU + O&M. " + C.FLAG_GRANT, C.CITES["equipment"]),
        ("Per-subscriber payment characterization / program income", "Mischaracterization risks disallowance or misreported program income.",
         "Obtain written program-income determination; label does not control. " + C.FLAG_GRANT, C.CITES["prog_income"]),
        ("Use of existing RIVR Tech assets as a free contribution", "Unpriced use of pre-existing or grant capacity risks disallowed costs.",
         "Value and cost-allocate existing assets; never treat as free. " + C.FLAG_BUSINESS, C.CITES["allowable"]),
        (f"{C.LREMC_SHORT} assets used without consent or joinder", "No lawful access to poles/fiber/easements/huts/power/land; construction blocked.",
         f"Obtain {C.LREMC_SHORT} consent/joinder and separate instruments (Schedule S9). " + C.FLAG_ATTORNEY, C.CITES["nc_pole"]),
        ("Sovereign-immunity / governing-law / forum unresolved", "Enforceability and dispute resolution uncertain; silent choice impermissible.",
         "Present bracketed alternatives for Council decision; any waiver must be express + approved. " + C.FLAG_ATTORNEY, C.CITES["sovereign_immunity"]),
        ("Lumbee federal-recognition threshold eligibility", "Eligibility failure is dispositive for TBCP participation.",
         "Confirm recognition status and TBCP eligibility before proceeding. " + C.FLAG_ATTORNEY + " " + C.FLAG_GRANT, C.CITES["lumbee_act"]),
        ("Procurement compliance / OCI (RIVR Tech affiliated with LREMC)", "Non-compliant selection risks disallowed costs and conflict findings.",
         "Document competition or sole-source justification; address OCI + conduct standards (Doc 05). " + C.FLAG_GRANT, C.CITES["procurement"]),
        ("BABA / Section 889 non-compliance", "Prohibited equipment or non-domestic content risks disallowance.",
         "Collect attestations, BOM review, and any waivers. " + C.FLAG_GRANT, C.CITES["baba"]),
        ("NEPA / NHPA not complete before construction", "Ground-disturbing work before clearance risks reimbursement loss.",
         "Complete environmental and Section 106 review before construction. " + C.FLAG_GRANT, C.CITES["nepa"] + "; " + C.CITES["nhpa"]),
        ("Federal interest and disposition on grant-funded assets", "Encumbrance/IRU may require federal approval; disposition rules apply.",
         "Map federal interest (Doc 07); obtain approvals for any encumbrance. " + C.FLAG_GRANT, C.CITES["real_property"] + "; " + C.CITES["equipment"]),
        ("Land / ROW / pole rights not secured", "Construction and operations cannot lawfully proceed.",
         f"Secure {C.LREMC_SHORT}/third-party instruments; confirm any BIA approvals. " + C.FLAG_ATTORNEY, C.CITES["indian_row"]),
        ("IRU to a contractor over federally funded assets", "Grant of long-term rights may need federal approval and program-income analysis.",
         "Request IRU determination; align term with federal-interest period. " + C.FLAG_GRANT, C.CITES["intangible"]),
        ("Commercial / non-project capacity and exclusivity", "Unapproved commercial use of grant capacity risks findings.",
         "Define reserved/commercial capacity; value it; seek approval (Schedule S10). " + C.FLAG_BUSINESS, C.CITES["prog_income"]),
        ("Customer-receipt retention and record access", "Inadequate records risk audit findings and payment disputes.",
         "Implement receipt retention and access procedures. " + C.FLAG_GRANT, C.CITES["records"]),
        ("Insurance adequacy and additional-insured status", "Underinsurance exposes the Tribe and grant-funded assets.",
         "Confirm limits and endorsements (Schedule S12). " + C.FLAG_ATTORNEY, C.CITES["insurance"]),
        ("Term ceiling vs. useful life and federal-interest period", "A 20-year + renewals term may exceed permissible ceilings.",
         "Cap term at " + C.DEAL["term_ceiling_note"] + ". " + C.FLAG_ATTORNEY, C.CITES["equipment"]),
    ]
    C.add_table(doc, rf_headers, [list(x) for x in rf], widths=rf_widths)
    C.spacer(doc, 1)

    C.article(doc, 3, "Required NTIA / Grants-Officer Written Determinations")
    C.para(doc, "The following written determinations should be requested from NTIA / the grants officer before "
                "the corresponding provisions are finalized. See Document 09 for the structured request.")
    det_headers = ["#", "Determination Needed", "Why It Matters", "Provisional Structure / Fallback"]
    det_widths = [0.35, 1.7, 2.0, 2.45]
    dets = [
        ("IRU grant to RIVR Tech over Grant-Funded Assets", "Long-term rights over federally funded assets may require approval.",
         "Structure as IRU limited to defined capacity within federal-interest period; fallback: capacity lease/license. " + C.CITES["intangible"]),
        ("Operating appointment / operator role", "Operator role over grant-funded network must be permissible.",
         "Appoint via O&M agreement; fallback: management agreement with Tribe retaining control. " + C.CITES["equipment"]),
        ("Interconnection arrangement", "Interconnection to existing networks affects asset boundaries and cost.",
         "Defined POIs and transport terms; fallback: arm's-length transport agreement. " + C.CITES["equipment"]),
        ("Use of existing (non-grant) RIVR Tech assets in the project", "Prevents free contribution / unallowable cost.",
         "Value and cost-allocate; fallback: exclude from project or lease at documented rate. " + C.CITES["allowable"]),
        ("Customer receipts / revenue handling", "Determines who collects and how receipts are retained.",
         "RIVR Tech collects as provider of record with segregated records; fallback: Tribe-controlled account. " + C.CITES["records"]),
        ("Per-subscriber payment characterization (program income)", "Program-income treatment turns on the Award, not the label.",
         "Option D per-Active-Subscriber operating payment; fallback: Options A/B/C/E. " + C.CITES["prog_income"]),
        ("Commercial / non-project capacity use", "Commercial use of grant capacity may generate program income or be prohibited.",
         "Define reserved commercial capacity, valued and approved; fallback: no commercial use. " + C.CITES["prog_income"]),
        ("Contractor pricing / margin reasonableness", "Costs must be reasonable and allowable.",
         "Cost/price analysis on record; fallback: renegotiate or competitively re-solicit. " + C.CITES["procurement"]),
        ("Out-of-scope service", "Services beyond approved scope may be unallowable or need approval.",
         "Confine to approved scope; fallback: separate non-grant commercial arrangement. " + C.CITES["allowable"]),
    ]
    C.add_table(doc, det_headers, [[str(i + 1), d[0], d[1], d[2]] for i, d in enumerate(dets)],
                widths=det_widths, col_align=["c", "l", "l", "l"])
    C.flag_para(doc, C.FLAG_GRANT, "No determination above has been obtained. Each is PENDING and must be secured "
                "in writing before the related clause is finalized.")

    doc_close(doc, 4)
    return C.save(doc, SUBDIR, "03_Red_Flag_and_Approval_Memorandum.docx")


# ===========================================================================
# 4. 04_Contractor_vs_Subrecipient_Analysis.docx
# ===========================================================================
def build_04():
    doc = doc_start("DOCUMENT 04 — CONTRACTOR VS. SUBRECIPIENT ANALYSIS",
                    "Contractor vs. Subrecipient Analysis",
                    "Phase 1 Diligence — 2 CFR 200.331 Substance Determination for RIVR Tech",
                    "Contractor/Subrecipient")
    C.article(doc, 1, "Purpose and Standard")
    C.para(doc, f"This document analyzes whether {C.OPERATOR_FULL} ({C.OPERATOR_SHORT}) should be treated as a "
                f"subrecipient or a contractor under {C.CITES['subrecipient']}. The classification turns on the "
                "substance of the relationship, not the form or title of the instrument. " + DOC_INTRO_NOTE)
    C.para(doc, "A subrecipient carries out a portion of the federal program for a public purpose and is subject "
                "to program compliance requirements; a contractor provides goods or services, within its normal "
                "business operations, to many purchasers, in a competitive environment, ancillary to the program.")

    C.article(doc, 2, "Substance Analysis (2 CFR 200.331)")
    hdr = ["Subrecipient Characteristic (200.331(a))", "Contractor Characteristic (200.331(b))", "Application to RIVR Tech", "Points Toward"]
    widths = [1.6, 1.6, 2.3, 1.0]
    rows = [
        ("Determines eligibility for federal assistance",
         "Provides goods/services within normal business operations",
         "RIVR Tech does not determine eligibility; the Tribe (recipient) does. RIVR Tech deploys and operates as its business. " + C.PH("confirm against approved application"),
         "CONTRACTOR"),
        ("Has performance measured against program objectives",
         "Provides similar goods/services to many purchasers",
         "RIVR Tech provides broadband deployment/operations of a kind sold to multiple customers; measured against a scope of work. " + C.PH("confirm existing customer base"),
         "CONTRACTOR"),
        ("Has responsibility for programmatic decision-making",
         "Normally operates in a competitive environment",
         "Programmatic decisions (scope, service commitments) remain with the Tribe; RIVR Tech executes. Competitive selection must be documented. " + C.FLAG_GRANT,
         "CONTRACTOR"),
        ("Responsible for adherence to federal program requirements",
         "Provides goods/services ancillary to program operation",
         "RIVR Tech must flow-down applicable clauses but is not the program steward; its services are ancillary to the Tribe's program.",
         "CONTRACTOR"),
        ("Uses federal funds to carry out a program for a public purpose (vs. providing goods/services for the recipient's own use)",
         "Is not subject to the program's compliance requirements (only to contract terms)",
         "If RIVR Tech were to assume programmatic responsibility (e.g., discretion over eligibility or service policy), a subrecipient characteristic could arise. " + C.FLAG_GRANT,
         "CONFIRM"),
    ]
    C.add_table(doc, hdr, [list(x) for x in rows], widths=widths, col_align=["l", "l", "l", "c"])
    C.spacer(doc, 1)

    C.article(doc, 3, "Recommendation")
    C.para(doc, "On the facts assumed, the weight of the substance factors points toward treating RIVR Tech as a "
                "CONTRACTOR / vendor (procured provider of design-engineering-procurement-construction and "
                "operations/maintenance services, and proposed retail provider of record), not a subrecipient.", bold=True)
    C.flag_para(doc, C.FLAG_GRANT, "This recommendation is provisional. The classification must be confirmed by a "
                "written contractor-vs-subrecipient determination by the recipient (with grants-officer concurrence "
                "as required) under " + C.CITES["subrecipient"] + ", supported by the approved application and scope.")

    C.article(doc, 4, "If Both Roles Are Present (Dual-Role Handling)")
    C.para(doc, "If any scope assigns RIVR Tech programmatic responsibility (a subrecipient characteristic), the "
                "scopes must be divided and documented separately, with controls appropriate to each role.")
    C.section(doc, "4.1", "Contractor scope (default)",
              "Design, engineering, procurement, construction, operations, maintenance, and retail service "
              "delivery as procured services. Controls: compliant procurement, cost/price reasonableness, "
              "flow-down clauses, and record access. " + C.CITES["procurement"])
    C.section(doc, "4.2", "Subrecipient scope (only if programmatic responsibility is assigned)",
              "Any delegated programmatic decision-making or public-purpose administration. Controls: a subaward "
              "with required data elements, risk assessment, monitoring, and closeout as a pass-through entity. "
              + C.CITES["pass_through"])
    C.flag_para(doc, C.FLAG_ATTORNEY, "Where any subrecipient scope exists, the Tribe becomes a pass-through entity "
                "under " + C.CITES["pass_through"] + " and must implement subrecipient monitoring; counsel and the "
                "grants officer must confirm the split before execution.")

    doc_close(doc, 5)
    return C.save(doc, SUBDIR, "04_Contractor_vs_Subrecipient_Analysis.docx")


# ===========================================================================
# 5. 05_Procurement_and_OCI_Checklist.xlsx
# ===========================================================================
def build_05():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Procurement and OCI Checklist — Phase 1 Diligence",
        "Document 05 of the Phase 1 Diligence set.",
        DOC_INTRO_NOTE,
        "",
        "## How to use",
        "• Sheet 'Procurement' documents how RIVR Tech's selection must be supported under " + C.CITES["procurement"] + ".",
        "• Sheet 'OCI & Conduct' covers organizational conflict of interest and standards of conduct under " + C.CITES["oci"] + ".",
        "• Because RIVR Tech (LREMC Technologies, LLC) is affiliated with LREMC, the OCI/standards-of-conduct analysis is critical.",
        "• Status: OPEN = not yet evidenced; IN PROGRESS; SATISFIED. Fill Evidence/Doc and Owner as records arrive.",
        "• Nearly all rows are OPEN because no procurement record was supplied - do not invent one.",
    ], title="Instructions")

    # Procurement sheet
    ws = wb.create_sheet("Procurement")
    C.xl_title(ws, "Procurement Compliance Checklist (2 CFR 200.317-200.327)",
               "Documents how RIVR Tech's selection must be supported; nearly all rows OPEN (no record supplied).", span=5)
    widths = [50, 40, 16, 34, 22]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    hdr = 5
    C.xl_header_row(ws, hdr, ["Requirement", "Citation", "Status", "Evidence/Doc", "Owner"])
    proc = [
        ("SECTION: General procurement standards", "", "", "", ""),
        ("Documented procurement procedures consistent with federal standards", C.CITES["procurement"], "OPEN", C.PH("procedures"), C.TRIBE_SHORT),
        ("Written standards of conduct / conflict-of-interest policy in force", C.CITES["conflict"], "OPEN", C.PH("policy"), C.TRIBE_SHORT),
        ("Necessity / avoidance of unnecessary or duplicative purchases", "2 CFR 200.318(d)", "OPEN", C.PH("analysis"), C.TRIBE_SHORT),
        ("SECTION: Competition", "", "", "", ""),
        ("Full and open competition; no unfair advantage to affiliates", C.CITES["competition"], "OPEN", C.PH("solicitation"), C.TRIBE_SHORT),
        ("No restrictive specifications or geographic preferences (except as allowed)", C.CITES["competition"], "OPEN", C.PH("spec review"), C.TRIBE_SHORT),
        ("Solicitation clearly describes requirements and evaluation basis", "2 CFR 200.319(c)", "OPEN", C.PH("solicitation"), C.TRIBE_SHORT),
        ("SECTION: Methods of procurement", "", "", "", ""),
        ("Method selected and documented (small purchase / sealed bid / proposals)", C.CITES["methods"], "OPEN", C.PH("method memo"), C.TRIBE_SHORT),
        ("If noncompetitive/sole-source: written justification meets a 200.320(c) exception", "2 CFR 200.320(c)", "OPEN", C.PH("sole-source justification"), C.TRIBE_SHORT),
        ("Sole-source exception analysis (unique source / urgency / competition inadequate / awarding-agency authorization)", "2 CFR 200.320(c)(1)-(4)", "OPEN", C.PH("exception analysis"), C.TRIBE_SHORT),
        ("SECTION: Cost / price and contractor responsibility", "", "", "", ""),
        ("Cost or price analysis performed and documented", "2 CFR 200.324", "OPEN", C.PH("cost/price analysis"), C.TRIBE_SHORT),
        ("Contractor responsibility determination (integrity, capacity, record)", "2 CFR 200.318(h)", "OPEN", C.PH("responsibility memo"), C.TRIBE_SHORT),
        ("Debarment / suspension (SAM exclusions) verification", C.CITES["debarment"], "OPEN", C.PH("SAM check"), C.TRIBE_SHORT),
        ("SECTION: Required contract provisions and preferences", "", "", "", ""),
        ("Required contract clauses included (Appendix II to Part 200)", "2 CFR Part 200, Appendix II", "OPEN", C.PH("clause matrix"), C.TRIBE_SHORT),
        ("Domestic preference / BABA compliance addressed", C.CITES["domestic"] + "; " + C.CITES["baba"], "OPEN", C.PH("BABA record"), C.OPERATOR_SHORT),
        ("Section 889 covered-equipment prohibition addressed", C.CITES["telecom_ban"], "OPEN", C.PH("889 attestation"), C.OPERATOR_SHORT),
        ("Procurement records retained and available for access", C.CITES["records"], "OPEN", C.PH("record file"), C.TRIBE_SHORT),
        ("SECTION: Approvals", "", "", "", ""),
        ("Awarding-agency pre-procurement review, if required", "2 CFR 200.325", "OPEN", C.PH("agency review"), C.TRIBE_SHORT),
    ]
    r = hdr + 1
    for req, cite, status, ev, owner in proc:
        if req.startswith("SECTION:"):
            C.xl_cell(ws, r, 1, req.replace("SECTION: ", ""), "section")
            for col in range(2, 6):
                C.xl_cell(ws, r, col, "", "section")
        else:
            C.xl_cell(ws, r, 1, req, "text", wrap=True)
            C.xl_cell(ws, r, 2, cite, "text", wrap=True)
            C.xl_cell(ws, r, 3, status, "output", bold=True, align="center")
            C.xl_cell(ws, r, 4, ev, "input", wrap=True)
            C.xl_cell(ws, r, 5, owner, "text", wrap=True)
        r += 1
    C.xl_legend(ws, r + 1)

    # OCI sheet
    wo = wb.create_sheet("OCI & Conduct")
    C.xl_title(wo, "Organizational Conflict of Interest & Standards of Conduct",
               "Critical because RIVR Tech (LREMC Technologies, LLC) is affiliated with LREMC, a separate asset owner.", span=5)
    for i, w in enumerate(widths):
        wo.column_dimensions[chr(65 + i)].width = w
    ohdr = 5
    C.xl_header_row(wo, ohdr, ["Requirement", "Citation", "Status", "Evidence/Doc", "Owner"])
    oci = [
        ("SECTION: Organizational conflict of interest", "", "", "", ""),
        ("Identify affiliation between RIVR Tech, LREMC Technologies LLC, and LREMC", C.CITES["oci"], "OPEN", C.PH("org chart / ownership"), C.OPERATOR_SHORT),
        ("Assess OCI arising from affiliate ownership of poles/fiber/easements/land", C.CITES["oci"], "OPEN", C.PH("OCI assessment"), C.TRIBE_SHORT),
        ("Mitigation plan for identified OCI (firewalls, disclosures, recusal)", C.CITES["oci"], "OPEN", C.PH("mitigation plan"), C.TRIBE_SHORT),
        ("SECTION: Standards of conduct / individual conflicts", "", "", "", ""),
        ("Written standards of conduct covering employees/officers/agents", C.CITES["conflict"], "OPEN", C.PH("policy"), C.TRIBE_SHORT),
        ("No employee with a real/apparent conflict participates in selection", C.CITES["conflict"], "OPEN", C.PH("recusal record"), C.TRIBE_SHORT),
        ("Gifts/gratuities prohibition and disciplinary provisions", "2 CFR 200.318(c)(1)", "OPEN", C.PH("policy"), C.TRIBE_SHORT),
        ("Mandatory disclosure of violations of law/fraud", C.CITES["disclosures"], "OPEN", C.PH("disclosure procedure"), C.TRIBE_SHORT),
        ("SECTION: Related-party economics", "", "", "", ""),
        ("Arm's-length pricing for any LREMC-owned facilities used", C.CITES["allowable"], "OPEN", C.PH("pricing basis"), C.OPERATOR_SHORT),
        ("Consideration/valuation for existing RIVR Tech assets (never free)", C.CITES["allowable"], "OPEN", C.PH("valuation"), C.OPERATOR_SHORT),
    ]
    r = ohdr + 1
    for req, cite, status, ev, owner in oci:
        if req.startswith("SECTION:"):
            C.xl_cell(wo, r, 1, req.replace("SECTION: ", ""), "section")
            for col in range(2, 6):
                C.xl_cell(wo, r, col, "", "section")
        else:
            C.xl_cell(wo, r, 1, req, "text", wrap=True)
            C.xl_cell(wo, r, 2, cite, "text", wrap=True)
            C.xl_cell(wo, r, 3, status, "output", bold=True, align="center")
            C.xl_cell(wo, r, 4, ev, "input", wrap=True)
            C.xl_cell(wo, r, 5, owner, "text", wrap=True)
        r += 1
    C.xl_legend(wo, r + 1)

    return C.xl_save(wb, SUBDIR, "05_Procurement_and_OCI_Checklist.xlsx")


# ===========================================================================
# 6. 06_Responsibility_and_Transaction_Matrix.xlsx
# ===========================================================================
def build_06():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Responsibility and Transaction Matrix — Phase 1 Diligence",
        "Document 06 of the Phase 1 Diligence set.",
        DOC_INTRO_NOTE,
        "",
        "## How to use",
        "• Sheet 'Transaction Diagram' describes the parties and the flows among them in text cells.",
        "• Sheet 'RACI Matrix' assigns Responsible / Accountable / Consulted / Informed for each key activity.",
        "• RACI is PROVISIONAL and must be confirmed against the controlling award and executed agreements.",
    ], title="Instructions")

    # Transaction diagram sheet (text-based)
    ws = wb.create_sheet("Transaction Diagram")
    C.xl_title(ws, "Transaction Diagram (Text) — Parties and Flows",
               "A text description of the parties and the principal flows among them.", span=4)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 60
    r = 5
    C.xl_cell(ws, r, 2, "PARTIES / NODES", "section", bold=True)
    C.xl_cell(ws, r, 3, "ROLE", "section", bold=True)
    C.xl_cell(ws, r, 4, "NOTES", "section", bold=True)
    r += 1
    nodes = [
        (C.TRIBE_FULL, "TBCP applicant; award recipient & grant steward; owner of Grant-Funded Assets", "Owns the " + C.NETWORK + " unless Award + written federal approval say otherwise."),
        ("NTIA / U.S. Dept. of Commerce", "Federal awarding agency; issues determinations", C.CITES["sac"]),
        (C.OPERATOR_FULL, "DEPC + operate/maintain + proposed retail provider of record (contractor)", "Subject to compliant procurement + written contractor-vs-subrecipient determination."),
        (C.LREMC_FULL, "Separate owner of poles/conduit/fiber/easements/huts/power/land", "Requires separate consent/joinder; not owned by RIVR Tech."),
        ("Subcontractors / vendors", "Construction, equipment, and specialty services", "Flow-down clauses; BABA/§889 compliance."),
        ("Landowners / pole owners / NCDOT", "Grant land, ROW, pole, and encroachment rights", C.CITES["nc_pole"] + "; " + C.CITES["nc_ncdot"]),
        ("Retail customers", "Subscribers within the Approved Project Area", "Basis for Active-Subscriber count and payment."),
        ("Billing / OSS-BSS / software vendors", "Billing, provisioning, and support systems", "Data ownership and portability to be confirmed."),
    ]
    for name, role, note in nodes:
        C.xl_cell(ws, r, 2, name, "text", wrap=True, bold=True)
        C.xl_cell(ws, r, 3, role, "text", wrap=True)
        C.xl_cell(ws, r, 4, note, "text", wrap=True)
        r += 1
    r += 1
    C.xl_cell(ws, r, 2, "FLOW", "section", bold=True)
    C.xl_cell(ws, r, 3, "FROM -> TO", "section", bold=True)
    C.xl_cell(ws, r, 4, "DESCRIPTION", "section", bold=True)
    r += 1
    flows = [
        ("Federal award funds", "NTIA/DOC -> Lumbee Tribe", "Reimbursement of allowable costs under the Award."),
        ("Written determinations", "NTIA/DOC -> Lumbee Tribe", "IRU, operator, interconnection, program income, etc. (Doc 09)."),
        ("DEPC + O&M services", "RIVR Tech -> Lumbee Tribe", "Procured services to build and operate the Tribal Network."),
        ("IRU / network-access rights", "Lumbee Tribe -> RIVR Tech", "Right to use defined capacity; subject to federal approval."),
        ("Per-Active-Subscriber payment", "RIVR Tech -> Lumbee Tribe", C.DEAL["per_subscriber_amount"] + "; characterization Option D (provisional)."),
        ("Pole/fiber/easement/hut/power access", "LREMC -> project (via consent/joinder)", "Separate instruments; arm's-length consideration."),
        ("Existing network capacity", "RIVR Tech Existing Network -> project", "Valued/cost-allocated; never a free contribution."),
        ("Retail service", "RIVR Tech -> customers", "Provider of record; billing and collections."),
        ("Subscriber revenue", "customers -> RIVR Tech", "Basis for Active-Subscriber count; receipts retained."),
        ("Land / ROW / encroachment rights", "landowners/NCDOT -> project", "Permits and easements for construction."),
        ("Subcontracted work / equipment", "subcontractors/vendors -> RIVR Tech", "Flow-down of federal clauses."),
        ("Billing / OSS-BSS services", "software vendors -> RIVR Tech", "Systems supporting retail operations."),
    ]
    for name, fromto, desc in flows:
        C.xl_cell(ws, r, 2, name, "text", wrap=True, bold=True)
        C.xl_cell(ws, r, 3, fromto, "text", wrap=True)
        C.xl_cell(ws, r, 4, desc, "text", wrap=True)
        r += 1

    # RACI sheet
    wr = wb.create_sheet("RACI Matrix")
    C.xl_title(wr, "RACI Responsibility Matrix (Provisional)",
               "R = Responsible, A = Accountable, C = Consulted, I = Informed.", span=5)
    wr.column_dimensions["A"].width = 44
    for col in ("B", "C", "D", "E"):
        wr.column_dimensions[col].width = 16
    rhdr = 5
    C.xl_header_row(wr, rhdr, ["Key Activity", C.TRIBE_SHORT, C.OPERATOR_SHORT, C.LREMC_SHORT, "NTIA/DOC"])
    raci = [
        ("Grant stewardship", "A", "C", "I", "C"),
        ("Procurement (of RIVR Tech and subs)", "A/R", "I", "I", "C"),
        ("Network design", "A", "R", "C", "I"),
        ("Environmental / NHPA review", "A", "R", "C", "C"),
        ("Land / ROW / pole rights", "A", "R", "R", "I"),
        ("Construction", "A", "R", "C", "I"),
        ("Testing / acceptance", "A", "R", "I", "I"),
        ("Interconnection", "A", "R", "C", "C"),
        ("Operations / NOC", "A", "R", "I", "I"),
        ("Billing / collections", "A", "R", "I", "I"),
        ("Customer service", "A", "R", "I", "I"),
        ("Program income", "A", "R", "I", "C"),
        ("Reporting", "A/R", "C", "I", "I"),
        ("Audit", "A/R", "C", "I", "C"),
        ("Step-in", "A", "C", "C", "I"),
        ("Disposition of assets", "A", "C", "C", "A/C"),
    ]
    r = rhdr + 1
    for activity, t, o, l, n in raci:
        C.xl_cell(wr, r, 1, activity, "text", wrap=True)
        C.xl_cell(wr, r, 2, t, "output", bold=True, align="center")
        C.xl_cell(wr, r, 3, o, "output", bold=True, align="center")
        C.xl_cell(wr, r, 4, l, "output", bold=True, align="center")
        C.xl_cell(wr, r, 5, n, "output", bold=True, align="center")
        r += 1
    r += 1
    C.xl_cell(wr, r, 1, "Legend:", "text", bold=True)
    for i, (code, mean) in enumerate([("R", "Responsible (does the work)"), ("A", "Accountable (owns the outcome)"),
                                      ("C", "Consulted (two-way input)"), ("I", "Informed (kept apprised)")]):
        C.xl_cell(wr, r + 1 + i, 1, f"{code} = {mean}", "text")
    C.xl_cell(wr, r + 6, 1, "PROVISIONAL - confirm against the executed agreements and award.", "text", bold=True)

    return C.xl_save(wb, SUBDIR, "06_Responsibility_and_Transaction_Matrix.xlsx")


# ===========================================================================
# 7. 07_Asset_Ownership_and_Rights_Matrix.xlsx
# ===========================================================================
def build_07():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Asset Ownership and Rights Matrix — Phase 1 Diligence",
        "Document 07 of the Phase 1 Diligence set.",
        DOC_INTRO_NOTE,
        "",
        "## How to use",
        "• One row per asset category, grouped into five sections.",
        "• 'Federal Interest Y/N' flags grant-funded assets subject to 2 CFR 200.311/.313 use and disposition rules.",
        "• 'Consideration/Valuation' must never be blank for RIVR Tech pre-existing or LREMC assets - never a free contribution.",
        "• Ownership/valuation values are placeholders because no title, inventory, or valuation was supplied.",
    ], title="Instructions")

    ws = wb.create_sheet("Asset & Rights Matrix")
    C.xl_title(ws, "Asset Ownership and Rights Matrix",
               "Grant-funded assets carry a federal interest; RIVR Tech and LREMC assets require rights + consideration.", span=8)
    widths = [40, 20, 24, 12, 34, 26, 26, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    hdr = 5
    C.xl_header_row(ws, hdr, ["Asset", "Owner", "Funding Source", "Federal Interest Y/N",
                              "Rights Needed by RIVR Tech", "Instrument", "Consideration/Valuation", "Notes"])

    def sect(name):
        return ("SECTION", name)

    rows = [
        sect("A. Grant-Funded Outside Plant"),
        ("Aerial/underground fiber cable, strands", C.TRIBE_SHORT, "TBCP grant", "Y", "IRU / capacity use", C.AGREEMENTS["iru"], C.DEAL["iru_prepaid_consideration"], C.CITES["equipment"] + "; owned by Tribe unless Award says otherwise."),
        ("Splice closures, handholes, vaults, pedestals", C.TRIBE_SHORT, "TBCP grant", "Y", "Operate/maintain", C.AGREEMENTS["om"], C.PH("in O&M scope"), C.CITES["equipment"]),
        ("Conduit / innerduct (grant-funded)", C.TRIBE_SHORT, "TBCP grant", "Y", "Use / access", C.AGREEMENTS["iru"], C.PH("valuation"), "Distinguish from LREMC conduit."),
        sect("B. Grant-Funded Electronics, Premises & Support Assets"),
        ("OLTs, switches, routers (grant-funded)", C.TRIBE_SHORT, "TBCP grant", "Y", "Operate/maintain", C.AGREEMENTS["om"], C.PH("valuation"), C.CITES["equipment"] + "; §889/BABA status. " + C.FLAG_TECH),
        ("ONTs / CPE / drops", C.TRIBE_SHORT, "TBCP grant", "Y", "Install/operate", C.AGREEMENTS["om"], C.PH("valuation"), "Confirm CPE ownership at demarcation."),
        ("Cabinets, huts (grant-funded), power/UPS", C.TRIBE_SHORT, "TBCP grant", "Y", "Use / operate", C.AGREEMENTS["om"], C.PH("valuation"), "Distinguish from LREMC huts/power (Section D)."),
        ("Spares, vehicles, tools (grant-funded)", C.TRIBE_SHORT, "TBCP grant", "Y", "Use", C.AGREEMENTS["om"], C.PH("valuation"), C.CITES["equipment"] + "/" + C.CITES["supplies"]),
        ("Software / licenses (grant-funded)", C.TRIBE_SHORT, "TBCP grant", "Y", "Use / license", C.AGREEMENTS["privacy"], C.PH("license terms"), C.CITES["intangible"]),
        sect("C. RIVR Tech Pre-Existing Assets (stay RIVR Tech's)"),
        ("Existing fiber / core network", C.OPERATOR_SHORT, "RIVR Tech (non-grant)", "N", "Interconnection / transport", C.AGREEMENTS["interconnect"], C.PH("valuation / rate - NEVER free"), C.OPERATOR_EXISTING + ". " + C.CITES["allowable"]),
        ("OLTs / routers / transport (existing)", C.OPERATOR_SHORT, "RIVR Tech (non-grant)", "N", "Interconnection", C.AGREEMENTS["interconnect"], C.PH("valuation / rate"), "Grant capacity never free commercial capacity."),
        ("IP transit / upstream connectivity", C.OPERATOR_SHORT, "RIVR Tech (non-grant)", "N", "Transit service", C.AGREEMENTS["interconnect"], C.PH("rate"), "Confirm contract and cost allocation."),
        ("NOC / systems / personnel", C.OPERATOR_SHORT, "RIVR Tech (non-grant)", "N", "Operations service", C.AGREEMENTS["om"], C.PH("service pricing"), "Cost-allocate shared use."),
        ("Billing / OSS-BSS / software", C.OPERATOR_SHORT, "RIVR Tech / vendors", "N", "Billing services", C.AGREEMENTS["retail"], C.PH("license / rate"), "Data ownership and portability. " + C.CITES["cpni"]),
        ("Customer data / relationships", C.OPERATOR_SHORT, "RIVR Tech", "CONFIRM", "Provider of record", C.AGREEMENTS["retail"], C.PH("n/a - governance terms"), "Project-customer data governance. " + C.FLAG_ATTORNEY),
        sect("D. LREMC Facilities (separate owner - consent/joinder required)"),
        ("Utility poles", C.LREMC_SHORT, "LREMC (non-grant)", "N", "Pole attachment", C.AGREEMENTS["land"], C.PH("attachment rate"), C.CITES["nc_pole"] + "; separate consent. " + C.FLAG_ATTORNEY),
        ("Conduit (LREMC)", C.LREMC_SHORT, "LREMC (non-grant)", "N", "Conduit use", C.AGREEMENTS["land"], C.PH("rate"), "Distinguish from grant-funded conduit."),
        ("Fiber (LREMC-owned)", C.LREMC_SHORT, "LREMC (non-grant)", "N", "Lease / IRU", C.AGREEMENTS["land"], C.PH("rate"), "Not owned by RIVR Tech."),
        ("Easements / ROW", C.LREMC_SHORT, "LREMC (non-grant)", "N", "Access / co-use", C.AGREEMENTS["land"], C.PH("consideration"), C.CITES["indian_row"] + " if Indian land. " + C.FLAG_ATTORNEY),
        ("Land / buildings / huts", C.LREMC_SHORT, "LREMC (non-grant)", "N", "Site lease / license", C.AGREEMENTS["land"], C.PH("lease rate"), C.CITES["indian_leasing"] + " if trust land."),
        ("Power / electrical service", C.LREMC_SHORT, "LREMC / utility", "N", "Power service", C.AGREEMENTS["land"], C.PH("service rate"), "Confirm consideration."),
        sect("E. Third-Party Property, Rights & Permits"),
        ("NCDOT / public ROW encroachment", "NCDOT / govt", "n/a", "N", "Encroachment permit", C.AGREEMENTS["land"], C.PH("permit fees"), C.CITES["nc_ncdot"]),
        ("Private landowner easements", "Landowners", "n/a", "N", "Easement", C.AGREEMENTS["land"], C.PH("consideration"), "Per parcel; confirm title."),
        ("Third-party pole owners (non-LREMC)", "Other utilities", "n/a", "N", "Attachment", C.AGREEMENTS["land"], C.PH("attachment rate"), C.CITES["nc_pole"]),
        ("Trust / restricted Indian land (if any)", C.PH("confirm"), "n/a", "CONFIRM", "Lease / ROW", C.AGREEMENTS["land"], C.PH("consideration"), C.CITES["bia_approval"] + ". " + C.FLAG_ATTORNEY),
    ]
    r = hdr + 1
    for row in rows:
        if row[0] == "SECTION":
            C.xl_cell(ws, r, 1, row[1], "section")
            for col in range(2, 9):
                C.xl_cell(ws, r, col, "", "section")
        else:
            asset, owner, funding, fed, rights, instr, consid, notes = row
            C.xl_cell(ws, r, 1, asset, "text", wrap=True)
            C.xl_cell(ws, r, 2, owner, "text", wrap=True)
            C.xl_cell(ws, r, 3, funding, "text", wrap=True)
            fedkind = "output" if fed in ("Y", "N", "CONFIRM") else "text"
            C.xl_cell(ws, r, 4, fed, fedkind, bold=True, align="center")
            C.xl_cell(ws, r, 5, rights, "text", wrap=True)
            C.xl_cell(ws, r, 6, instr, "text", wrap=True)
            C.xl_cell(ws, r, 7, consid, "input", wrap=True)
            C.xl_cell(ws, r, 8, notes, "text", wrap=True)
        r += 1
    C.xl_legend(ws, r + 1)
    C.xl_cell(ws, r + 3, 1, "Reminder: Consideration/Valuation must never be blank for Section C or D assets - grant-funded capacity is never free commercial capacity.", "text", bold=True, wrap=True)

    return C.xl_save(wb, SUBDIR, "07_Asset_Ownership_and_Rights_Matrix.xlsx")


# ===========================================================================
# 8. 08_Term_Sheet_and_Decision_Register.docx
# ===========================================================================
def build_08():
    doc = doc_start("DOCUMENT 08 — TERM SHEET AND DECISION REGISTER",
                    "Term Sheet and Decision Register",
                    "Phase 1 Diligence — Recommended Positions, Alternatives, and Open Decisions",
                    "Term Sheet / Decisions")
    C.article(doc, 1, "Purpose")
    C.para(doc, "This term sheet records, for each open decision, a recommended position, alternatives, required "
                "approvals, a fallback, and the responsible party. It is non-binding and provisional. " + DOC_INTRO_NOTE)

    decisions = [
        ("Payment amount + characterization",
         f"Recommend {C.DEAL['per_subscriber_amount']} as a Per-Active-Subscriber Operating Payment (Option D).",
         "Options A (asset-use/license), B (IRU consideration), C (revenue share), E (hybrid).",
         "Written program-income determination (label does not control). " + C.CITES["prog_income"],
         "If Option D is disallowed, fall back to Option B (defined IRU consideration) or E (hybrid).",
         f"{C.TRIBE_SHORT} / {C.OPERATOR_SHORT}; grants officer"),
        ("Term / renewal",
         f"Recommend {C.IRU_TERM_RECOMMENDED} years from {C.TERM_TRIGGER} + {C.DEAL['iru_renewal']}.",
         "Shorter base term; single renewal; renewal tied to compliance.",
         "Term cannot exceed " + C.DEAL["term_ceiling_note"] + ". " + C.CITES["equipment"],
         "Cap term at the shortest of useful life, land rights, federal-interest period, and the Award.",
         f"{C.TRIBE_SHORT} counsel; grants officer"),
        ("Exclusivity limits",
         "Recommend non-exclusive operating rights within approved scope; no exclusivity over grant capacity beyond the Award.",
         "Limited exclusivity for retail within the Approved Project Area (if permissible).",
         "Confirm no conflict with program requirements. " + C.CITES["prog_income"],
         "No exclusivity if it impairs program purpose or competition.",
         f"{C.TRIBE_SHORT}; grants officer"),
        ("Commercial / non-project capacity",
         "Recommend defining reserved commercial capacity separately, valued and approved before use.",
         "No commercial use; or commercial use only on non-grant assets.",
         "Written commercial-capacity determination. " + C.CITES["prog_income"],
         "If not approved, prohibit commercial use of grant capacity.",
         f"{C.OPERATOR_SHORT}; grants officer"),
        ("Reserved fiber",
         "Recommend reserving defined dark-fiber/strand capacity for the Tribe and public use.",
         "No reservation; reservation by segment.",
         "Confirm against approved design. " + C.PH("strand count - not supplied"),
         "Reserve at least the Tribe's operational needs.",
         f"{C.TRIBE_SHORT}; {C.OPERATOR_SHORT}"),
        ("LREMC scope",
         f"Recommend a separate {C.LREMC_SHORT} consent/joinder for all {C.LREMC_ASSETS} used.",
         "Bundle into master agreement (not recommended - LREMC is a separate owner).",
         "LREMC board approval; arm's-length pricing. " + C.CITES["nc_pole"],
         "No access without executed LREMC instruments.",
         f"{C.LREMC_SHORT}; {C.TRIBE_SHORT}"),
        ("Sovereign-immunity approach",
         "Present bracketed alternatives (no waiver / limited express waiver for defined disputes) for Council decision.",
         "Limited waiver to arbitration with a damages cap; no waiver.",
         "Council approval; express, unequivocal waiver required. " + C.CITES["sovereign_immunity"],
         "Default to no waiver absent Council action.",
         f"{C.TRIBE_SHORT} counsel; Council"),
        ("Governing law / forum",
         "Present bracketed alternatives (Tribal law/forum; federal law; NC law) - do not choose silently.",
         "Arbitration seat; choice-of-law variations.",
         "Consistency with federal award and enforceability. " + C.FLAG_ATTORNEY,
         "Default to negotiated neutral arbitration if unresolved.",
         f"{C.TRIBE_SHORT} counsel"),
        ("Provider of record",
         "Recommend RIVR Tech as retail provider of record, subject to written determination.",
         "Tribe as provider of record with RIVR Tech as operator.",
         "Written operating-appointment determination. " + C.CITES["equipment"],
         "If not approved, Tribe-controlled retail with RIVR Tech operations only.",
         f"{C.TRIBE_SHORT}; {C.OPERATOR_SHORT}; grants officer"),
        ("Affordability",
         "Recommend a low-cost service option consistent with approved commitments.",
         "Program-specific affordability tiers.",
         "Confirm against approved service commitments. " + C.PH("commitments - not supplied"),
         "Adopt the Award's minimum affordability commitment.",
         f"{C.OPERATOR_SHORT}; {C.TRIBE_SHORT}"),
        ("Data rights",
         "Recommend Tribe access to project-customer data; RIVR Tech custodian with portability on transition.",
         "Shared ownership; RIVR Tech ownership with license to Tribe.",
         "CPNI and privacy compliance. " + C.CITES["cpni"],
         "Ensure portability and access on step-in/transition regardless of ownership.",
         f"{C.TRIBE_SHORT}; {C.OPERATOR_SHORT}"),
    ]

    n = 2
    for title, rec, alts, appr, fb, owner in decisions:
        C.article(doc, n, title)
        C.para(doc, "Recommended position: " + rec)
        C.para(doc, "Alternatives: " + alts)
        C.para(doc, "Required approvals: " + appr)
        C.para(doc, "Fallback: " + fb)
        C.para(doc, "Responsible party: " + owner)
        n += 1

    C.article(doc, n, "Decision Register")
    dr_headers = ["ID", "Decision", "Options", "Recommended", "Owner", "Approval Needed", "Status"]
    dr_widths = [0.4, 1.1, 1.5, 1.3, 0.8, 0.9, 0.5]
    dr_rows = []
    approvals_short = ["Program income", "Term ceiling", "Program reqs", "Commercial cap.", "Design",
                       "LREMC board", "Council", "Counsel", "Operating appt.", "Commitments", "CPNI/privacy"]
    rec_short = ["Option D per-sub", f"{C.IRU_TERM_RECOMMENDED}y + 2x5y", "Non-exclusive", "Reserved + valued",
                 "Reserve capacity", "Separate joinder", "Bracketed alts", "Bracketed alts",
                 "RIVR Tech PoR", "Low-cost option", "Tribe access"]
    for i, (title, rec, alts, appr, fb, owner) in enumerate(decisions):
        dr_rows.append([
            f"D{i+1}", title,
            "A-E" if i == 0 else C.PH("see section") if i in (4, 9) else "see section",
            rec_short[i], "Tribe/RIVR" if "/" in owner else owner.split(";")[0],
            approvals_short[i], "OPEN",
        ])
    C.add_table(doc, dr_headers, dr_rows, widths=dr_widths, col_align=["c", "l", "l", "l", "l", "l", "c"])
    C.flag_para(doc, C.FLAG_BUSINESS, "Every decision above is OPEN and non-binding until confirmed by the "
                "responsible party and any required approver.")

    doc_close(doc, n + 1)
    return C.save(doc, SUBDIR, "08_Term_Sheet_and_Decision_Register.docx")


# ===========================================================================
# 9. 09_Federal_Approval_Request.docx
# ===========================================================================
def build_09():
    doc = doc_start("DOCUMENT 09 — FEDERAL APPROVAL REQUEST",
                    "Federal Approval Request",
                    "Phase 1 Diligence — Written Determinations Requested from the Grants Officer",
                    "Federal Approval Request")
    C.article(doc, 1, "Addressee and Purpose")
    C.para(doc, f"To: {C.PH('Federal Program Officer / Grants Officer name and office')}, "
                f"{C.AGENCY_FULL} ({C.AGENCY_SHORT}), {C.DOC_FULL}.")
    C.para(doc, f"From: {C.TRIBE_FULL} ({C.TRIBE_SHORT}), TBCP applicant / recipient, with respect to the "
                f"proposed engagement of {C.OPERATOR_FULL} ({C.OPERATOR_SHORT}).")
    C.para(doc, "Purpose: to request the specific written determinations below before the corresponding "
                "provisions of the definitive agreements are finalized. " + DOC_INTRO_NOTE)
    C.flag_para(doc, C.FLAG_GRANT, "No determination requested below has been obtained. All are PENDING. "
                "This request is a draft for attorney and grants review before submission.")
    C.flag_para(doc, C.FLAG_ATTORNEY, "Confirm the controlling round and award before submission. " + C.NOFO_LIVE_NOTE)

    requests = [
        ("IRU grant to RIVR Tech over Grant-Funded Assets",
         "Long-term rights over federally funded assets may require federal approval and program-income analysis.",
         "IRU limited to defined capacity within the federal-interest period, with the Tribe retaining title.",
         "If not approved as an IRU, a shorter capacity lease/license.",
         C.CITES["intangible"] + "; " + C.CITES["equipment"]),
        ("Operating appointment / operator role",
         "The operator role over grant-funded network must be permissible under the Award.",
         "RIVR Tech operates and maintains under an O&M agreement; the Tribe retains programmatic control.",
         "Management agreement with narrower delegated authority.",
         C.CITES["equipment"]),
        ("Interconnection arrangement",
         "Interconnection to existing networks affects asset boundaries, cost allocation, and federal interest.",
         "Defined points of interconnection and transport terms at documented rates.",
         "Arm's-length transport/transit agreement outside the grant scope.",
         C.CITES["equipment"] + "; " + C.CITES["allowable"]),
        ("Use of existing (non-grant) RIVR Tech assets",
         "Prevents treating pre-existing assets as a free contribution and ensures cost reasonableness.",
         "Value and cost-allocate existing assets used for the project; never free.",
         "Exclude existing assets or lease them at a documented rate.",
         C.CITES["allowable"]),
        ("Customer receipts / revenue handling",
         "Determines who collects retail receipts and how they are retained and accessed.",
         "RIVR Tech collects as provider of record with segregated records and Tribe access.",
         "Tribe-controlled account with RIVR Tech remitting collections.",
         C.CITES["records"] + "; " + C.CITES["prog_income"]),
        ("Per-subscriber payment characterization (program income)",
         "Program-income treatment turns on the Award and this determination, not the contractual label.",
         C.DEAL["per_subscriber_amount"] + " as a Per-Active-Subscriber Operating Payment (Option D).",
         "Options A/B/C/E if Option D is not accepted.",
         C.CITES["prog_income"]),
        ("Commercial / non-project capacity use",
         "Commercial use of grant capacity may generate program income or be prohibited.",
         "Defined reserved commercial capacity, valued and approved before use.",
         "No commercial use of grant capacity.",
         C.CITES["prog_income"] + "; " + C.CITES["equipment"]),
        ("Contractor pricing / margin reasonableness",
         "Costs charged to the award must be reasonable, allocable, and allowable.",
         "Cost/price analysis supporting RIVR Tech pricing and margin on record.",
         "Renegotiate pricing or competitively re-solicit.",
         C.CITES["procurement"] + "; " + C.CITES["allowable"]),
        ("Out-of-scope service",
         "Services beyond the approved scope may be unallowable or require prior approval.",
         "Confine funded work to the approved scope; document any additions.",
         "Handle additions through a separate non-grant commercial arrangement.",
         C.CITES["allowable"]),
    ]

    C.article(doc, 2, "Determinations Requested")
    hdr = ["#", "Determination Requested", "Why It Matters", "Proposed Structure", "Fallback", "Citation", "Status"]
    widths = [0.3, 1.2, 1.4, 1.3, 1.1, 0.9, 0.3]
    rows = []
    for i, (title, why, struct, fb, cite) in enumerate(requests):
        rows.append([str(i + 1), title, why, struct, fb, cite, "PENDING"])
    C.add_table(doc, hdr, rows, widths=widths, col_align=["c", "l", "l", "l", "l", "l", "c"])

    C.article(doc, 3, "Requested Action and Next Steps")
    C.para(doc, "The Tribe requests written determinations on each item above, or guidance on the information "
                "NTIA requires to issue them. The Tribe will supply the executed Award, approved application, "
                "budget, routes, and supporting records upon request.")
    C.flag_para(doc, C.FLAG_GRANT, "Threshold matter: confirm the Tribe's TBCP eligibility in light of its "
                "federal-recognition status. " + C.CITES["lumbee_act"])

    doc_close(doc, 4)
    return C.save(doc, SUBDIR, "09_Federal_Approval_Request.docx")


# ===========================================================================
# MAIN
# ===========================================================================
if __name__ == "__main__":
    builders = [build_01, build_02, build_03, build_04, build_05,
                build_06, build_07, build_08, build_09]
    paths = []
    for b in builders:
        p = b()
        paths.append(p)
        print("BUILT ->", p)
    print("\nALL 9 PHASE 1 FILES GENERATED.")
