"""
gen_schedules.py — Generator for the 14 SCHEDULES (S1–S14) of the coordinated
TBCP fiber package for the Lumbee Tribe of North Carolina and RIVR Tech.

Uses the canonical engine (Scripts/common.py) for ALL formatting, defined terms,
placeholders, flags, citations, and deal mechanics. Single source of truth.

Run:  python3 gen_schedules.py
Output:  ../05_Schedules/*.docx + *.xlsx
"""
from __future__ import annotations

import sys
sys.path.insert(0, '/home/user/claude-skills/Lumbee_RIVR_TBCP_Agreement_Package/Scripts')
import common as C  # noqa: E402

from openpyxl import Workbook  # noqa: E402


# ---------------------------------------------------------------------------
# Shared helpers (thin wrappers over the canonical engine)
# ---------------------------------------------------------------------------
DISCLAIMER = (
    "This Schedule is a working draft prepared for discussion only. It is not "
    "executed, is not final legal advice, and does not create binding obligations. "
    "No controlling TBCP award or source documents were supplied; the provisional "
    "baseline is the " + C.NOFO_NAME + ". " + C.NOFO_LIVE_NOTE + " Every bracketed "
    "placeholder, flag, and alternative must be resolved by the responsible legal, "
    "grant, financial, and technical reviewers before finalization."
)


def start_docx(sn: str):
    """New Word doc with cover, header/footer, banner, and precedence/definitions note."""
    title = C.SCHEDULES[sn]
    doc = C.new_doc()
    C.add_cover(doc, f"SCHEDULE {sn} — {title}", title)
    C.setup_header_footer(doc, f"Schedule {sn}")
    C.status_banner(doc)
    C.spacer(doc, 1)
    C.para(doc,
           f"This is Schedule {sn} ({title}) to the {C.AGREEMENTS['master']}. "
           f"{C.SHARED_DEFINITIONS_RULE}",
           italic=True)
    C.para(doc, C.PRECEDENCE_RULE, italic=True)
    C.spacer(doc, 1)
    return doc


def end_docx(doc):
    C.spacer(doc, 1)
    p = C.section(doc, "", "Draft Disclaimer")
    C.para(doc, DISCLAIMER, italic=True)
    C.status_banner(doc)


def start_xlsx(sn: str, extra_lines=None):
    """New workbook whose Instructions sheet carries context + baseline + legend."""
    title = C.SCHEDULES[sn]
    wb = Workbook()
    lines = [
        f"## SCHEDULE {sn} — {title}",
        C.DRAFT_STATUS,
        "",
        "## Provenance / caveat",
        "• This register is NOT derived from a supplied TBCP award or source documents.",
        f"• Provisional baseline: {C.NOFO_NAME}.",
        f"• {C.NOFO_LIVE_NOTE}",
        f"• {C.SHARED_DEFINITIONS_RULE}",
        f"• Order-of-precedence rule: {C.PRECEDENCE_RULE}",
        "",
        "## How to use this workbook",
        "• Cells marked [■ ...] are placeholders — replace with confirmed data; do not invent values.",
        "• Amber = INPUT (edit). Blue = FORMULA (do not edit). Green = OUTPUT/RESULT.",
        "• All figures, counts, dates, prices, and descriptions are placeholders pending confirmation.",
    ]
    if extra_lines:
        lines += [""] + extra_lines
    C.xl_instructions(wb, lines)
    return wb, title


def ws_new(wb, name, tab="1F3864"):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    return ws


def fill_register(ws, title, headers, rows, widths, subtitle=None, start=5):
    """Standard register layout: title (1-3), header row @start, PH-aware body."""
    C.xl_title(ws, title, subtitle=subtitle, span=len(headers))
    C.xl_header_row(ws, start, headers)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[C.get_column_letter(j)].width = w
    r = start + 1
    for row in rows:
        for j, val in enumerate(row, start=1):
            kind = "input" if isinstance(val, str) and val.startswith("[■") else "text"
            C.xl_cell(ws, r, j, val, kind=kind, wrap=True, align="left")
        r += 1
    return r


# ===========================================================================
#  S1 — Award, Authority, Resolutions, Signatories, and Document Precedence
# ===========================================================================
def build_s1():
    sn = "S1"
    doc = start_docx(sn)

    C.article(doc, 1, "Award and Authority Summary")
    C.para(doc,
           "The following summary is PRELIMINARY. No controlling TBCP award, NOFO "
           "package, or source documents were supplied. Every field below is a "
           "placeholder to be populated from the executed award and confirmed before "
           "any definitive clause is finalized.")
    C.add_table(
        doc,
        ["Item", "Value (to be confirmed)"],
        [
            ["Recipient / Award Steward", f"{C.TRIBE_FULL} (“{C.TRIBE_SHORT}”)"],
            ["Federal program", f"{C.PROGRAM_FULL} ({C.PROGRAM_SHORT})"],
            ["Awarding agency", f"{C.AGENCY_FULL} ({C.AGENCY_SHORT}), {C.DOC_FULL}"],
            ["Controlling NOFO (provisional)", C.NOFO_NAME],
            ["Federal Award Identification Number", C.PH("award number — NOT SUPPLIED; do not invent")],
            ["Award amount", C.DEAL["ph_award_amount"]],
            ["Approved Project Area", C.PH("approved project area / service territory")],
            ["Locations / homes passed", C.DEAL["ph_homes_passed"]],
            ["Period of performance", C.PH("period of performance start/end")],
            ["Federal Interest period", C.PH("Federal Interest period / useful life")],
            ["Operator / provider of record", f"{C.OPERATOR_FULL} (“{C.OPERATOR_SHORT}”)"],
            ["Separate facilities owner", f"{C.LREMC_FULL} (“{C.LREMC_SHORT}”)"],
        ],
        widths=[2.4, 4.1],
    )
    C.flag_para(doc, C.FLAG_GRANT,
                "Confirm every award field against the executed award, the controlling "
                "round/NOFO, and the Specific Award Conditions before finalizing. "
                f"See {C.CITES['sac']} and {C.CITES['nofo']}.")
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Confirm the recipient of record and any designation of "
                f"{C.TRIBE_ENTITY_ALT} as the contracting Tribal party.")

    C.article(doc, 2, "Authorized Signatories")
    C.para(doc, "Signature authority and titles are placeholders pending governing "
                "resolutions and corporate consents.")
    C.add_table(
        doc,
        ["Party", "Signatory", "Title", "Authority basis"],
        [
            [C.TRIBE_FULL, C.PH("authorized Tribal signatory"),
             C.PH("e.g., Chairman / Tribal Administrator"),
             C.PH("Tribal Council resolution no. & date")],
            [C.OPERATOR_FULL, C.PH("authorized RIVR Tech signatory"),
             C.PH("e.g., CEO / Manager"),
             C.PH("board/member consent no. & date")],
            [C.LREMC_FULL, C.PH("authorized LREMC signatory"),
             C.PH("e.g., CEO / General Manager"),
             C.PH("LREMC board resolution no. & date")],
        ],
        widths=[2.3, 1.5, 1.4, 1.6],
    )
    C.flag_para(doc, C.FLAG_ATTORNEY,
                f"{C.LREMC_SHORT} signs only as owner of the {C.LREMC_ASSETS} via a "
                "separate owner-approved agreement or joinder; it is not a party to the "
                "core Tribe/RIVR Tech commercial terms except where it consents to use "
                "of its facilities.")

    C.article(doc, 3, "Governing Resolutions and Consents")
    C.para(doc, "The following authorizing instruments are required and are all "
                "PENDING; none has been supplied or confirmed:")
    for item in [
        f"{C.TRIBE_SHORT} — Tribal Council resolution authorizing the {C.PROGRAM_SHORT} "
        f"partnership, the {C.AGREEMENTS['master']}, and the acceptance/stewardship of "
        f"the {C.GRANT_FUNDED} " + C.PH("resolution number & date"),
        f"{C.OPERATOR_SHORT} — member/board consent authorizing execution and performance "
        + C.PH("consent number & date"),
        f"{C.LREMC_SHORT} — board resolution authorizing use of the {C.LREMC_ASSETS} and "
        f"any joinder " + C.PH("resolution number & date"),
        "Any lender consent, subordination, or collateral-access acknowledgment "
        + C.PH("lender & instrument"),
    ]:
        C.bullet(doc, item)
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Confirm quorum, notice, and authority for each resolution; attach "
                "certified copies as sub-exhibits when obtained.")

    C.article(doc, 4, "Order of Precedence")
    C.para(doc, "In the event of a conflict among the instruments comprising the "
                "transaction, the following order of precedence controls, from highest "
                "to lowest:")
    for i, item in enumerate(C.PRECEDENCE, start=1):
        C.numbered(doc, item)
    C.section(doc, "4.1", "Precedence Rule", C.PRECEDENCE_RULE)
    C.section(doc, "4.2", "Shared Definitions", C.SHARED_DEFINITIONS_RULE)
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Sovereign immunity, limited waiver (if any), dispute resolution, and "
                "governing law are bracketed ALTERNATIVES to be selected by counsel. "
                f"See {C.CITES['sovereign_immunity']}.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S1_Award_Authority_Resolutions_and_Precedence.docx")


# ===========================================================================
#  S2 — Approved Project Scope, Eligible Locations, Routes, Milestones, Service
# ===========================================================================
def build_s2():
    sn = "S2"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Scope note",
        f"• Service floor is {C.SPEED_FLOOR} per TBCP guidance (confirm against the controlling award).",
        "• Locations, routes, and milestones are NOT from a supplied award — all placeholders.",
    ])

    ws = ws_new(wb, "Eligible Locations")
    fill_register(
        ws, f"S2.1 — Eligible Locations ({C.SERVICE_TERRITORY})",
        ["Location ID", "Address / Description", "Location Type", "Eligibility Basis", "Status"],
        [[C.PH("loc ID"), C.PH("address / parcel description"),
          C.PH("residential / business / community anchor"),
          C.PH("unserved / underserved basis per award"), C.PH("planned")]
         for _ in range(6)],
        widths=[14, 44, 22, 28, 14],
        subtitle="Placeholder location inventory — populate from the approved application.")

    ws = ws_new(wb, "Routes")
    fill_register(
        ws, "S2.2 — Fiber Routes / Segments",
        ["Route/Segment ID", "From (POI/Node)", "To (Node/Location)", "Route Miles",
         "Fiber Count", "Build Type", "Status"],
        [[C.PH("route ID"), C.PH("origin"), C.PH("terminus"), C.PH("miles"),
          C.PH("count"), C.PH("aerial / underground"), C.PH("planned")]
         for _ in range(6)],
        widths=[16, 22, 24, 12, 12, 20, 12],
        subtitle="Placeholder route table — populate from approved engineering routes.")

    ws = ws_new(wb, "Milestones")
    fill_register(
        ws, "S2.3 — Deployment Milestones",
        ["Milestone", "Description", "Target Date", "Dependency", "Owner", "Status"],
        [["M1 — Design complete", C.PH("scope"), C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")],
         ["M2 — Permits/consents secured", C.PH("scope"), C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")],
         ["M3 — Construction start", C.PH("scope"), C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")],
         ["M4 — Segment acceptance", "Triggers term per segment ("
          + C.DEAL["term_trigger"] + ")", C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")],
         ["M5 — Service commencement", C.PH("scope"), C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")],
         ["M6 — Full deployment / closeout", C.PH("scope"), C.PH("date"), C.PH("dep"), C.PH("owner"), C.PH("planned")]],
        widths=[26, 30, 14, 20, 16, 12],
        subtitle="Placeholder milestone schedule — dates from the approved award timeline.")

    ws = ws_new(wb, "Service Commitments")
    fill_register(
        ws, "S2.4 — Service Commitments",
        ["Commitment", "Required Value", "Source / Basis", "Status"],
        [["Minimum broadband service floor", C.SPEED_FLOOR,
          "TBCP guidance (confirm vs. award)", C.PH("committed")],
         ["Affordable service offering", "Includes a " + C.SPEED_FLOOR + " affordable tier",
          "TBCP affordability requirement", C.PH("committed")],
         ["Latency / performance", C.PH("latency & performance commitments"),
          "Award / NOFO", C.PH("committed")],
         ["Deployment completion", C.PH("required completion date"),
          "Award period of performance", C.PH("committed")],
         ["Locations served", C.DEAL["ph_homes_passed"],
          "Approved application", C.PH("committed")]],
        widths=[30, 34, 30, 14],
        subtitle=f"Service floor {C.SPEED_FLOOR}; all values placeholder pending the controlling award.")

    return C.xl_save(wb, "05_Schedules", "S2_Approved_Scope_Locations_Routes_Milestones.xlsx")


# ===========================================================================
#  S3 — Asset Ownership, Funding Source, Federal Interest, IRU/Network-Rights
# ===========================================================================
def build_s3():
    sn = "S3"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Ownership rule",
        f"• {C.GRANT_FUNDED} are owned by the {C.TRIBE_SHORT} and carry Federal Interest = YES.",
        f"• {C.OPERATOR_EXISTING} stays {C.OPERATOR_SHORT}'s property and is NEVER free to the project.",
        f"• {C.LREMC_ASSETS} are owned by {C.LREMC_SHORT} (separate owner) — used only by owner-approved agreement/joinder.",
    ])
    ws = ws_new(wb, "Asset Register")
    headers = ["Asset / Facility", "Owner", "Funding Source", "Federal Interest (Y/N)",
               "IRU / Network Right", "Consideration / Valuation", "Notes"]
    rows = [
        ["Grant-funded FTTH plant (fiber, ONTs, cabinets)", C.TRIBE_SHORT,
         "TBCP grant", "YES", "IRU granted to " + C.OPERATOR_SHORT + " to operate",
         C.DEAL["iru_prepaid_consideration"],
         "Federal Interest applies; disposition per 2 CFR 200.311/.313"],
        ["Grant-funded electronics / OLT / core (grant-funded portion)", C.TRIBE_SHORT,
         "TBCP grant", "YES", "IRU / operating right",
         C.PH("valuation / cost basis"), "Federal Interest applies"],
        ["Grant-funded huts / shelters / equipment (grant-funded)", C.TRIBE_SHORT,
         "TBCP grant", "YES", "IRU / use right", C.PH("valuation"),
         "Federal Interest applies"],
        [C.OPERATOR_EXISTING + " (fiber/core/transport/NOC/billing)", C.OPERATOR_SHORT,
         "RIVR Tech private capital", "NO",
         "Retained by RIVR Tech; access to project via interconnection agreement",
         "Not free — priced/allocated per " + C.AGREEMENTS["interconnect"],
         "Never contributed free; RIVR Tech Assets remain RIVR Tech's"],
        [C.LREMC_ASSETS + " (poles/conduit/fiber/easements/huts/power/land)", C.LREMC_SHORT,
         "LREMC / non-grant", "NO",
         "Used only under separate owner-approved agreement or joinder",
         C.PH("attachment / lease / IRU consideration to LREMC"),
         "Separate owner; not assumed owned by RIVR Tech or the Tribe"],
        [C.JOINT_ASSETS + " (if any cost-shared)", C.PH("Tribe / RIVR Tech allocation"),
         C.PH("grant + private split"), C.PH("Y/N by portion"),
         C.PH("allocated right"), C.PH("allocated basis"),
         "Federal Interest tracks the grant-funded portion; document cost allocation"],
        ["Third-party assets (carrier/backhaul/tower)", C.PH("third-party owner"),
         C.PH("third-party"), "NO", C.PH("lease / IRU / interconnection"),
         C.PH("consideration"), "Confirm no Federal Interest attaches"],
    ]
    fill_register(
        ws, "S3 — Asset Ownership, Funding, Federal Interest & Rights Register",
        headers, rows,
        widths=[34, 14, 20, 14, 30, 30, 34],
        subtitle="Grant-funded = Tribe-owned, Federal Interest YES. Placeholders pending confirmed asset inventory.")
    # legend already on Instructions; add an inline note row
    note_row = 5 + 1 + len(rows) + 1
    C.xl_cell(ws, note_row, 1,
              "NOTE: Federal Interest disposition/encumbrance is governed by 2 CFR 200.311 "
              "(real property), 200.313 (equipment), and 200.316 (property trust relationship). "
              "Consideration for grant-funded capacity is never $0 where an IRU is granted.",
              kind="text", wrap=True)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    return C.xl_save(wb, "05_Schedules", "S3_Asset_Ownership_FederalInterest_and_IRU_Register.xlsx")


# ===========================================================================
#  S4 — Engineering Standards, BOM, Equipment, Interconnection, Demarcation
# ===========================================================================
def build_s4():
    sn = "S4"
    doc = start_docx(sn)

    C.article(doc, 1, "Engineering and Technology Standards")
    C.para(doc, "The Grant-Funded Assets shall be designed, engineered, and constructed "
                "as a fiber-to-the-home (FTTH) passive optical network meeting the "
                "following standards, subject to technical confirmation:")
    for item in [
        "FTTH architecture using XGS-PON (10 Gbps symmetrical) as the target PON "
        "generation, backward-compatible with GPON where required " + C.FLAG_TECH,
        "Calix-class (or functionally equivalent, §889-compliant) OLT/ONT platform "
        + C.FLAG_TECH,
        "ITU-T G.9807.1 (XGS-PON) and ITU-T G.984 (GPON) conformance " + C.FLAG_TECH,
        "Split ratios, wavelength plan, and power budget " + C.PH("split ratio / wavelength / budget"),
        "BABA / domestic-preference and §889 supply-chain compliance for all equipment "
        f"(see {C.CITES['baba']}; {C.CITES['telecom_ban']})",
    ]:
        C.bullet(doc, item)
    C.flag_para(doc, C.FLAG_TECH,
                "Confirm PON generation, vendor/platform, split ratios, and power budget "
                "with the design engineer of record.")

    C.article(doc, 2, "Bill of Materials (Categories)")
    C.para(doc, "BOM categories (line-item quantities and part numbers are placeholders "
                "pending final design):")
    C.add_table(
        doc,
        ["Category", "Representative Items", "Quantity"],
        [
            ["Outside plant fiber", "Feeder/distribution/drop cable, strand, hardware", C.PH("qty")],
            ["Passive optics", "Splitters, splice enclosures, terminals, pedestals", C.PH("qty")],
            ["Active electronics", "OLT chassis/cards, ONTs/ONUs, optics", C.PH("qty")],
            ["Cabinets/huts/power", "Cabinets, shelters, rectifiers, batteries, generators", C.PH("qty")],
            ["Interconnection", "Cross-connects, patch panels, transport optics", C.PH("qty")],
            ["Test/acceptance", "OTDR, power meters, launch fiber", C.PH("qty")],
        ],
        widths=[1.8, 3.4, 1.3],
    )

    C.article(doc, 3, "Interconnection Points and Demarcation")
    C.section(doc, "3.1", "Points of Interconnection",
              f"The {C.NETWORK} interconnects with the {C.OPERATOR_EXISTING} and any "
              f"third-party transport at defined {C.POI} locations. Each {C.POI} and its "
              "hand-off specification is a placeholder pending design: "
              + C.PH("POI locations and hand-off specs") + ".")
    C.section(doc, "3.2", "Demarcation",
              f"The {C.DEMARCATION} between {C.GRANT_FUNDED} and the {C.OPERATOR_EXISTING} "
              f"(and between the {C.NETWORK} and {C.LREMC_ASSETS}) shall be defined at each "
              "interface so that ownership, Federal Interest, maintenance, and risk are "
              "unambiguous. " + C.PH("demarcation schematic / interface list") + ".")
    C.flag_para(doc, C.FLAG_TECH,
                "Attach the demarcation schematic and interface matrix; confirm every "
                "ownership boundary against Schedule S3.")

    C.article(doc, 4, "Optical Acceptance Testing")
    C.para(doc, "Acceptance testing of the fiber plant shall meet the following optical "
                "criteria before segment acceptance (which triggers the term per "
                + C.DEAL["term_trigger"] + "):")
    C.add_table(
        doc,
        ["Parameter", "Acceptance Criterion"],
        [
            ["Test wavelengths", C.OPTICAL["wavelengths"]],
            ["Test method", C.OPTICAL["otdr"]],
            ["Splice loss (max)", C.OPTICAL["splice_loss_max"]],
            ["Connector loss (max)", C.OPTICAL["connector_loss_max"]],
            ["Span-loss budget", C.OPTICAL["span_margin"]],
            ["Reflectance (max)", C.OPTICAL["reflectance_max"]],
        ],
        widths=[2.0, 4.5],
    )
    C.flag_para(doc, C.FLAG_TECH,
                "Optical acceptance results are prerequisites to segment acceptance and to "
                "any construction milestone payment under Schedule S5.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S4_Engineering_Standards_BOM_Interconnection_Demarcation.docx")


# ===========================================================================
#  S5 — Construction Pricing, Payment Milestones, Retainage, Warranties
# ===========================================================================
def build_s5():
    sn = "S5"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Cost / duplication rule",
        "• Prices are eligible-cost based per 2 CFR Part 200 Subpart E (allowable/allocable/reasonable).",
        "• No cost may be charged twice — construction cost recovered here is not also recovered via the",
        "  per-Active-Subscriber operating payment or via RIVR Tech Existing Network charges.",
        "• All prices, percentages, and warranty terms are placeholders pending competitive procurement.",
    ])
    ws = ws_new(wb, "Construction Pricing")
    C.xl_title(ws, "S5 — Construction Pricing, Milestones, Retainage & Warranties",
               subtitle="Eligible-cost based; placeholders pending procurement. Totals computed by formula.",
               span=6)
    headers = ["Milestone", "Eligible-Cost Basis", "Price", "Retainage %", "Retainage Held", "Warranty"]
    start = 5
    C.xl_header_row(ws, start, headers)
    widths = [30, 34, 16, 12, 16, 26]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[C.get_column_letter(j)].width = w

    milestones = [
        ("M1 — Design & engineering", "Design labor, survey, permitting (allocable)"),
        ("M2 — Materials on site", "BABA-compliant materials at documented cost"),
        ("M3 — Construction 50%", "Installed plant, verified progress"),
        ("M4 — Substantial completion", "Segment ready for acceptance testing"),
        ("M5 — Final acceptance", "OTDR acceptance passed (Schedule S4)"),
    ]
    r = start + 1
    first_data = r
    for name, basis in milestones:
        C.xl_cell(ws, r, 1, name, kind="text", wrap=True, align="left")
        C.xl_cell(ws, r, 2, basis, kind="text", wrap=True, align="left")
        C.xl_cell(ws, r, 3, C.PH("price"), kind="input", align="left")
        C.xl_cell(ws, r, 4, C.PH("%"), kind="input", align="center")
        # Retainage held = price * retainage% — formula (placeholders keep it inert but colored)
        C.xl_cell(ws, r, 5, f"=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),C{r}*D{r},\"\")",
                  kind="formula", fmt=C.FMT_USD, align="right")
        C.xl_cell(ws, r, 6, C.PH("warranty term"), kind="input", wrap=True, align="left")
        r += 1
    last_data = r - 1

    # Totals row (formula/output)
    C.xl_cell(ws, r, 1, "TOTAL", kind="section", align="left")
    C.xl_cell(ws, r, 2, "", kind="section")
    C.xl_cell(ws, r, 3, f"=IF(COUNT(C{first_data}:C{last_data})>0,SUM(C{first_data}:C{last_data}),\"\")",
              kind="output", fmt=C.FMT_USD, align="right", bold=True)
    C.xl_cell(ws, r, 4, "", kind="section")
    C.xl_cell(ws, r, 5, f"=IF(COUNT(E{first_data}:E{last_data})>0,SUM(E{first_data}:E{last_data}),\"\")",
              kind="output", fmt=C.FMT_USD, align="right", bold=True)
    C.xl_cell(ws, r, 6, "", kind="section")
    total_row = r

    # Net payable
    r += 1
    C.xl_cell(ws, r, 1, "NET PAYABLE (Total less Retainage held)", kind="section", align="left")
    for cc in (2, 4, 6):
        C.xl_cell(ws, r, cc, "", kind="section")
    C.xl_cell(ws, r, 3, f"=IF(ISNUMBER(C{total_row}),C{total_row}-E{total_row},\"\")",
              kind="output", fmt=C.FMT_USD, align="right", bold=True)
    C.xl_cell(ws, r, 5, "", kind="section")

    # No-duplication check
    r += 2
    C.xl_check(ws, r, 1,
               "MANUAL VERIFY: no cost recovered here is also recovered via the per-Active-Subscriber "
               "payment (S8) or RIVR Tech Existing Network charges (S3).",
               "No-duplicate-charge check:")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

    return C.xl_save(wb, "05_Schedules", "S5_Construction_Pricing_Milestones_Retainage_Warranties.xlsx")


# ===========================================================================
#  S6 — Operations, Service Levels, Lifecycle, Spares, Refresh, Reserves
# ===========================================================================
def build_s6():
    sn = "S6"
    doc = start_docx(sn)
    sla = C.DEAL["sla"]

    C.article(doc, 1, "Service Level Agreement (Severity Matrix)")
    C.para(doc, f"{C.OPERATOR_SHORT} shall operate and maintain the {C.NETWORK} to the "
                f"following service levels. The Network Operations Center operates "
                f"{sla['noc']}.")
    C.add_table(
        doc,
        ["Severity", "Definition", "Response", "Dispatch", "Restore"],
        [
            ["P1 — Critical", "Total/major outage; Essential Services down",
             f"{sla['P1_response_min']} min", f"{sla['P1_dispatch_hr']} hr", f"{sla['P1_restore_hr']} hr"],
            ["P2 — High", "Partial outage / significant degradation",
             f"{sla['P2_response_min']} min", f"{sla['P2_dispatch_hr']} hr", f"{sla['P2_restore_hr']} hr"],
            ["P3 — Medium", "Limited impact; workaround available",
             f"{sla['P3_response_hr']} hr", f"{sla['P3_dispatch_hr']} hr", f"{sla['P3_restore_days']} days"],
            ["P4 — Low", "Minor / single-subscriber / cosmetic",
             f"{sla['P4_response_hr']} hr", "scheduled", f"{sla['P4_restore_days']} days"],
        ],
        widths=[1.3, 2.6, 1.0, 0.9, 0.9],
    )
    C.section(doc, "1.1", "Performance Targets",
              f"Availability target {sla['availability_target']}; latency ≤ {sla['latency_ms']} ms; "
              f"packet loss ≤ {sla['packet_loss']}; jitter ≤ {sla['jitter_ms']} ms. "
              "Measurement methodology and service credits are "
              + C.PH("SLA measurement window & service-credit schedule") + ".")

    C.article(doc, 2, "Lifecycle Management")
    for item in [
        "Preventive maintenance program with documented intervals and records "
        + C.PH("PM intervals"),
        "Fiber and electronics lifecycle tracking against manufacturer useful life "
        + C.PH("useful life by asset class"),
        "Configuration, firmware, and patch management aligned with Schedule S11",
        "Asset records reconciled to the Schedule S3 register and Federal Interest tracking",
    ]:
        C.bullet(doc, item)

    C.article(doc, 3, "Spares and Refresh")
    C.section(doc, "3.1", "Spares",
              "A critical-spares inventory (OLT cards, ONTs, optics, splice materials, "
              "power components) shall be maintained at levels sufficient to meet the P1/P2 "
              "restore targets. Minimum spare quantities: "
              + C.PH("spares levels by part") + ".")
    C.section(doc, "3.2", "Technology Refresh",
              "Active electronics shall follow a refresh cycle of "
              + C.PH("refresh cycle, e.g., 7–10 years") + ", coordinated with the useful "
              "life of the Grant-Funded Assets and any Federal Interest constraints on "
              "disposition of replaced equipment.")

    C.article(doc, 4, "Replacement Reserve")
    C.para(doc, "A replacement reserve shall be funded to ensure long-term sustainability "
                "of the Grant-Funded Assets across the term ("
                + C.DEAL["iru_term_recommended"].__str__() + " years from "
                + C.DEAL["term_trigger"] + ", plus " + C.DEAL["iru_renewal"] + ").")
    C.add_table(
        doc,
        ["Reserve Element", "Basis", "Amount / Rate"],
        [
            ["Annual reserve contribution", C.PH("basis"), C.PH("$ or % of revenue")],
            ["Reserve funding source", "Operating revenue / per-subscriber payment", C.PH("allocation")],
            ["Reserve floor / target balance", C.PH("basis"), C.PH("target balance")],
            ["Governance of reserve draws", "Joint approval per Master Agreement", C.PH("approval threshold")],
        ],
        widths=[2.2, 2.6, 1.7],
    )
    C.flag_para(doc, C.FLAG_GRANT,
                "Reserve funding and any use of Program Income for reserves must comply "
                f"with {C.CITES['prog_income']} and the controlling award.")
    C.flag_para(doc, C.FLAG_BUSINESS,
                "Reserve rate, refresh cycle, and spares levels are business decisions "
                "requiring RIVR Tech and Tribe agreement.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S6_Operations_SLA_Lifecycle_Spares_Refresh_Reserves.docx")


# ===========================================================================
#  S7 — Provider-of-Record, Products, Affordability, Support, Billing, Complaints
# ===========================================================================
def build_s7():
    sn = "S7"
    doc = start_docx(sn)

    C.article(doc, 1, "Provider of Record")
    C.para(doc, f"{C.OPERATOR_FULL} (“{C.OPERATOR_SHORT}”) is the retail provider "
                f"of record for broadband service delivered over the {C.NETWORK} within the "
                f"{C.SERVICE_TERRITORY}. {C.OPERATOR_SHORT} holds the retail customer "
                "relationship, billing, and regulatory responsibility, subject to the "
                f"{C.AGREEMENTS['retail']} and the order of precedence.")

    C.article(doc, 2, "Customer Products and Affordability")
    C.para(doc, "Retail product tiers (pricing is placeholder pending business decision "
                "and any award affordability commitments):")
    C.add_table(
        doc,
        ["Tier", "Speed", "Target Segment", "Monthly Price"],
        [
            ["Affordable / essential", C.SPEED_FLOOR + " (meets service floor)",
             "Eligible / affordability", C.PH("affordable price")],
            ["Standard residential", C.PH("speed"), "Residential", C.PH("price")],
            ["Premium residential", C.PH("speed"), "Residential", C.PH("price")],
            ["Business / anchor", C.PH("speed"), "Business / community anchor", C.PH("price")],
        ],
        widths=[1.7, 2.1, 1.9, 1.3],
    )
    C.flag_para(doc, C.FLAG_BUSINESS,
                f"A {C.SPEED_FLOOR} affordable tier is required; confirm price points, "
                "eligibility, and any low-cost option mandated by the controlling award.")

    C.article(doc, 3, "Customer Support")
    for item in [
        "Support channels and hours " + C.PH("support hours / channels"),
        "Installation, provisioning, and truck-roll standards coordinated with Schedule S6 SLAs",
        "Escalation path to the Network Operations Center for service-affecting issues",
    ]:
        C.bullet(doc, item)

    C.article(doc, 4, "Billing and Collections")
    for item in [
        "Billing cycle, invoicing, payment methods, and delinquency handling "
        + C.PH("billing terms"),
        "Active Subscriber counting for the operating payment is governed by Schedule S8 "
        "and the " + C.AGREEMENTS["retail"],
        "Taxes, surcharges, and regulatory fees separately stated per Schedule S8",
    ]:
        C.bullet(doc, item)

    C.article(doc, 5, "Complaints and Regulatory")
    C.section(doc, "5.1", "Complaint Handling",
              "A documented complaint intake, tracking, and resolution process shall be "
              "maintained, with reporting to the Tribe " + C.PH("complaint SLA & reporting") + ".")
    C.section(doc, "5.2", "CPNI",
              "Customer Proprietary Network Information shall be handled in compliance with "
              f"{C.CITES['cpni']}.")
    C.section(doc, "5.3", "ETC / Lifeline",
              "If Lifeline or ETC-supported service is offered, compliance with "
              f"{C.CITES['usac_lifeline']} is required; ETC designation status is "
              + C.PH("ETC designation status") + ".")
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Confirm state (NCUC) and federal (FCC) regulatory status, tariffing, and "
                "any ETC obligations with regulatory counsel.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S7_Provider_of_Record_Products_Affordability_Billing.docx")


# ===========================================================================
#  S8 — Subscriber Payment, Revenue, Program Income, Taxes, Cash Controls, Audit
# ===========================================================================
def build_s8():
    sn = "S8"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Payment model",
        f"• Primary model: {C.DEAL['pay_options']['D']}.",
        f"• Rate: {C.DEAL['per_subscriber_amount']} (placeholder).",
        "• Characterization alternatives A–E do NOT control program-income treatment.",
        f"• {C.DEAL['program_income_note']} ({C.CITES['prog_income']}).",
    ])

    # Active Subscriber definition
    ws = ws_new(wb, "Active Subscriber Def")
    C.xl_title(ws, "S8.1 — Active Subscriber Definition",
               subtitle="Elements are canonical from common.py; confirm against the Retail Agreement.",
               span=2)
    C.xl_header_row(ws, 5, ["#", "Active Subscriber Element"])
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 110
    r = 6
    for i, el in enumerate(C.DEAL["active_subscriber_elements"], start=1):
        C.xl_cell(ws, r, 1, i, kind="text", align="center")
        C.xl_cell(ws, r, 2, el, kind="text", wrap=True, align="left")
        r += 1

    # Payment mechanics
    ws = ws_new(wb, "Payment Mechanics")
    fill_register(
        ws, "S8.2 — Payment Mechanics",
        ["Element", "Term", "Status"],
        [["Payment basis", C.DEAL["pay_options"]["D"], C.PH("agreed")],
         ["Rate", C.DEAL["per_subscriber_amount"], C.PH("agreed")],
         ["Counting date", "Activation at first provisioning; disconnect at cease", C.PH("agreed")],
         ["Proration", "Partial months prorated daily [BUSINESS DECISION REQUIRED]", C.PH("agreed")],
         ["Payment frequency", C.PH("monthly / quarterly"), C.PH("agreed")],
         ["Reconciliation", C.PH("reconciliation & true-up cadence"), C.PH("agreed")]],
        widths=[26, 60, 16],
        subtitle="Placeholders pending business decision.")

    # Revenue & program income
    ws = ws_new(wb, "Revenue & Program Income")
    fill_register(
        ws, "S8.3 — Revenue Categories & Program-Income Treatment",
        ["Revenue Category", "Description", "Program Income (2 CFR 200.307)?", "Notes"],
        [["Retail broadband revenue", "Subscriber recurring charges",
          C.PH("determination pending"), "Treatment turns on the award, not the label"],
         ["Installation / activation fees", C.PH("scope"), C.PH("pending"), C.PH("note")],
         ["Commercial / non-project capacity", "Approved commercial use (S10)",
          C.PH("pending"), "Grant-funded capacity never free"],
         ["Other fees", C.PH("scope"), C.PH("pending"), C.PH("note")]],
        widths=[28, 34, 30, 40],
        subtitle="Program-income treatment requires a written NTIA determination against the controlling award.")

    # Taxes / cash controls / audit
    ws = ws_new(wb, "Taxes Cash Audit")
    fill_register(
        ws, "S8.4 — Taxes, Fees, Cash Controls & Audit",
        ["Topic", "Requirement", "Status"],
        [["Taxes & regulatory fees", "Separately stated; excluded from Active Subscriber calc", C.PH("confirm")],
         ["Cash controls", "Segregated accounting; documented internal controls", C.PH("confirm")],
         ["Records retention", str(C.DEAL["records_retention_years"]) + " years ("
          + C.CITES["records"] + ")", C.PH("confirm")],
         ["Audit rights", "Access on " + str(C.DEAL["audit_notice_business_days"])
          + " business days' notice; single audit if threshold met (" + C.CITES["single_audit"] + ")",
          C.PH("confirm")]],
        widths=[24, 70, 16],
        subtitle="Cash controls and audit rights align to 2 CFR Part 200 Subparts D/F.")

    # Worked example
    ws = ws_new(wb, "Worked Example")
    C.xl_title(ws, "S8.5 — Worked Example (Per-Active-Subscriber Payment)",
               subtitle="Illustrative only. Rate is a placeholder; formula computes when a number is entered.",
               span=3)
    C.xl_header_row(ws, 5, ["Input", "Value", "Notes"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 44
    C.xl_cell(ws, 6, 1, "Active Subscribers (count)", kind="text", align="left")
    C.xl_cell(ws, 6, 2, C.PH("subscriber count"), kind="input", align="right")
    C.xl_cell(ws, 6, 3, "Per Schedule S8.1 definition", kind="text", wrap=True, align="left")
    C.xl_cell(ws, 7, 1, "Rate per Active Subscriber / month", kind="text", align="left")
    C.xl_cell(ws, 7, 2, C.PH("rate"), kind="input", align="right")
    C.xl_cell(ws, 7, 3, C.DEAL["per_subscriber_amount"], kind="text", wrap=True, align="left")
    C.xl_cell(ws, 8, 1, "Monthly operating payment", kind="text", bold=True, align="left")
    C.xl_cell(ws, 8, 2, "=IF(AND(ISNUMBER(B6),ISNUMBER(B7)),B6*B7,\"\")",
              kind="formula", fmt=C.FMT_USD, align="right")
    C.xl_cell(ws, 8, 3, "= Subscribers x Rate", kind="text", align="left")
    C.xl_cell(ws, 9, 1, "Annual operating payment", kind="text", bold=True, align="left")
    C.xl_cell(ws, 9, 2, "=IF(ISNUMBER(B8),B8*12,\"\")", kind="output", fmt=C.FMT_USD, align="right")
    C.xl_cell(ws, 9, 3, "= Monthly x 12", kind="text", align="left")

    return C.xl_save(wb, "05_Schedules", "S8_Subscriber_Payment_Revenue_ProgramIncome_Tax_Audit.xlsx")


# ===========================================================================
#  S9 — Land, Poles, Sites, Power, Permits, Shared Facilities, Owner Consents
# ===========================================================================
def build_s9():
    sn = "S9"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Land / access rule",
        f"• {C.LREMC_ASSETS} (poles/conduit/easements/huts/power/land) are owned by {C.LREMC_SHORT}.",
        "• Where trust/restricted Indian land is used or crossed, BIA approval may be required:",
        f"    - {C.CITES['indian_leasing']}",
        f"    - {C.CITES['indian_row']}",
        f"    - {C.CITES['bia_approval']}",
        "• All interests, instruments, and consent statuses are placeholders — none obtained.",
    ])
    ws = ws_new(wb, "Land & Facilities Register")
    headers = ["Item / Facility", "Owner", "Interest Type", "Instrument",
               "BIA Approval Needed?", "Term", "Consent Status"]
    rows = [
        ["Pole attachments", C.LREMC_SHORT, "Pole attachment",
         C.PH("attachment agreement"), C.PH("likely N (non-trust) — confirm"),
         C.PH("term"), "PENDING"],
        ["Conduit / duct", C.LREMC_SHORT, "License / IRU",
         C.PH("conduit license"), C.PH("confirm"), C.PH("term"), "PENDING"],
        ["Easements / ROW (utility)", C.LREMC_SHORT, "Easement / ROW",
         C.PH("easement instrument"), C.PH("confirm if trust land crossed"),
         C.PH("term"), "PENDING"],
        ["Fee land parcel(s)", C.PH("owner"), "Fee",
         C.PH("deed / lease"), "N (fee land)", C.PH("term"), "PENDING"],
        ["Trust / restricted land", C.PH("Tribe / allottee"), "Trust / restricted",
         C.PH("BIA lease / ROW"), "YES — BIA approval (25 CFR 162/169)",
         C.PH("term"), "PENDING"],
        ["Hut / shelter sites", C.PH("owner"), C.PH("fee / lease / license"),
         C.PH("site agreement"), C.PH("confirm"), C.PH("term"), "PENDING"],
        ["Power service points", C.LREMC_SHORT, "Service agreement",
         C.PH("power agreement"), "N", C.PH("term"), "PENDING"],
        ["NCDOT encroachment / ROW", "NCDOT", "Encroachment permit",
         C.PH("encroachment agreement"), "N",
         C.PH("term"), "PENDING"],
        ["Railroad / other crossings", C.PH("owner"), C.PH("crossing agreement"),
         C.PH("instrument"), C.PH("confirm"), C.PH("term"), "PENDING"],
    ]
    fill_register(
        ws, "S9 — Land, Poles, Sites, Power, Permits & Consents Register",
        headers, rows,
        widths=[26, 14, 22, 26, 34, 14, 16],
        subtitle="All consents PENDING; none obtained. LREMC-owned facilities used only by owner-approved agreement.")
    note_row = 5 + 1 + len(rows) + 1
    C.xl_cell(ws, note_row, 1,
              "NOTE: NC-specific: pole attachment (N.C. Gen. Stat. § 62-350); 811/locates "
              "(N.C. Gen. Stat. Ch. 87 Art. 8A); NCDOT encroachment (N.C. Gen. Stat. § 136-18). "
              "Trust/restricted Indian land triggers 25 CFR Parts 162/169 and BIA approval.",
              kind="text", wrap=True)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    return C.xl_save(wb, "05_Schedules", "S9_Land_Poles_Sites_Power_Permits_Consents.xlsx")


# ===========================================================================
#  S10 — Approved Commercial Capacity, Non-Project Use, Valuation, Cost Allocation
# ===========================================================================
def build_s10():
    sn = "S10"
    doc = start_docx(sn)

    C.article(doc, 1, "Approved Commercial Capacity")
    C.para(doc, f"Any commercial or non-project use of the {C.GRANT_FUNDED} is permitted "
                "ONLY where all of the following conditions are satisfied:")
    for item in [
        "the use is within the approved geography (Approved Project Area) or otherwise "
        "expressly permitted by the award",
        "the use is consistent with the controlling award and any written NTIA/DOC direction",
        "cost allocation is documented so grant-funded capacity is not subsidized by the "
        f"federal award ({C.CITES['prog_income']}; {C.CITES['allowable']})",
        "federal-property use restrictions are honored ("
        + C.CITES["real_property"] + "; " + C.CITES["intangible"] + ")",
        "written agency direction is obtained where required",
    ]:
        C.bullet(doc, item)
    C.flag_para(doc, C.FLAG_GRANT,
                "Grant-funded capacity is NEVER provided free. Any commercial use must be "
                "priced and cost-allocated, and program-income treatment determined against "
                "the controlling award.")

    C.article(doc, 2, "Non-Project vs. Out-of-Scope Use")
    C.section(doc, "2.1", "Non-Tribal Occupant at an Approved Location",
              "A non-Tribal customer or occupant AT an approved, eligible location "
              "(e.g., a renter, business, or anchor within the Approved Project Area) is "
              "a permitted retail relationship, not an out-of-scope extension, provided "
              "the location itself is award-eligible.")
    C.section(doc, "2.2", "Out-of-Scope Extension",
              "Extending the network to serve locations or areas OUTSIDE the approved "
              "scope is an out-of-scope use that requires separate authority, separate "
              "(non-grant) funding, documented cost allocation, and written agency "
              "direction where required. " + C.PH("out-of-scope extension policy") + ".")
    C.flag_para(doc, C.FLAG_GRANT,
                "Distinguish (a) a non-Tribal occupant at an approved location [permitted] "
                "from (b) an out-of-scope geographic extension [requires separate authority "
                "and non-grant funding].")

    C.article(doc, 3, "Valuation and Cost Allocation")
    C.add_table(
        doc,
        ["Element", "Method / Basis", "Value"],
        [
            ["Commercial capacity pricing", "Market-based, cost-allocated", C.PH("price / rate")],
            ["Cost allocation methodology", "Documented allocation base per 2 CFR 200 Subpart E", C.PH("methodology")],
            ["Program-income treatment", "Per award determination (200.307)", C.PH("determination")],
            ["Approval / documentation", "Written agency direction where required", C.PH("status")],
        ],
        widths=[2.2, 3.0, 1.3],
    )
    C.flag_para(doc, C.FLAG_BUSINESS,
                "Commercial capacity pricing and any wholesale/dark-fiber offering are "
                "business decisions requiring Tribe and RIVR Tech agreement.")
    C.flag_para(doc, C.FLAG_ATTORNEY,
                f"{C.OPERATOR_EXISTING} used to serve commercial demand remains "
                f"{C.OPERATOR_SHORT}'s and is never contributed free; keep it distinct from "
                "grant-funded capacity in every cost allocation.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S10_Commercial_Capacity_NonProject_Use_Valuation.docx")


# ===========================================================================
#  S11 — Data, Privacy, Cybersecurity, Supply-Chain, Incident Response, DR
# ===========================================================================
def build_s11():
    sn = "S11"
    doc = start_docx(sn)
    C.para(doc, f"This Schedule cross-refers to the {C.AGREEMENTS['privacy']}.", italic=True)

    C.article(doc, 1, "Data Classification and Rights by Category")
    C.add_table(
        doc,
        ["Data Category", "Sensitivity", "Owner / Rights", "Handling"],
        [
            ["Subscriber PII / CPNI", "High", "Subscriber; steward per " + C.AGREEMENTS["retail"],
             "CPNI controls (" + C.CITES["cpni"] + ")"],
            ["Network / operational data", "Medium", "Joint per Master Agreement", C.PH("handling")],
            ["Grant / compliance records", "Medium", C.TRIBE_SHORT + " (recipient)",
             "Retain " + str(C.DEAL["records_retention_years"]) + " yrs (" + C.CITES["records"] + ")"],
            ["Tribal data / sovereignty-sensitive", "High", C.TRIBE_SHORT,
             C.PH("Tribal data governance policy")],
        ],
        widths=[1.9, 1.0, 2.2, 2.4],
    )
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Confirm data-ownership, Tribal data sovereignty, and any breach-notice "
                "obligations by category with counsel.")

    C.article(doc, 2, "Cybersecurity Controls")
    for item in [
        "Control framework " + C.PH("e.g., NIST CSF / 800-53 baseline"),
        "Access control, MFA, least privilege, and logging",
        "Encryption in transit and at rest for sensitive data",
        "Vulnerability management, patching (coordinated with Schedule S6 lifecycle)",
        "Security awareness and personnel controls",
    ]:
        C.bullet(doc, item)

    C.article(doc, 3, "Supply-Chain Risk and Section 889")
    C.para(doc, "No covered telecommunications or video-surveillance equipment or services "
                "(Section 889) shall be procured, used, or connected to the Network. "
                f"See {C.CITES['telecom_ban']}. BABA / domestic-preference compliance per "
                f"{C.CITES['baba']}.")
    C.flag_para(doc, C.FLAG_GRANT,
                "Vendor and equipment certifications for §889 and BABA are prerequisites "
                "to procurement; retain certifications with the compliance records.")

    C.article(doc, 4, "Incident Response")
    for item in [
        "Incident classification, escalation, and notification timelines "
        + C.PH("IR notification SLAs"),
        "Breach-notification to affected subscribers and the Tribe per law and the "
        + C.AGREEMENTS["privacy"],
        "Coordination with the Network Operations Center and Schedule S6 severity matrix",
        "Post-incident review and remediation tracking",
    ]:
        C.bullet(doc, item)

    C.article(doc, 5, "Disaster Recovery (RTO / RPO)")
    C.add_table(
        doc,
        ["System / Service", "RTO", "RPO", "Notes"],
        [
            ["Core network / transport", C.PH("RTO"), C.PH("RPO"), "Tied to S6 P1 restore targets"],
            ["OSS/BSS / billing", C.PH("RTO"), C.PH("RPO"), C.PH("note")],
            ["Subscriber data / records", C.PH("RTO"), C.PH("RPO"), "Backup & retention aligned"],
        ],
        widths=[2.4, 1.3, 1.3, 2.5],
    )
    C.flag_para(doc, C.FLAG_TECH,
                "Confirm RTO/RPO targets, backup architecture, and DR test cadence with the "
                "operations team.")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S11_Data_Privacy_Cyber_SupplyChain_IR_DR.docx")


# ===========================================================================
#  S12 — Insurance, Indemnity, Liability Allocation, Claims Procedures
# ===========================================================================
def build_s12():
    sn = "S12"
    doc = start_docx(sn)
    ins = C.DEAL["insurance"]

    C.article(doc, 1, "Insurance Requirements")
    C.para(doc, "Each Party shall maintain, at minimum, the following coverages "
                "(limits are placeholder defaults for negotiation and confirmation against "
                f"the controlling award; see {C.CITES['insurance']}):")
    C.add_table(
        doc,
        ["Coverage", "Minimum Limit"],
        [
            ["Commercial General Liability (per occurrence)", ins["cgl_occurrence"]],
            ["Commercial General Liability (aggregate)", ins["cgl_aggregate"]],
            ["Automobile Liability", ins["auto"]],
            ["Umbrella / Excess Liability", ins["umbrella"]],
            ["Workers' Compensation", ins["workers_comp"]],
            ["Employer's Liability", ins["employers_liability"]],
            ["Professional / Technology E&O", ins["professional_tech_eo"]],
            ["Cyber Liability", ins["cyber"]],
            ["Property / Builder's Risk", ins["property_builders_risk"]],
            ["Pollution / Environmental", ins["pollution"]],
        ],
        widths=[3.6, 2.9],
    )
    C.flag_para(doc, C.FLAG_BUSINESS,
                "Confirm limits, additional-insured and waiver-of-subrogation requirements, "
                "and who insures the Grant-Funded Assets (builder's risk / property).")

    C.article(doc, 2, "Indemnification")
    C.para(doc, "Each Party (the “Indemnifying Party”) shall indemnify, defend, and "
                "hold harmless the other Party from third-party claims to the extent arising "
                "from the Indemnifying Party's negligence, willful misconduct, or breach, "
                "subject to the liability allocation below and to any bracketed limitations "
                "selected by counsel.")
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Uncapped vs. capped indemnity, IP indemnity, and the interaction with "
                "tribal sovereign immunity and any limited waiver are bracketed ALTERNATIVES "
                f"for counsel. See {C.CITES['sovereign_immunity']}.")

    C.article(doc, 3, "Liability Allocation")
    C.para(doc, "Liability is allocated in proportion to each Party's control over, and "
                "responsibility for, the facility, activity, or breach giving rise to the "
                "loss:")
    for item in [
        f"{C.OPERATOR_SHORT} bears liability for operations, retail service, and its "
        f"{C.OPERATOR_EXISTING} it controls",
        f"the {C.TRIBE_SHORT} bears liability tied to its ownership/stewardship role to the "
        "extent within its control",
        f"{C.LREMC_SHORT} bears liability for the {C.LREMC_ASSETS} it owns and controls "
        "(via its separate agreement/joinder)",
        "consequential-damages waiver and any liability cap are "
        + C.PH("cap / carve-outs") + " [ATTORNEY REVIEW REQUIRED]",
    ]:
        C.bullet(doc, item)

    C.article(doc, 4, "Claims Procedure")
    for item in [
        "Prompt written notice of claim with reasonable detail " + C.PH("notice period"),
        "Control of defense and settlement, with cooperation and consent-not-unreasonably-"
        "withheld standards",
        "Coordination with insurers and preservation of coverage",
        "Reservation of rights and no admission by mere notice",
    ]:
        C.numbered(doc, item)

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S12_Insurance_Indemnity_Liability_Claims.docx")


# ===========================================================================
#  S13 — Default, Cure, Step-In, Continuity, Transition, Termination Assistance
# ===========================================================================
def build_s13():
    sn = "S13"
    doc = start_docx(sn)
    C.para(doc, f"This Schedule cross-refers to the {C.AGREEMENTS['transition']}.", italic=True)

    C.article(doc, 1, "Events of Default and Cure")
    C.para(doc, "An Event of Default occurs upon a material breach not cured within the "
                "applicable cure period after written notice:")
    C.add_table(
        doc,
        ["Default Type", "Cure Period", "Extension"],
        [
            ["Monetary default", f"{C.DEAL['cure_monetary_days']} days", "None (unless agreed)"],
            ["Non-monetary default", f"{C.DEAL['cure_nonmonetary_days']} days",
             C.DEAL["cure_nonmonetary_extension"]],
            ["Notice of default", f"{C.DEAL['notice_default_days']} days' written notice", "—"],
        ],
        widths=[2.2, 2.3, 2.9],
    )

    C.article(doc, 2, "Step-In Rights")
    C.section(doc, "2.1", "Triggers",
              "Step-in may be exercised upon an uncured Event of Default that threatens "
              "continuity of Essential Services, public safety, or compliance with the "
              f"Award, and " + C.DEAL["stepin_emergency"] + ".")
    C.section(doc, "2.2", "Limits",
              "Step-in is limited to the scope and duration reasonably necessary to restore "
              "continuity and compliance; it does not transfer ownership of the "
              f"{C.GRANT_FUNDED} and is subject to the order of precedence and the award. "
              + C.PH("step-in scope, duration, and cost-recovery terms") + ".")
    C.flag_para(doc, C.FLAG_ATTORNEY,
                "Step-in scope, cost recovery, and interaction with lender/agency rights and "
                "tribal sovereign immunity are bracketed ALTERNATIVES for counsel.")

    C.article(doc, 3, "Continuity of Service")
    for item in [
        "Essential Services must continue during any dispute, cure, step-in, or transition",
        "No self-help that disconnects subscribers except as permitted by law and the award",
        "Coordination with the Network Operations Center and Schedule S6 SLAs",
    ]:
        C.bullet(doc, item)

    C.article(doc, 4, "Transition Assistance")
    C.para(doc, "Upon expiration or termination, the outgoing operator shall provide "
                "transition assistance for a period of "
                f"{C.DEAL['transition_assistance_months']} months, including:")
    for item in [
        "orderly hand-off of operations, records, and subscriber relationships",
        "knowledge transfer, credentials, configurations, and documentation",
        "continued operation during transition at agreed rates " + C.PH("transition rates"),
        "return/transfer of Grant-Funded Assets and Federal-Interest compliance",
    ]:
        C.bullet(doc, item)
    C.flag_para(doc, C.FLAG_GRANT,
                "Any transfer or disposition of Grant-Funded Assets at transition must comply "
                f"with {C.CITES['real_property']}, {C.CITES['equipment']}, and the award "
                "(Federal Interest survives operator changes).")

    end_docx(doc)
    return C.save(doc, "05_Schedules", "S13_Default_Cure_StepIn_Continuity_Transition.docx")


# ===========================================================================
#  S14 — Required Federal, Tribal, Corporate, Landowner, Lender, Regulatory
# ===========================================================================
def build_s14():
    sn = "S14"
    wb, title = start_xlsx(sn, extra_lines=[
        "## Approval status",
        "• ALL approvals below are PENDING. NONE has been obtained.",
        "• This register does not itself confer any approval.",
    ])
    ws = ws_new(wb, "Required Approvals")
    headers = ["Approval", "Source", "Why Required", "Status", "Owner"]
    rows = [
        ["Award acceptance / any amendment", "NTIA / DOC",
         "Federal award authority and conditions", "PENDING", C.PH("Tribe grant lead")],
        ["Prior approval for arrangement (if required)", "NTIA",
         "Subaward/contract/property arrangements per 2 CFR 200", "PENDING", C.PH("Tribe grant lead")],
        ["BIA approval (trust/restricted land)", "BIA",
         "25 CFR 162/169 leases/ROW on Indian land", "PENDING", C.PH("Tribe realty")],
        ["Tribal Council resolution", "Lumbee Tribal Council",
         "Authorize partnership & agreements", "PENDING", C.PH("Tribal Council")],
        ["RIVR Tech corporate consent", "RIVR Tech (member/board)",
         "Authorize execution & performance", "PENDING", C.PH("RIVR Tech officer")],
        ["LREMC board approval / joinder", "LREMC",
         "Use of LREMC Facilities (poles/fiber/easements)", "PENDING", C.PH("LREMC officer")],
        ["Landowner consents", "Landowner(s) / NCDOT",
         "Easements/ROW/encroachment for the routes", "PENDING", C.PH("project realty")],
        ["Lender consent / subordination", "Lender(s)",
         "Collateral/access acknowledgment where financed", "PENDING", C.PH("finance lead")],
        ["Regulatory (state)", "NCUC",
         "Any state authorization/registration as applicable", "PENDING", C.PH("regulatory counsel")],
        ["Regulatory (federal)", "FCC",
         "ETC/USF/Lifeline & any FCC obligations as applicable", "PENDING", C.PH("regulatory counsel")],
        ["Environmental / historic review", "NTIA / SHPO / THPO",
         "NEPA/NHPA §106 where applicable", "PENDING", C.PH("compliance lead")],
    ]
    fill_register(
        ws, "S14 — Required Approvals Register (ALL PENDING)",
        headers, rows,
        widths=[34, 26, 40, 14, 24],
        subtitle="Status column is PENDING for every row; none obtained. Placeholders for owners.")
    return C.xl_save(wb, "05_Schedules", "S14_Required_Approvals_Federal_Tribal_Corporate_Landowner_Lender.xlsx")


# ---------------------------------------------------------------------------
BUILDERS = [
    ("S1", build_s1), ("S2", build_s2), ("S3", build_s3), ("S4", build_s4),
    ("S5", build_s5), ("S6", build_s6), ("S7", build_s7), ("S8", build_s8),
    ("S9", build_s9), ("S10", build_s10), ("S11", build_s11), ("S12", build_s12),
    ("S13", build_s13), ("S14", build_s14),
]


def main():
    print("Generating 14 SCHEDULES (S1–S14) into 05_Schedules/ ...\n")
    results = []
    for sn, fn in BUILDERS:
        path = fn()
        results.append((sn, path))
        print(f"  {sn:>3}  ->  {path}")
    print(f"\nDone. {len(results)} files written.")
    return results


if __name__ == "__main__":
    main()
