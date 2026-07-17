"""
finish_workbooks.py — post-generation finisher.

Injects a parties-identification line into the Instructions sheet of every
workbook so the Tribe and RIVR Tech are named in each spreadsheet (spreadsheets
otherwise identify parties only by role). Idempotent: skips a workbook that
already names the Tribe. Run after regenerating the package.
"""
import os
import glob
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C
from openpyxl import load_workbook
from openpyxl.styles import Font

parties = f"Parties: {C.TRIBE_FULL} (the “{C.TRIBE_SHORT}”) and {C.OPERATOR_FULL} (“{C.OPERATOR_SHORT}”)"
patched = []
for path in glob.glob(os.path.join(C.PKG_DIR, "**", "*.xlsx"), recursive=True):
    wb = load_workbook(path)
    ws = wb["Instructions"] if "Instructions" in wb.sheetnames else wb.worksheets[0]
    present = any(
        v and C.TRIBE_FULL in str(v)
        for row in ws.iter_rows(values_only=True) for v in row
    )
    if present:
        continue
    cell = ws.cell(row=3, column=2, value=parties)
    cell.font = Font(italic=True, size=9, color="595959")
    wb.save(path)
    patched.append(os.path.relpath(path, C.PKG_DIR))

print(f"Patched {len(patched)} workbook(s) with parties line.")
for p in patched:
    print("  +", p)
