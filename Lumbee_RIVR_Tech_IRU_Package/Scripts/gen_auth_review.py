"""
gen_auth_review.py — Generator for Deliverable 11 (Tribal & Corporate
Authorizations, 05_Authorizations/) and Deliverable 12 (Review & Negotiation
Materials, 06_Review_Materials/) of the Lumbee Tribe of North Carolina / RIVR
Tech broadband IRU transaction package.

Builds 10 files (3 DOCX authorizations + 5 XLSX review workbooks + 2 DOCX
review documents). Every artifact is a DRAFT FOR DISCUSSION and imports the
canonical common.py module so party names, defined terms, citations, flags and
formatting stay identical across the whole package.

Run:  python3 gen_auth_review.py
"""

from __future__ import annotations

import sys
import os

sys.path.insert(0, "/home/user/claude-skills/Lumbee_RIVR_Tech_IRU_Package/Scripts")
import common as C  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402


# ---------------------------------------------------------------------------
# small local docx helpers (built on top of common.py primitives)
# ---------------------------------------------------------------------------
def whereas(doc, text):
    """A 'WHEREAS, ...' recital paragraph with an inline-styled body."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = C.Pt(6)
    r = p.add_run("WHEREAS, ")
    r.bold = True
    C._emit_runs(p, text)
    return p


def resolved(doc, text, bold_lead="NOW, THEREFORE, BE IT RESOLVED"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = C.Pt(6)
    r = p.add_run(bold_lead + " ")
    r.bold = True
    C._emit_runs(p, text)
    return p


def plain_bold(doc, text, size=11, color=None, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = C.WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = C.Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return p


def sig_line(doc, label, value, indent=0.25):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = C.Inches(indent)
    r = p.add_run(f"{label}  ")
    r.bold = True
    C._emit_runs(p, value)
    return p


LINE = "____________________________________"
GENERATED = []


# ===========================================================================
#  DELIVERABLE 11 — 05_Authorizations
# ===========================================================================
def build_tribal_resolution():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "05.01",
        "Draft Resolution of the Lumbee Tribal Council",
        "Authorizing the TBCP Round 3 Award, Tribal Ownership of Grant-Funded "
        "Infrastructure, and Execution of the RIVR Tech Partnership Agreements",
    )
    C.setup_header_footer(doc, "Lumbee Tribal Council Resolution")
    C.status_banner(doc)
    C.spacer(doc, 1)

    plain_bold(doc, "LUMBEE TRIBE OF NORTH CAROLINA", size=15, color=C.NAVY, center=True)
    plain_bold(doc, "LUMBEE TRIBAL COUNCIL", size=13, color=C.NAVY, center=True)
    plain_bold(doc, f"RESOLUTION NO. {C.PH('resolution number, e.g., 2026-__')}",
               size=12, color=C.BLACK, center=True)
    ttl = doc.add_paragraph()
    ttl.alignment = C.WD_ALIGN_PARAGRAPH.CENTER
    rt = ttl.add_run(
        "A RESOLUTION AUTHORIZING PARTICIPATION IN THE TRIBAL BROADBAND "
        "CONNECTIVITY PROGRAM (ROUND 3); TRIBAL OWNERSHIP OF GRANT-FUNDED "
        "BROADBAND INFRASTRUCTURE; AND EXECUTION OF THE MASTER DEVELOPMENT, "
        "CONSTRUCTION AND OPERATING AGREEMENT, THE INDEFEASIBLE RIGHT OF USE "
        "AGREEMENT, AND RELATED ANCILLARY AGREEMENTS WITH LREMC TECHNOLOGIES, "
        "LLC d/b/a RIVR TECH"
    )
    rt.italic = True
    rt.font.size = C.Pt(10)
    rt.font.color.rgb = C.GREY
    C.spacer(doc, 1)

    # ---- Threshold flag, surfaced prominently -----------------------------
    C.flag_para(
        doc, C.FLAG_GRANT,
        "THRESHOLD ELIGIBILITY MATTER — surfaced before adoption. The "
        f"{C.TRIBE_FULL} holds a distinct federal status under the "
        f"{C.CITES['lumbee_act']}. Whether the Lumbee Tribe qualifies as an eligible "
        "\"Tribal Government\" (or must apply through a designated/affiliated "
        "eligible entity, or seek NTIA confirmation of eligibility) for the "
        "TBCP Round 3 Award is a gating question that must be confirmed before "
        "or concurrently with the actions authorized below. This Resolution is "
        "adopted subject to that confirmation.",
    )
    C.flag_para(
        doc, C.FLAG_ATTORNEY,
        "Tribal counsel must confirm Council quorum, adoption procedure, and "
        "the Lumbee Tribe's constitutional/governing-document authority to take each "
        "action below, and must separately review any waiver of sovereign "
        "immunity (see RESOLVED clause (g), which reserves — and does NOT "
        "grant — any such waiver).",
    )
    C.spacer(doc, 1)

    # ---- Recitals ---------------------------------------------------------
    plain_bold(doc, "RECITALS", size=12, color=C.NAVY)
    whereas(doc,
        f"the {C.TRIBE_FULL} (the \"{C.TRIBE_SHORT}\") is committed to closing "
        f"the digital divide and expanding access to affordable, high-speed "
        f"broadband service ({C.SPEED_FLOOR} or better) for Tribal members and "
        f"households throughout {C.GEOGRAPHY}; and")
    whereas(doc,
        f"the {C.AGENCY_FULL} ({C.AGENCY_SHORT}) has made available the "
        f"{C.PROGRAM_FULL} ({C.PROGRAM_SHORT}) pursuant to the {C.CITES['iija']}, "
        f"and has published the {C.NOFO_NAME} (the \"{C.NOFO_SHORT}\") offering "
        f"funding for broadband infrastructure deployment on Tribal lands and "
        f"for eligible Tribal entities; and")
    whereas(doc,
        f"the {C.TRIBE_SHORT} intends to pursue and, if awarded, accept a "
        f"{C.PROGRAM_SHORT} Round 3 grant award (the \"Award\") and to serve as "
        f"the RECIPIENT of the Award and the OWNER of the broadband "
        f"infrastructure funded in whole or in part by the Award (the "
        f"\"{C.GRANT_FUNDED}\" or \"{C.TRIBAL_ASSETS}\"); and")
    whereas(doc,
        f"the {C.GRANT_FUNDED} will be subject to a continuing {C.FEDERAL_INTEREST} "
        f"and to the federal requirements of {C.CITES['ug_part']}, including "
        f"property-management, procurement, program-income, and disposition "
        f"requirements, which the {C.TRIBE_SHORT} as Recipient is responsible "
        f"for satisfying; and")
    whereas(doc,
        f"the {C.TRIBE_SHORT} has identified {C.OPERATOR_FULL} (\"{C.OPERATOR_SHORT}\") "
        f"as a qualified developer and operator to design, construct, operate, "
        f"and maintain a hybrid broadband {C.NETWORK} combining the {C.TRIBAL_ASSETS} "
        f"with {C.OPERATOR_SHORT}'s existing middle-mile and core facilities (the "
        f"\"{C.OPERATOR_EXISTING}\"); and")
    whereas(doc,
        f"the proposed commercial structure contemplates that the {C.TRIBE_SHORT} "
        f"will retain ownership of the {C.TRIBAL_ASSETS} and grant {C.OPERATOR_SHORT} "
        f"an indefeasible right of use (\"{C.IRU_TERM_DEFINED}\") for a term of "
        f"{C.PH('20 / 25 / 30 years — 30 recommended')}, together with operating, "
        f"maintenance, and revenue-sharing obligations, all as more particularly "
        f"set forth in the transaction documents; and")
    whereas(doc,
        f"the material economic terms (including {C.IRU_TERM_DEFINED} term, "
        f"revenue-share model, fixed-payment and revenue-share amounts, reserved "
        f"Tribal strand count, and pricing) remain "
        f"{C.PH('under negotiation and not yet finalized')} and will be brought "
        f"back to the Council or its authorized representatives for confirmation "
        f"consistent with this Resolution; and")
    whereas(doc,
        f"the Council has reviewed, or directed its representatives and counsel "
        f"to review, the forms of the Master Development, Construction and "
        f"Operating Agreement, the {C.IRU_TERM_DEFINED} Agreement, and the "
        f"related ancillary agreements, schedules, and exhibits (collectively, "
        f"the \"Transaction Documents\"); and")
    whereas(doc,
        f"the Council finds that authorizing the actions below, subject to the "
        f"eligibility confirmation and counsel reviews noted herein, is in the "
        f"best interest of the {C.TRIBE_SHORT} and its members;")
    C.spacer(doc, 1)

    # ---- Resolved clauses -------------------------------------------------
    resolved(doc,
        f"by the Lumbee Tribal Council that the following actions are hereby "
        f"AUTHORIZED and APPROVED, each subject to the caveats stated herein:")
    C.subsection(doc, "a",
        f"Grant Application / Award Acceptance. The {C.TRIBE_SHORT} is authorized "
        f"to prepare, submit, and pursue an application for, and to accept if "
        f"awarded, a {C.PROGRAM_SHORT} Round 3 Award, and to execute the "
        f"{C.AGENCY_SHORT} award documents, Standard Terms and Conditions, and "
        f"Specific Award Conditions, provided that the eligibility matter flagged "
        f"above is confirmed. {C.FLAG_GRANT}")
    C.subsection(doc, "b",
        f"Tribal Ownership; Federal Interest. The {C.TRIBE_SHORT} is authorized "
        f"to take title to and own the {C.TRIBAL_ASSETS} as Recipient, and to "
        f"accept and undertake the continuing {C.FEDERAL_INTEREST} and the "
        f"obligations of {C.CITES['ug_part']}, including {C.CITES['trust']} and "
        f"the property-use, encumbrance, and disposition restrictions of "
        f"{C.CITES['real_property']}. {C.FLAG_GRANT}")
    C.subsection(doc, "c",
        f"Master Development, Construction and Operating Agreement. The "
        f"{C.TRIBE_SHORT} is authorized to negotiate, execute, and deliver the "
        f"Master Development, Construction and Operating Agreement with "
        f"{C.OPERATOR_SHORT}, substantially in the form reviewed by counsel, with "
        f"such changes as the Authorized Representatives (below) approve on advice "
        f"of counsel.")
    C.subsection(doc, "d",
        f"IRU Agreement. The {C.TRIBE_SHORT} is authorized to negotiate, execute, "
        f"and deliver the Indefeasible Right of Use ({C.IRU_TERM_DEFINED}) "
        f"Agreement granting {C.OPERATOR_SHORT} an {C.IRU_TERM_DEFINED} in the "
        f"{C.TRIBAL_ASSETS} for the negotiated term "
        f"({C.PH('20 / 25 / 30 years — 30 recommended')}), consistent with the "
        f"Federal Interest and disposition restrictions.")
    C.subsection(doc, "e",
        f"Ancillary Agreements and Exhibits. The {C.TRIBE_SHORT} is authorized to "
        f"execute and deliver the ancillary agreements, schedules, and exhibits "
        f"comprising the Transaction Documents, including the operations & "
        f"maintenance / SLA schedule, the financial schedule, the grant-compliance "
        f"agreement, the privacy/data agreement, the transition agreement, and "
        f"related exhibits.")
    C.subsection(doc, "f",
        f"Authorized Representatives. The following are designated as \"Authorized "
        f"Representatives\" empowered, acting singly or as the Council directs, to "
        f"finalize, execute, and deliver the Transaction Documents and take "
        f"related actions: {C.PH('name / title — e.g., Chairman of the Lumbee Tribal Council')}; "
        f"{C.PH('name / title — e.g., Tribal Administrator')}; and "
        f"{C.PH('name / title — e.g., Chief Financial Officer or designee')}.")
    C.subsection(doc, "g",
        f"Limited Approvals Reserved to Separate, Specific Action (Sovereign "
        f"Immunity NOT Waived Herein). Any limited waiver of the {C.TRIBE_SHORT}'s "
        f"sovereign immunity, any consent to a particular forum, governing law, or "
        f"dispute-resolution mechanism, and any grant of remedies affecting Tribal "
        f"assets are EXPRESSLY RESERVED and are NOT granted by this Resolution. Such "
        f"matters, if approved at all, shall be the subject of a separate, specific "
        f"Council action, narrowly tailored, on the written advice of Tribal "
        f"counsel. {C.FLAG_ATTORNEY}")
    C.subsection(doc, "h",
        f"Grant-Compliance Responsibilities. The Authorized Representatives are "
        f"directed to establish and maintain the systems necessary for the "
        f"{C.TRIBE_SHORT} to meet its Recipient obligations, including procurement "
        f"standards ({C.CITES['procurement']}), program-income handling "
        f"({C.CITES['prog_income']}), subrecipient-vs-contractor determinations "
        f"({C.CITES['subrecipient']}), record retention ({C.CITES['records']}), "
        f"single-audit ({C.CITES['single_audit']}), and environmental/historic "
        f"review ({C.CITES['nepa']}; {C.CITES['nhpa']}). {C.FLAG_GRANT}")
    C.spacer(doc, 1)
    resolved(doc,
        f"that this Resolution shall take effect immediately upon its adoption, "
        f"and that all prior actions of the {C.TRIBE_SHORT}'s officers and "
        f"representatives consistent herewith are ratified and confirmed.",
        bold_lead="BE IT FURTHER RESOLVED")
    C.spacer(doc, 2)

    # ---- Certification block ----------------------------------------------
    plain_bold(doc, "CERTIFICATION", size=12, color=C.NAVY)
    p = doc.add_paragraph()
    C._emit_runs(p,
        f"The undersigned, serving as Secretary of the Lumbee Tribal Council, "
        f"hereby certifies that the foregoing Resolution No. "
        f"{C.PH('resolution number')} was duly introduced, considered, and ADOPTED "
        f"by the Lumbee Tribal Council at a duly called meeting held on "
        f"{C.PH('meeting date')}, at which a quorum was present and acting "
        f"throughout, by the following vote:")
    sig_line(doc, "For:", C.PH("number of votes in favor"))
    sig_line(doc, "Against:", C.PH("number of votes against"))
    sig_line(doc, "Abstaining:", C.PH("number abstaining"))
    sig_line(doc, "Absent / Recused:", C.PH("number absent or recused"))
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    C._emit_runs(p, "and that said Resolution has not been amended, rescinded, "
                    "or modified and remains in full force and effect as of the "
                    "date below.")
    C.spacer(doc, 2)
    sig_line(doc, "By:", LINE, indent=0.0)
    sig_line(doc, "Name:", C.PH("Secretary of the Lumbee Tribal Council"), indent=0.0)
    sig_line(doc, "Title:", "Secretary, Lumbee Tribal Council", indent=0.0)
    sig_line(doc, "Date:", LINE, indent=0.0)
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = C.Inches(0.0)
    r = p.add_run("[ SEAL of the Lumbee Tribe of North Carolina ]")
    r.italic = True
    r.font.color.rgb = C.GREY

    path = C.save(doc, "05_Authorizations", "01_Draft_Lumbee_Tribal_Council_Resolution.docx")
    GENERATED.append(path)


def build_rivr_corporate_authorization():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "05.02",
        "Draft Written Consent of the Managers and Members",
        "LREMC Technologies, LLC d/b/a RIVR Tech — Authorization to Execute the "
        "Lumbee Tribe Broadband Partnership Agreements",
    )
    C.setup_header_footer(doc, "RIVR Tech Corporate Authorization")
    C.status_banner(doc)
    C.spacer(doc, 1)

    plain_bold(doc, C.OPERATOR_FULL, size=14, color=C.NAVY, center=True)
    plain_bold(doc, "ACTION BY UNANIMOUS WRITTEN CONSENT", size=12, color=C.NAVY, center=True)
    plain_bold(doc, "OF THE MANAGER(S) AND MEMBER(S) IN LIEU OF A MEETING",
               size=11, color=C.GREY, center=True)
    C.spacer(doc, 1)

    C.flag_para(
        doc, C.FLAG_ATTORNEY,
        "Company counsel to confirm the correct governance body (manager-managed "
        "vs. member-managed), the requisite consent threshold under the LLC's "
        "operating agreement and the North Carolina Limited Liability Company Act, "
        "and any parent/affiliate consents required (LREMC).",
    )
    C.spacer(doc, 1)

    p = doc.add_paragraph()
    C._emit_runs(p,
        f"The undersigned, constituting the {C.PH('manager(s) / all members / requisite percentage in interest')} "
        f"of {C.OPERATOR_FULL}, a {C.PH('North Carolina')} limited liability "
        f"company (the \"Company\" or \"{C.OPERATOR_SHORT}\"), acting by written "
        f"consent in lieu of a meeting pursuant to the Company's operating "
        f"agreement and applicable law, hereby adopt the following resolutions "
        f"as of {C.PH('consent date')}:")
    C.spacer(doc, 1)

    plain_bold(doc, "RECITALS", size=12, color=C.NAVY)
    whereas(doc,
        f"the Company desires to partner with the {C.TRIBE_FULL} (the "
        f"\"{C.TRIBE_SHORT}\") to develop, construct, operate, and maintain a "
        f"hybrid broadband {C.NETWORK} in {C.GEOGRAPHY}, in which the "
        f"{C.TRIBE_SHORT} will own the {C.PROGRAM_SHORT}-funded infrastructure and "
        f"grant the Company an indefeasible right of use ({C.IRU_TERM_DEFINED}); and")
    whereas(doc,
        f"the Company has reviewed the forms of the Master Development, "
        f"Construction and Operating Agreement, the {C.IRU_TERM_DEFINED} Agreement, "
        f"and the related ancillary agreements, schedules, and exhibits "
        f"(collectively, the \"Transaction Documents\"); and")
    whereas(doc,
        f"the managers/members have determined that entering into the Transaction "
        f"Documents is in the best interest of the Company;")
    C.spacer(doc, 1)

    plain_bold(doc, "RESOLUTIONS", size=12, color=C.NAVY)
    C.subsection(doc, "1",
        f"Approval of Transaction. The Company's entry into the Transaction "
        f"Documents with the {C.TRIBE_SHORT}, substantially in the forms reviewed, "
        f"with such changes as the Authorized Signatory approves on advice of "
        f"counsel, is hereby APPROVED.")
    C.subsection(doc, "2",
        f"Authorized Signatory. {C.PH('name')}, serving as "
        f"{C.PH('title — e.g., Chief Executive Officer / Manager / President')} of "
        f"the Company, is designated the \"Authorized Signatory\" and is authorized "
        f"to negotiate, execute, and deliver the Transaction Documents and any "
        f"related certificates, notices, and instruments on behalf of the Company.")
    C.subsection(doc, "3",
        f"Authority and Good Standing. The Company confirms that it is a limited "
        f"liability company duly organized, validly existing, and in good standing "
        f"under the laws of {C.PH('State of formation — North Carolina')}, is "
        f"qualified to do business in {C.STATE}, and has full power and authority "
        f"to enter into and perform the Transaction Documents.")
    C.subsection(doc, "4",
        f"Insurance Commitments. The Company is authorized and directed to obtain "
        f"and maintain the insurance coverages required by the Transaction "
        f"Documents (Exhibit I), including commercial general liability "
        f"({C.DEAL['insurance']['cgl_occurrence']} per occurrence / "
        f"{C.DEAL['insurance']['cgl_aggregate']} aggregate), automobile liability "
        f"({C.DEAL['insurance']['auto']}), umbrella/excess "
        f"({C.DEAL['insurance']['umbrella']}), workers' compensation "
        f"({C.DEAL['insurance']['workers_comp']}) and employers' liability "
        f"({C.DEAL['insurance']['employers_liability']}), professional/technology "
        f"E&O ({C.DEAL['insurance']['professional_tech_eo']}), cyber "
        f"({C.DEAL['insurance']['cyber']}), and builders' risk/property "
        f"({C.DEAL['insurance']['property_builders_risk']}), naming the "
        f"{C.TRIBE_SHORT} as additional insured where required.")
    C.subsection(doc, "5",
        f"Bonding Commitments. The Company is authorized and directed to procure "
        f"the payment and performance bonds required for the construction work "
        f"({C.PH('bond amounts / percentage of construction cost — confirm with surety')}), "
        f"with the {C.TRIBE_SHORT} as obligee/dual-obligee as required.")
    C.subsection(doc, "6",
        f"Grant Flow-Down. The Company acknowledges that, as a "
        f"contractor/subrecipient to a federal Recipient, it will comply with the "
        f"applicable flow-down provisions of {C.CITES['ug_part']}, including "
        f"{C.CITES['telecom_ban']} (§889 covered equipment), {C.CITES['baba']} "
        f"(Build America, Buy America), and {C.CITES['debarment']} (debarment/"
        f"suspension). {C.FLAG_GRANT}")
    C.subsection(doc, "7",
        f"Further Assurances. The Authorized Signatory is authorized to take all "
        f"further actions and execute all further documents reasonably necessary "
        f"to carry out the intent of these resolutions.")
    C.spacer(doc, 2)

    p = doc.add_paragraph()
    r = p.add_run("IN WITNESS WHEREOF, the undersigned have executed this written "
                  "consent as of the date first written above.")
    r.italic = True
    C.spacer(doc, 1)
    plain_bold(doc, C.OPERATOR_FULL, size=11, color=C.NAVY)
    for _ in range(2):
        sig_line(doc, "By:", LINE, indent=0.0)
        sig_line(doc, "Name:", C.PH("manager / member name"), indent=0.0)
        sig_line(doc, "Title:", C.PH("Manager / Member / Managing Member"), indent=0.0)
        sig_line(doc, "Date:", LINE, indent=0.0)
        C.spacer(doc, 1)

    path = C.save(doc, "05_Authorizations", "02_Draft_RIVR_Tech_Corporate_Authorization.docx")
    GENERATED.append(path)


def _cert_section(doc, party_full, party_short, cert_officer_ph, signatories):
    plain_bold(doc, f"CERTIFICATE OF AUTHORITY AND INCUMBENCY — {party_short.upper()}",
               size=12, color=C.NAVY)
    p = doc.add_paragraph()
    C._emit_runs(p,
        f"The undersigned, {cert_officer_ph}, hereby certifies on behalf of "
        f"{party_full} (the \"Entity\") that:")
    C.subsection(doc, "1",
        f"The undersigned is the duly elected/appointed and acting "
        f"{C.PH('certifying officer title — e.g., Secretary / Tribal Secretary')} "
        f"of the Entity and is authorized to execute this Certificate.")
    C.subsection(doc, "2",
        f"Each person named below has been duly authorized to execute and deliver "
        f"the Master Development, Construction and Operating Agreement, the "
        f"{C.IRU_TERM_DEFINED} Agreement, and the related ancillary agreements, "
        f"schedules, and exhibits (the \"Transaction Documents\") on behalf of the "
        f"Entity, holds the office or title set forth opposite such person's name, "
        f"and the specimen signature set forth is the genuine signature of such "
        f"person:")
    C.spacer(doc, 1)
    C.add_table(
        doc,
        ["Name", "Title / Office", "Specimen Signature"],
        [[C.PH("authorized signatory name"), C.PH("title"), "____________________"]
         for _ in range(signatories)],
        widths=[2.1, 2.3, 2.1],
        col_align=["l", "l", "c"],
    )
    C.subsection(doc, "3",
        f"The resolutions/consents authorizing the Transaction Documents "
        f"(namely, {C.PH('reference to Tribal Council Resolution No. / RIVR Tech written consent dated ___')}) "
        f"were duly adopted, have not been amended, rescinded, or revoked, and "
        f"remain in full force and effect as of the date hereof.")
    C.subsection(doc, "4",
        f"The Entity is validly existing and in good standing under its governing "
        f"law/charter, and the execution and delivery of the Transaction Documents "
        f"do not violate its governing documents. {C.FLAG_ATTORNEY}")
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    C._emit_runs(p, f"IN WITNESS WHEREOF, the undersigned has executed this "
                    f"Certificate as of {C.PH('date')}.")
    C.spacer(doc, 1)
    sig_line(doc, "By:", LINE, indent=0.0)
    sig_line(doc, "Name:", C.PH("certifying officer name"), indent=0.0)
    sig_line(doc, "Title:", C.PH("certifying officer title"), indent=0.0)
    sig_line(doc, "Date:", LINE, indent=0.0)
    # cross-certification of the certifier's own incumbency
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    r = p.add_run("Countersignature (certifying the incumbency of the officer above):")
    r.italic = True
    r.font.color.rgb = C.GREY
    sig_line(doc, "By:", LINE, indent=0.0)
    sig_line(doc, "Name / Title:", C.PH("second officer certifying the certifier"), indent=0.0)


def build_certificate_of_authority():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "05.03",
        "Certificate of Authority and Incumbency",
        "Combined Form for Both Parties — Lumbee Tribe of North Carolina and "
        "LREMC Technologies, LLC d/b/a RIVR Tech",
    )
    C.setup_header_footer(doc, "Certificate of Authority")
    C.status_banner(doc)
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    r = p.add_run(
        "This combined instrument contains two certificates — one for each Party. "
        "The Parties may execute separate certificates or this combined form.")
    r.italic = True
    r.font.color.rgb = C.GREY
    C.spacer(doc, 1)

    _cert_section(
        doc, C.TRIBE_FULL, C.TRIBE_SHORT,
        f"acting as {C.PH('Secretary of the Lumbee Tribal Council / Tribal Secretary')}",
        signatories=3,
    )
    doc.add_page_break()
    C.status_banner(doc)
    C.spacer(doc, 1)
    _cert_section(
        doc, C.OPERATOR_FULL, C.OPERATOR_SHORT,
        f"acting as {C.PH('Secretary / Manager / authorized officer of the Company')}",
        signatories=2,
    )

    path = C.save(doc, "05_Authorizations", "03_Certificate_of_Authority.docx")
    GENERATED.append(path)


# ===========================================================================
#  DELIVERABLE 12 — 06_Review_Materials
# ===========================================================================
def build_attorney_issue_list():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "06.01",
        "Attorney Review Issue List",
        "Open Legal Issues, Recommended Positions, and Owners for the "
        "Lumbee Tribe / RIVR Tech Broadband Partnership",
    )
    C.setup_header_footer(doc, "Attorney Review Issue List")
    C.status_banner(doc)
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    C._emit_runs(p,
        "This issue list is prepared for counsel review. It is not legal advice. "
        "Recommended positions are drafting starting points only; governing law, "
        "jurisdiction, sovereign immunity, and subrecipient-vs-contractor status "
        "are OPEN ITEMS reserved for the Parties' counsel and are not resolved "
        "here. Priority: H = high / gating, M = medium, L = lower.")
    C.spacer(doc, 1)
    C.flag_para(
        doc, C.FLAG_GRANT,
        "Lead item #1 (Lumbee federal recognition / TBCP eligibility) is a "
        "threshold, potentially deal-gating question and is flagged for grant and "
        "attorney review before other items are finalized.",
    )
    C.spacer(doc, 1)

    Cx = C.CITES
    # Issue | Description/Risk | Citation | Recommended Position | Owner | Priority
    issues = [
        ("1. Lumbee federal recognition / TBCP eligibility (LEAD)",
         "Whether the Lumbee Tribe qualifies as an eligible \"Tribal Government\" for a "
         "TBCP Award, given the Lumbee Act's limited-recognition language. If "
         "ineligible or contested, the entire funding premise fails or must be "
         "restructured through an eligible designated entity. " + C.FLAG_GRANT,
         Cx["lumbee_act"] + "; " + Cx["nofo"],
         "Confirm eligibility with NTIA in writing before award acceptance; "
         "identify an eligible designated/affiliated entity as fallback; condition "
         "all agreements on eligibility confirmation.",
         "Tribal counsel + Grant counsel + NTIA", "H"),
        ("2. Sovereign immunity",
         "Extent to which the Lumbee Tribe waives immunity for enforcement; RIVR Tech "
         "needs enforceable remedies; Lumbee Tribe must protect assets and self-governance.",
         "Common-law tribal sovereign immunity; " + C.FLAG_ATTORNEY,
         "OPEN — do not pre-decide. If any waiver, make it a limited, express, "
         "narrowly-scoped waiver by separate specific Council action, capped and "
         "asset-limited; pair with dispute-resolution clause.",
         "Tribal counsel (lead) + RIVR counsel", "H"),
        ("3. Governing law",
         "Which body of law governs the agreements (Tribal law, N.C. law, federal). "
         "Affects interpretation and enforceability.",
         "Choice-of-law principles; " + C.FLAG_ATTORNEY,
         "OPEN — reserved for counsel. Commonly negotiated to N.C. law for "
         "commercial terms with Tribal law respected for Tribal-governance matters; "
         "confirm consistency with any immunity waiver.",
         "Both parties' counsel", "H"),
        ("4. Tribal court jurisdiction",
         "Whether disputes may/must be heard in a Tribal forum; adequacy and "
         "neutrality of forum for a commercial counterparty.",
         C.FLAG_ATTORNEY,
         "OPEN. Consider Tribal court for Tribal-interest matters; confirm rules "
         "and appeal rights; coordinate with exhaustion doctrine (#6).",
         "Tribal counsel + RIVR counsel", "H"),
        ("5. Federal court jurisdiction",
         "Availability of a federal forum; diversity/federal-question basis is "
         "uncertain for tribal commercial disputes.",
         C.FLAG_ATTORNEY,
         "OPEN. Do not assume federal jurisdiction exists; if desired, draft an "
         "express consent tied to a limited immunity waiver.",
         "Both parties' counsel", "M"),
        ("6. Arbitration",
         "Binding arbitration as a neutral mechanism; seat, rules, and "
         "enforceability against a sovereign.",
         "FAA / applicable arbitration law; " + C.FLAG_ATTORNEY,
         "Recommend binding arbitration (e.g., AAA commercial) with a defined seat "
         "and limited immunity waiver for the purpose of compelling arbitration and "
         "enforcing awards; carve out injunctive/step-in relief.",
         "Both parties' counsel", "M"),
        ("7. Exhaustion of Tribal remedies",
         "Doctrine may require exhaustion in Tribal court before other fora, "
         "affecting speed/cost of enforcement.",
         C.FLAG_ATTORNEY,
         "OPEN. Address expressly — either require exhaustion or waive/limit it in "
         "the dispute clause to avoid uncertainty.",
         "Tribal counsel + RIVR counsel", "M"),
        ("8. Federal grant restrictions (general)",
         "Award terms constrain use, transfer, and operation of grant assets; "
         "IRU/operating structure must fit within them.",
         Cx["ug_part"] + "; " + Cx["sac"] + "; " + C.FLAG_GRANT,
         "Map every commercial term against award terms; add a grant-compliance "
         "supremacy clause so federal terms control on conflict.",
         "Grant counsel", "H"),
        ("9. Program income (2 CFR 200.307)",
         "Revenue from operating grant-funded assets may be Program Income subject "
         "to federal use rules (deduction/addition/cost-sharing methods).",
         Cx["prog_income"] + "; " + C.FLAG_GRANT,
         "Determine method with NTIA; structure revenue-share so Lumbee Tribe can meet "
         "Program Income obligations; document treatment in the financial schedule.",
         "Grant counsel + Finance", "H"),
        ("10. Contractor vs subrecipient status (2 CFR 200.331)",
         "Whether RIVR Tech is a contractor or a subrecipient — drives flow-down, "
         "audit, and compliance burden.",
         Cx["subrecipient"] + "; " + Cx["pass_through"] + "; " + C.FLAG_GRANT,
         "OPEN — perform the 200.331 checklist analysis; do not assume. Likely "
         "contractor for operations, but confirm; set flow-downs accordingly.",
         "Grant counsel", "H"),
        ("11. Procurement compliance (2 CFR 200.317–.327)",
         "Selection of RIVR Tech and downstream purchases must satisfy federal "
         "procurement standards (competition, methods, domestic preference).",
         Cx["procurement"] + "; " + Cx["competition"] + "; " + C.FLAG_GRANT,
         "Document the procurement basis (competition or sole-source justification) "
         "for RIVR Tech selection; require compliant subprocurement.",
         "Grant counsel + Lumbee Tribe procurement", "H"),
        ("12. Impermissible private benefit from grant assets",
         "Risk that a private operator derives disproportionate benefit from "
         "publicly funded assets, jeopardizing the award.",
         Cx["ug_part"] + "; " + C.FLAG_GRANT,
         "Ensure consideration to the Lumbee Tribe is fair and documented; benchmark IRU "
         "and revenue-share terms; retain public-purpose protections and reversion.",
         "Grant counsel + Finance", "H"),
        ("13. Asset encumbrance",
         "Federal interest restricts liens/encumbrances on grant assets; lenders "
         "or the IRU itself may be viewed as encumbrances.",
         Cx["real_property"] + "; " + Cx["trust"] + "; " + C.FLAG_GRANT,
         "Obtain NTIA position on whether the IRU is a permissible use vs. "
         "encumbrance; avoid liens on Tribal Assets; use RIVR assets for security.",
         "Grant counsel + Tribal counsel", "H"),
        ("14. Federal interest (2 CFR 200.316)",
         "Property trust relationship — federal interest persists through the "
         "useful life / disposition.",
         Cx["trust"] + "; " + C.FLAG_GRANT,
         "Acknowledge federal interest in all conveyancing docs; align IRU term and "
         "disposition with federal-interest period.",
         "Grant counsel", "M"),
        ("15. BIA or other federal approval questions",
         "Whether BIA/land-status or other federal approvals are triggered by the "
         "transaction or the underlying land.",
         "25 U.S.C./25 CFR (BIA) as applicable; " + C.FLAG_ATTORNEY,
         "OPEN — confirm land status and whether any BIA approval or right-of-way "
         "consent is required; may be N/A if no trust land involved.",
         "Tribal counsel", "M"),
        ("16. Easements",
         "Rights to place/maintain fiber across public and private parcels; "
         "durability across the IRU term.",
         Cx["nc_pole"] + "; " + C.FLAG_ATTORNEY,
         "Secure recordable easements of sufficient duration; confirm assignability "
         "and survival on default/transition.",
         "Real-estate counsel + RIVR", "M"),
        ("17. Rights-of-way",
         "DOT/municipal/utility ROW and pole-attachment access; permitting risk.",
         Cx["nc_pole"] + "; " + Cx["nc_dig"] + "; " + C.FLAG_ATTORNEY,
         "Confirm ROW/permitting path and pole-attachment agreements; allocate "
         "permitting risk and cost in the construction schedule.",
         "Real-estate counsel + RIVR", "M"),
        ("18. Tax treatment",
         "Property, sales/use, and income tax treatment of Tribal-owned assets and "
         "IRU consideration; possible exemptions.",
         C.FLAG_ATTORNEY,
         "OPEN — obtain tax analysis on Tribal ownership exemptions, IRU "
         "characterization, and any UBIT-type exposure for Tribal enterprises.",
         "Tax counsel", "M"),
        ("19. Regulatory authority (FCC / USAC / ETC)",
         "Which entity holds regulatory authorizations, ETC designation, and "
         "USAC/USF obligations; CPNI compliance.",
         Cx["usac_lifeline"] + "; " + Cx["cpni"] + "; " + C.FLAG_ATTORNEY,
         "Clarify that RIVR Tech (operator) holds the carrier authorizations and "
         "ETC status; assign CPNI compliance; document in the privacy agreement.",
         "Regulatory counsel + RIVR", "M"),
        ("20. Customer ownership",
         "Who owns the customer relationship, contracts, and data on default or "
         "transition — critical for continuity and value.",
         Cx["cpni"] + "; " + C.FLAG_BUSINESS,
         "Recommend Lumbee Tribe holds a reversionary/step-in right to customer "
         "relationships tied to the Tribal Assets; RIVR operates during the term.",
         "Both parties + business", "M"),
        ("21. Rate-setting",
         "Who sets retail rates and affordability commitments (100/20 affordable "
         "offering); limits on price increases.",
         Cx["nofo"] + "; " + C.FLAG_BUSINESS,
         "Set affordability floor per award; give Lumbee Tribe consultation/consent on "
         "material rate changes; document in financial + O&M schedules.",
         "Business + Grant counsel", "M"),
        ("22. Assignment",
         "Ability of either party to assign; protecting the Lumbee Tribe from unwanted "
         "successors to the operator role.",
         C.FLAG_ATTORNEY,
         "Prohibit assignment without Tribal consent (not unreasonably withheld); "
         "permit collateral assignment to lenders with recognition/cure rights.",
         "Both parties' counsel", "M"),
        ("23. Change in control",
         "Change of control of RIVR Tech/parent as a consent or termination "
         "trigger.",
         C.FLAG_ATTORNEY,
         "Treat change of control as a consent event with Tribal approval and "
         "vetting; define control threshold.",
         "Both parties' counsel", "M"),
        ("24. Default",
         "Definitions, cure periods, and cross-default across the agreement suite.",
         "See DEAL cure periods (30/60 days); " + C.FLAG_ATTORNEY,
         "Use consistent monetary (30-day) / non-monetary (60-day + extension) cure "
         "periods; cross-default within the transaction suite only.",
         "Both parties' counsel", "M"),
        ("25. Step-in rights",
         "Lumbee Tribe's ability to step in / appoint a replacement operator to preserve "
         "service and award compliance.",
         Cx["remedies"] + "; " + C.FLAG_ATTORNEY,
         "Grant Lumbee Tribe emergency step-in for essential-service/award-compliance "
         "threats; define scope, duration, and cost allocation.",
         "Tribal counsel", "H"),
        ("26. Grant clawback exposure",
         "Risk NTIA recovers funds for noncompliance, disposition, or "
         "underperformance — who bears the loss.",
         Cx["remedies"] + "; " + Cx["real_property"] + "; " + C.FLAG_GRANT,
         "Allocate clawback risk by fault; require RIVR indemnity for "
         "operator-caused noncompliance; maintain compliance reserves.",
         "Grant counsel + Finance", "H"),
        ("27. Transition obligations",
         "End-of-term / termination handover of assets, records, customers, and "
         "operations without service disruption.",
         Cx["closeout"] + "; " + C.FLAG_ATTORNEY,
         "Require a 12-month transition-assistance period, records/customer "
         "handover, and continued service at defined pricing; tie to transition "
         "agreement.",
         "Both parties' counsel", "M"),
    ]

    C.add_table(
        doc,
        ["Issue", "Description / Risk", "Citation", "Recommended Position", "Owner", "Pri."],
        [list(r) for r in issues],
        widths=[1.15, 1.9, 1.15, 1.85, 0.95, 0.35],
        font_size=8,
        col_align=["l", "l", "l", "l", "l", "c"],
    )

    path = C.save(doc, "06_Review_Materials", "01_Attorney_Review_Issue_List.docx")
    GENERATED.append(path)


# ------- XLSX helpers -------------------------------------------------------
def _set_widths(ws, widths, start=1):
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(start + i)].width = w


def _write_row(ws, row, values, kinds=None, wraps=None, fmts=None):
    for i, v in enumerate(values):
        kind = (kinds[i] if kinds else "text")
        wrap = (wraps[i] if wraps else False)
        fmt = (fmts[i] if fmts else None)
        C.xl_cell(ws, row, i + 1, v, kind=kind, wrap=wrap, fmt=fmt)


def build_open_business_decisions():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Open Business Decisions Log",
        "This log tracks the material, unresolved BUSINESS decisions for the "
        "Lumbee Tribe / RIVR Tech broadband partnership. Economics are not "
        "finalized; recommended positions are starting points for negotiation.",
        "• Filter/sort on the header row. Status values: Open, In Discussion, "
        "Tentative, Decided, Deferred.",
        "• 'Recommended' reflects the drafting team's default; it is not a "
        "decision. Legal/grant items appear on the Attorney Review Issue List.",
        "• Placeholder economics are marked and must be confirmed with the "
        "financial model (04_Financial_Model).",
    ])
    ws = wb.create_sheet("Business Decisions")
    ws.sheet_properties.tabColor = "1F3864"
    C.xl_title(ws, "Open Business Decisions Log",
               subtitle="Lumbee Tribe of North Carolina / RIVR Tech — material open commercial items",
               span=8)
    headers = ["ID", "Decision", "Options / Alternatives", "Recommended",
               "Owner", "Status", "Impacted Documents", "Due"]
    hr = 5
    C.xl_header_row(ws, hr, headers)
    _set_widths(ws, [7, 30, 40, 30, 18, 14, 32, 14])

    rows = [
        ("BD-01", "IRU term length",
         "20 / 25 / 30 years (+ renewals: two 10-yr terms)",
         "30 years (align with federal-interest period)",
         "Lumbee Tribe + RIVR", "Open",
         "Master (02.01); IRU (02.02); Financial (02.04)", C.PH("date")),
        ("BD-02", "Revenue-share model",
         "A Fixed / B % Gross Rev / C % Cash Flow / D Hybrid / E Payment holiday+stepped",
         "Option D — Hybrid (fixed base + revenue share)",
         "Lumbee Tribe + RIVR + Finance", "In Discussion",
         "Financial (02.04); Master (02.01)", C.PH("date")),
        ("BD-03", "Grant award amount",
         "Depends on NTIA award; model uses placeholder",
         C.PH("$25,000,000 placeholder — confirm on award"),
         "Grant + Finance", "Open",
         "Grant (02.05); Financial (02.04)", C.PH("award date")),
        ("BD-04", "Homes passed / footprint",
         "Depends on final network design",
         C.PH("8,500 placeholder — confirm from design"),
         "RIVR + Lumbee Tribe", "Open",
         "Financial (02.04); Exhibits (maps)", C.PH("date")),
        ("BD-05", "Retail pricing / affordability",
         "Affordable 100/20 tier + market tiers; price-increase caps",
         "Affordable tier per award floor; consult Lumbee Tribe on increases",
         "Business + Grant", "Open",
         "Financial (02.04); O&M/SLA (02.03)", C.PH("date")),
        ("BD-06", "Reserved Tribal strand count",
         "Number of dark strands reserved to the Lumbee Tribe on the Tribal Assets",
         C.PH("reserved strand count — e.g., 24–48 strands; confirm"),
         "Lumbee Tribe + RIVR", "Open",
         "IRU (02.02); Exhibits (fiber)", C.PH("date")),
        ("BD-07", "Fixed-payment amount (Option A / D base)",
         "Annual fixed IRU/operating payment to the Lumbee Tribe",
         C.PH("fixed annual payment $ — confirm from model"),
         "Finance", "Open",
         "Financial (02.04)", C.PH("date")),
        ("BD-08", "Revenue-share percentage",
         "% of gross revenue (Opt B) or adjusted cash flow (Opt C/D)",
         C.PH("5% gross / 20% cash-flow placeholders — confirm"),
         "Finance", "Open",
         "Financial (02.04)", C.PH("date")),
        ("BD-09", "Governing law",
         "Tribal law / N.C. law / federal (see Attorney Issue #3)",
         C.PH("OPEN — reserved for counsel; not pre-decided"),
         "Counsel (both)", "Open",
         "All agreements", C.PH("date")),
        ("BD-10", "Sovereign-immunity approach",
         "No waiver / limited express waiver / arbitration-only waiver",
         C.PH("OPEN — separate specific Council action; not pre-decided"),
         "Tribal counsel", "Open",
         "Master (02.01); IRU (02.02)", C.PH("date")),
        ("BD-11", "Wholesale / open-access allowed?",
         "Permit RIVR to sell wholesale/open access on the Tribal Assets? Terms?",
         C.PH("business decision — recommend allow with Lumbee Tribe revenue participation"),
         "Business + Grant", "Open",
         "Master (02.01); Financial (02.04)", C.PH("date")),
        ("BD-12", "Capital-replacement reserve %",
         "% of revenue funded to a reserve for lifecycle asset replacement",
         C.PH("reserve % — e.g., 3–5% of revenue; confirm"),
         "Finance + RIVR", "Open",
         "Financial (02.04); O&M/SLA (02.03)", C.PH("date")),
        ("BD-13", "Transition pricing",
         "Pricing for continued service/handover during 12-mo transition",
         C.PH("transition pricing basis — cost-plus vs. then-current; confirm"),
         "Business + Counsel", "Open",
         "Transition (02.07)", C.PH("date")),
        ("BD-14", "Insurance / bonding levels",
         "Confirm limits (Exhibit I) and construction bond amounts",
         "Use canonical limits (Exhibit I); confirm bond % with surety",
         "Risk + RIVR", "Tentative",
         "Master (02.01); Exhibit I", C.PH("date")),
        ("BD-15", "Customer ownership on default/transition",
         "Lumbee Tribe reversion/step-in to customer relationships vs. RIVR retains",
         "Recommend Lumbee Tribe reversionary right tied to Tribal Assets",
         "Business + Counsel", "Open",
         "Master (02.01); Transition (02.07); Privacy (02.06)", C.PH("date")),
    ]
    r = hr + 1
    for row in rows:
        _write_row(ws, r, list(row),
                   wraps=[False, True, True, True, True, False, True, False])
        r += 1

    # status dropdown
    dv = DataValidation(type="list",
                        formula1='"Open,In Discussion,Tentative,Decided,Deferred"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F{hr+1}:F{r-1}")

    path = C.xl_save(wb, "06_Review_Materials", "02_Open_Business_Decisions_Log.xlsx")
    GENERATED.append(path)


def build_negotiation_risk_matrix():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Negotiation Risk Matrix",
        "Scores each material risk in the Lumbee Tribe / RIVR Tech transaction. "
        "Score = Likelihood x Impact (each 1–5, so 1–25).",
        "• Enter Likelihood and Impact (amber input cells, 1–5). Score is a "
        "formula (blue) and auto-highlights: red when >= 15 (high), amber 8–14 "
        "(medium), green <= 7 (low).",
        "• Likelihood/Impact scale: 1 = very low, 3 = moderate, 5 = very high.",
        "• Mitigation and Owner columns capture the plan and accountable party.",
    ])
    ws = wb.create_sheet("Risk Matrix")
    ws.sheet_properties.tabColor = "C00000"
    C.xl_title(ws, "Negotiation Risk Matrix",
               subtitle="Likelihood x Impact scoring — draft for counsel and business review",
               span=7)
    headers = ["Risk", "Category", "Likelihood (1-5)", "Impact (1-5)",
               "Score (=LxI)", "Mitigation", "Owner"]
    hr = 5
    C.xl_header_row(ws, hr, headers)
    _set_widths(ws, [34, 18, 12, 12, 12, 44, 20])

    # (risk, category, L, I, mitigation, owner)
    risks = [
        ("Lumbee federal recognition / TBCP eligibility", "Grant / Threshold", 3, 5,
         "Confirm eligibility with NTIA in writing before acceptance; fallback "
         "eligible designated entity; condition agreements on confirmation.",
         "Tribal + Grant counsel"),
        ("Grant clawback / recapture", "Grant", 3, 5,
         "Fault-based allocation; RIVR indemnity for operator-caused noncompliance; "
         "compliance reserve; strict property/disposition controls.",
         "Grant counsel + Finance"),
        ("Construction cost overrun", "Construction / Financial", 4, 4,
         "GMP/fixed-price where possible; contingency; change-order controls; "
         "performance/payment bonds.",
         "RIVR + Finance"),
        ("Take-rate / adoption shortfall", "Commercial", 4, 4,
         "Conservative model; affordability tier + marketing plan; revenue-share "
         "sized to downside; milestone-based obligations.",
         "Business + Finance"),
        ("Sovereign-immunity impasse", "Legal", 3, 4,
         "Early principled negotiation; limited express waiver for arbitration only; "
         "asset-limited caps; separate Council action.",
         "Both parties' counsel"),
        ("Environmental / historic-review delay", "Regulatory", 3, 3,
         "Start NEPA/NHPA §106 early; route around sensitive sites; schedule "
         "float; SHPO/THPO consultation.",
         "Grant counsel + RIVR"),
        ("BABA sourcing non-compliance / delay", "Grant / Supply", 3, 4,
         "Early BABA-compliant sourcing; supplier certifications; waiver process "
         "if needed; document domestic content.",
         "RIVR + Grant counsel"),
        ("Program-income treatment dispute", "Grant / Financial", 3, 3,
         "Confirm method with NTIA up front; structure revenue-share to satisfy "
         "200.307; document in financial schedule.",
         "Grant counsel + Finance"),
        ("Impermissible private benefit", "Grant", 2, 5,
         "Fair, benchmarked, documented consideration to Lumbee Tribe; public-purpose "
         "protections; reversion; arm's-length terms.",
         "Grant counsel + Finance"),
        ("Operator default / step-in event", "Operational / Legal", 2, 5,
         "Clear default/cure; Lumbee Tribe emergency step-in; replacement-operator plan; "
         "escrow of records/credentials.",
         "Tribal counsel + RIVR"),
        ("Asset encumbrance conflict with federal interest", "Grant / Legal", 3, 4,
         "No liens on Tribal Assets; confirm IRU is permissible use; security from "
         "RIVR assets only; NTIA position letter.",
         "Grant + Tribal counsel"),
        ("Easement / ROW / pole-attachment failure", "Real Estate", 3, 3,
         "Secure recordable long-duration easements/ROW; pole agreements; allocate "
         "permitting risk; survival on transition.",
         "Real-estate counsel + RIVR"),
        ("Governing-law / jurisdiction dispute", "Legal", 2, 4,
         "Resolve choice-of-law + forum + exhaustion together; consistent across "
         "the agreement suite.",
         "Both parties' counsel"),
        ("Regulatory / ETC / USAC compliance gap", "Regulatory", 2, 3,
         "RIVR holds carrier authorizations/ETC; assign CPNI; compliance calendar.",
         "Regulatory counsel + RIVR"),
        ("Insurance / bonding inadequacy", "Risk", 2, 3,
         "Enforce Exhibit I limits; additional-insured/waiver of subrogation; "
         "annual COI verification.",
         "Risk + RIVR"),
        ("Transition / end-of-term service disruption", "Operational", 2, 4,
         "12-mo transition assistance; records/customer handover; defined "
         "transition pricing; step-in fallback.",
         "Both parties + Business"),
        ("Cybersecurity / data breach (CPNI)", "Operational / Privacy", 2, 4,
         "Cyber insurance ($5M); CPNI/privacy controls; incident response; §889 "
         "equipment ban compliance.",
         "RIVR + Privacy counsel"),
    ]
    r = hr + 1
    first = r
    for name, cat, L, I, mit, owner in risks:
        C.xl_cell(ws, r, 1, name, kind="text", wrap=True)
        C.xl_cell(ws, r, 2, cat, kind="text", wrap=True)
        C.xl_cell(ws, r, 3, L, kind="input", align="center")
        C.xl_cell(ws, r, 4, I, kind="input", align="center")
        C.xl_cell(ws, r, 5, f"=C{r}*D{r}", kind="formula", align="center", bold=True)
        C.xl_cell(ws, r, 6, mit, kind="text", wrap=True)
        C.xl_cell(ws, r, 7, owner, kind="text", wrap=True)
        r += 1
    last = r - 1

    score_range = f"E{first}:E{last}"
    ws.conditional_formatting.add(score_range, CellIsRule(
        operator="greaterThanOrEqual", formula=["15"],
        fill=PatternFill("solid", fgColor=C.XL_CHECK_BAD),
        font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(score_range, CellIsRule(
        operator="between", formula=["8", "14"],
        fill=PatternFill("solid", fgColor=C.XL_INPUT)))
    ws.conditional_formatting.add(score_range, CellIsRule(
        operator="lessThanOrEqual", formula=["7"],
        fill=PatternFill("solid", fgColor=C.XL_CHECK_OK)))

    # legend for score bands
    lr = last + 2
    ws.cell(row=lr, column=1, value="Score bands:").font = Font(bold=True, size=9)
    for i, (lbl, fill) in enumerate([("High (>=15)", C.XL_CHECK_BAD),
                                     ("Medium (8-14)", C.XL_INPUT),
                                     ("Low (<=7)", C.XL_CHECK_OK)]):
        cc = ws.cell(row=lr, column=2 + i, value=lbl)
        cc.fill = PatternFill("solid", fgColor=fill)
        cc.font = Font(size=9)
        cc.border = C.BORDER

    path = C.xl_save(wb, "06_Review_Materials", "03_Negotiation_Risk_Matrix.xlsx")
    GENERATED.append(path)


def build_cross_reference_matrix():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Document Cross-Reference Matrix",
        "Maps key defined terms and provisions across the transaction suite so "
        "counsel can confirm consistency. Cell values give the governing "
        "section in each document (placeholders where a document does not "
        "address the item, or where numbering is not yet fixed).",
        "• Columns are the seven core documents (02.01–02.07) plus Exhibits.",
        "• Confirm every cross-reference during redlining; section numbers are "
        "indicative and marked as placeholders.",
    ])
    ws = wb.create_sheet("Cross-Reference")
    ws.sheet_properties.tabColor = "1F3864"
    C.xl_title(ws, "Document Cross-Reference Matrix",
               subtitle="Defined terms & key provisions across the transaction documents",
               span=9)
    headers = ["Provision / Defined Term", "Master 02.01", "IRU 02.02",
               "O&M/SLA 02.03", "Financial 02.04", "Grant 02.05",
               "Privacy 02.06", "Transition 02.07", "Exhibits"]
    hr = 5
    C.xl_header_row(ws, hr, headers)
    _set_widths(ws, [30, 15, 15, 15, 15, 15, 15, 15, 16])

    NA = C.PH("not addressed — confirm")
    def S(n):  # section placeholder
        return C.PH(f"§ {n} — confirm")

    rows = [
        ("Parties / recitals", S("1.1"), S("1.1"), S("1"), S("1"), S("1"), S("1"), S("1"), "Cover pages"),
        ("Definitions (defined terms)", S("2"), S("2"), S("2"), S("2"), S("2"), S("2"), S("2"), "As used"),
        ("IRU grant & term (20/25/30)", S("4"), S("3"), NA, S("Sch. A"), NA, NA, NA, "Fiber exhibit"),
        ("Tribal ownership of assets", S("3"), S("3.2"), NA, NA, S("2"), NA, S("2"), "Asset schedule"),
        ("Federal interest (200.316)", S("3.4"), S("3.3"), NA, NA, S("3"), NA, S("4"), C.PH("n/a")),
        ("Reserved Tribal strands", S("4.3"), S("4"), NA, NA, NA, NA, NA, "Fiber exhibit"),
        ("Revenue share / payments", S("6"), S("5"), NA, S("Core"), NA, NA, NA, "Financial exhibit"),
        ("Insurance", S("11"), S("9"), S("8"), NA, NA, NA, S("6"), "Exhibit I"),
        ("Bonding", S("11.5"), NA, NA, NA, NA, NA, NA, "Exhibit I"),
        ("Demarcation point", S("5"), S("4.2"), S("3"), NA, NA, NA, NA, "Network exhibit"),
        ("SLA / severity levels", S("7"), NA, S("Core"), NA, NA, NA, S("5"), "SLA exhibit"),
        ("Cure periods (30/60)", S("14.2"), S("10"), S("9"), NA, S("7"), S("8"), S("7"), C.PH("n/a")),
        ("Default & remedies", S("14"), S("10"), S("9"), NA, S("7"), S("8"), S("7"), C.PH("n/a")),
        ("Step-in rights", S("15"), S("11"), S("10"), NA, S("7.3"), NA, S("3"), C.PH("n/a")),
        ("Program income (200.307)", S("6.6"), NA, NA, S("Rev"), S("4"), NA, NA, C.PH("n/a")),
        ("Procurement (200.317-.327)", S("8"), NA, NA, NA, S("5"), NA, NA, C.PH("n/a")),
        ("Contractor vs subrecipient (200.331)", S("8.4"), NA, NA, NA, S("6"), NA, NA, C.PH("n/a")),
        ("Asset disposition / reversion", S("16"), S("12"), NA, NA, S("9"), NA, S("2"), "Asset schedule"),
        ("Records & audit (200.334)", S("9"), NA, S("11"), S("Audit"), S("8"), S("7"), NA, C.PH("n/a")),
        ("Data / privacy / CPNI", S("12"), NA, S("12"), NA, NA, S("Core"), S("8"), C.PH("n/a")),
        ("Customer ownership", S("6.8"), NA, NA, NA, NA, S("5"), S("3.4"), C.PH("n/a")),
        ("Assignment / change of control", S("17"), S("13"), NA, NA, NA, NA, S("9"), C.PH("n/a")),
        ("Sovereign immunity / dispute resolution", S("18"), S("14"), NA, NA, NA, NA, NA, C.PH("n/a")),
        ("Governing law / jurisdiction", S("18.1"), S("14.1"), S("13"), S("13"), S("10"), S("9"), S("10"), C.PH("n/a")),
        ("Transition obligations", S("16.4"), S("12.3"), NA, NA, NA, NA, S("Core"), C.PH("n/a")),
        ("BABA / §889 compliance", S("8.6"), NA, NA, NA, S("5.4"), NA, NA, C.PH("n/a")),
        ("Environmental / NHPA", S("8.7"), NA, NA, NA, S("5.5"), NA, NA, "Permitting exhibit"),
    ]
    r = hr + 1
    for row in rows:
        _write_row(ws, r, list(row),
                   wraps=[True] + [True] * 8)
        r += 1

    path = C.xl_save(wb, "06_Review_Materials", "04_Document_Cross_Reference_Matrix.xlsx")
    GENERATED.append(path)


def build_grant_compliance_checklist():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Grant Compliance Checklist",
        "Tracks the TBCP / 2 CFR Part 200 compliance obligations for the Lumbee Tribe as "
        "Recipient (with RIVR Tech flow-downs). Status column has a dropdown: "
        "Not Started / In Progress / Complete.",
        "• Assign a Responsible owner and record Evidence / Doc Ref for each item.",
        "• Citations reference the canonical citation library (common.py CITES).",
        "• This checklist supports, but does not replace, grant counsel review.",
    ])
    ws = wb.create_sheet("Grant Compliance")
    ws.sheet_properties.tabColor = "7F007F"
    C.xl_title(ws, "Grant Compliance Checklist",
               subtitle="TBCP Round 3 / 2 CFR Part 200 — Recipient obligations and flow-downs",
               span=6)
    headers = ["Requirement", "Citation", "Responsible", "Status",
               "Evidence / Doc Ref", "Notes"]
    hr = 5
    C.xl_header_row(ws, hr, headers)
    _set_widths(ws, [40, 44, 18, 14, 24, 30])

    Cx = C.CITES
    rows = [
        ("Confirm TBCP eligibility (Tribal Government status)",
         Cx["lumbee_act"] + "; " + Cx["nofo"], "Tribal + Grant counsel", "Not Started",
         C.PH("NTIA confirmation letter"), "Threshold / gating item"),
        ("Environmental review (NEPA)",
         Cx["nepa"], "Grant counsel + RIVR", "Not Started",
         C.PH("environmental determination"), "Start early"),
        ("Historic preservation review (NHPA §106)",
         Cx["nhpa"], "Grant counsel + SHPO/THPO", "Not Started",
         C.PH("§106 clearance"), "SHPO/THPO consultation"),
        ("Build America, Buy America (BABA) compliance",
         Cx["baba"], "RIVR + Grant counsel", "Not Started",
         C.PH("supplier certifications / domestic content"), "Waiver if needed"),
        ("Procurement standards compliance",
         Cx["procurement"] + "; " + Cx["competition"] + "; " + Cx["methods"],
         "Lumbee Tribe procurement + Grant counsel", "Not Started",
         C.PH("procurement file / sole-source justification"), "Document RIVR selection"),
        ("Domestic preference for procurements",
         Cx["domestic"], "Lumbee Tribe procurement", "Not Started",
         C.PH("domestic-preference documentation"), ""),
        ("Property records & inventory",
         Cx["equipment"] + "; " + Cx["intangible"], "Lumbee Tribe + RIVR", "Not Started",
         C.PH("asset register / inventory"), "Physical inventory cadence"),
        ("Real property use / encumbrance controls",
         Cx["real_property"] + "; " + Cx["trust"], "Grant + Tribal counsel", "Not Started",
         C.PH("federal-interest acknowledgments"), "No prohibited liens"),
        ("Program income handling",
         Cx["prog_income"], "Grant counsel + Finance", "Not Started",
         C.PH("program-income method + ledger"), "Confirm method with NTIA"),
        ("Single audit / audit requirements",
         Cx["single_audit"], "Lumbee Tribe Finance + auditor", "Not Started",
         C.PH("single-audit engagement"), "Threshold $1,000,000"),
        ("Record retention & access",
         Cx["records"], "Lumbee Tribe + RIVR", "Not Started",
         C.PH("records-retention policy"), "3 years from final report"),
        ("Reporting cadence (performance & financial)",
         Cx["sac"], "Lumbee Tribe + RIVR", "Not Started",
         C.PH("reporting calendar"), "Per Specific Award Conditions"),
        ("Affordability / 100-20 service commitment",
         Cx["nofo"] + "; " + Cx["usac_lifeline"], "RIVR + Business", "Not Started",
         C.PH("affordable tier documentation"), "100/20 Mbps floor"),
        ("Disposition restrictions",
         Cx["real_property"] + "; " + Cx["closeout"], "Grant counsel", "Not Started",
         C.PH("disposition plan"), "Federal interest persists"),
        ("Conflict of interest policy",
         Cx["conflict"], "Lumbee Tribe + RIVR", "Not Started",
         C.PH("COI policy + disclosures"), "Mandatory disclosures 200.113"),
        ("Debarment / suspension screening",
         Cx["debarment"], "Lumbee Tribe procurement", "Not Started",
         C.PH("SAM.gov exclusion checks"), "Screen all subcontractors"),
        ("§889 covered-telecom equipment ban",
         Cx["telecom_ban"], "RIVR", "Not Started",
         C.PH("§889 compliance certification"), "No covered equipment"),
        ("Flow-down of federal terms to subcontracts",
         Cx["pass_through"] + "; " + Cx["subrecipient"], "Grant counsel + RIVR", "Not Started",
         C.PH("flow-down clause library"), "Contractor vs subrecipient analysis"),
        ("Closeout & post-closeout obligations",
         Cx["closeout"], "Lumbee Tribe + Grant counsel", "Not Started",
         C.PH("closeout package"), "Continuing responsibilities"),
        ("Remedies for noncompliance / termination readiness",
         Cx["remedies"], "Grant counsel", "Not Started",
         C.PH("compliance escalation plan"), "Clawback exposure"),
    ]
    r = hr + 1
    for row in rows:
        _write_row(ws, r, list(row),
                   wraps=[True, True, True, False, True, True])
        r += 1
    last = r - 1

    dv = DataValidation(type="list",
                        formula1='"Not Started,In Progress,Complete"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D{hr+1}:D{last}")

    path = C.xl_save(wb, "06_Review_Materials", "05_Grant_Compliance_Checklist.xlsx")
    GENERATED.append(path)


def build_closing_checklist():
    wb = Workbook()
    C.xl_instructions(wb, [
        "## Closing Checklist",
        "Conditions to signing/closing the Lumbee Tribe / RIVR Tech broadband "
        "transaction. Status dropdown: Not Started / In Progress / Complete / "
        "Waived / N/A.",
        "• Assign Responsible and record the executed Doc Ref for each item.",
        "• Authorizations (05_Authorizations) and review materials "
        "(06_Review_Materials) feed this list.",
    ])
    ws = wb.create_sheet("Closing Checklist")
    ws.sheet_properties.tabColor = "1F3864"
    C.xl_title(ws, "Closing Checklist",
               subtitle="Conditions to execution / closing — draft for counsel review",
               span=6)
    headers = ["Item", "Category", "Responsible", "Status", "Doc Ref", "Notes"]
    hr = 5
    C.xl_header_row(ws, hr, headers)
    _set_widths(ws, [42, 20, 20, 14, 24, 30])

    rows = [
        ("TBCP eligibility confirmation obtained", "Threshold / Grant",
         "Tribal + Grant counsel", "Not Started", C.PH("NTIA letter"), "Gating"),
        ("TBCP award accepted / award documents signed", "Grant",
         "Lumbee Tribe", "Not Started", C.PH("signed award"), "If awarded"),
        ("Lumbee Tribal Council resolution adopted", "Authorization",
         "Lumbee Tribe", "Not Started", "05.01", "This package"),
        ("RIVR Tech corporate authorization executed", "Authorization",
         "RIVR", "Not Started", "05.02", "This package"),
        ("Certificates of authority / incumbency (both)", "Authorization",
         "Both", "Not Started", "05.03", "This package"),
        ("Master Development, Construction & Operating Agreement executed", "Agreement",
         "Both", "Not Started", "02.01", ""),
        ("IRU Agreement executed", "Agreement",
         "Both", "Not Started", "02.02", "Term confirmed"),
        ("O&M / SLA schedule finalized & executed", "Agreement",
         "Both", "Not Started", "02.03", ""),
        ("Financial schedule finalized & executed", "Agreement",
         "Both", "Not Started", "02.04", "Economics confirmed"),
        ("Grant-compliance agreement executed", "Agreement",
         "Both", "Not Started", "02.05", ""),
        ("Privacy / data agreement executed", "Agreement",
         "Both", "Not Started", "02.06", "CPNI"),
        ("Transition agreement executed", "Agreement",
         "Both", "Not Started", "02.07", ""),
        ("Exhibits finalized (fiber, network, insurance, maps)", "Exhibits",
         "Both", "Not Started", C.PH("exhibit set"), "Incl. Exhibit I"),
        ("Insurance certificates delivered", "Risk",
         "RIVR", "Not Started", C.PH("COIs"), "Additional insured"),
        ("Performance / payment bonds delivered", "Risk",
         "RIVR", "Not Started", C.PH("bonds"), "Construction"),
        ("NTIA approvals / consents (as required)", "Grant",
         "Grant counsel", "Not Started", C.PH("consents"), "If required"),
        ("BIA / land-status approvals (if applicable)", "Regulatory",
         "Tribal counsel", "Not Started", C.PH("BIA consent"), "May be N/A"),
        ("Environmental / NHPA clearance", "Regulatory",
         "Grant counsel", "Not Started", C.PH("clearances"), ""),
        ("Easements / rights-of-way secured", "Real Estate",
         "Real-estate counsel", "Not Started", C.PH("recorded easements"), ""),
        ("Pole-attachment agreements in place", "Real Estate",
         "RIVR", "Not Started", C.PH("pole agreements"), ""),
        ("Legal opinion letters delivered", "Legal",
         "Both counsel", "Not Started", C.PH("opinions"), "Authority / enforceability"),
        ("Good-standing certificates (both entities)", "Legal",
         "Both", "Not Started", C.PH("good-standing"), ""),
        ("Sovereign-immunity / dispute-resolution terms resolved", "Legal",
         "Both counsel", "Not Started", C.PH("final clause"), "Separate Council action"),
        ("Regulatory authorizations / ETC status confirmed", "Regulatory",
         "Regulatory counsel", "Not Started", C.PH("authorizations"), ""),
        ("Funds-flow / accounts established", "Financial",
         "Finance", "Not Started", C.PH("account setup"), "Program income"),
    ]
    r = hr + 1
    for row in rows:
        _write_row(ws, r, list(row),
                   wraps=[True, False, False, False, True, True])
        r += 1
    last = r - 1

    dv = DataValidation(type="list",
                        formula1='"Not Started,In Progress,Complete,Waived,N/A"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D{hr+1}:D{last}")

    path = C.xl_save(wb, "06_Review_Materials", "06_Closing_Checklist.xlsx")
    GENERATED.append(path)


def build_due_diligence_request_list():
    doc = C.new_doc()
    C.add_cover(
        doc,
        "06.07",
        "Due Diligence Request List",
        "Documents and Information Requested from Each Party — Lumbee Tribe / "
        "RIVR Tech Broadband Partnership",
    )
    C.setup_header_footer(doc, "Due Diligence Request List")
    C.status_banner(doc)
    C.spacer(doc, 1)
    p = doc.add_paragraph()
    C._emit_runs(p,
        "This list requests documents and information from each Party to support "
        "the transaction, grant compliance, and legal review. Please provide items "
        "electronically to the shared data room, noting any item that is not "
        "applicable or not available. Bracketed placeholders indicate items to be "
        "completed by the requesting/responding party.")
    C.spacer(doc, 1)
    C.flag_para(
        doc, C.FLAG_GRANT,
        "Priority item: Lumbee Tribe to provide documentation bearing on federal-"
        "recognition / TBCP eligibility (Lead item on the Attorney Review Issue "
        "List).",
    )

    def category(doc, name):
        plain_bold(doc, name, size=12, color=C.NAVY)

    def req_table(doc, items):
        C.add_table(
            doc,
            ["#", "Document / Item Requested", "Purpose / Notes"],
            [[str(i + 1), it[0], it[1]] for i, it in enumerate(items)],
            widths=[0.4, 3.4, 2.7],
            font_size=9,
            col_align=["c", "l", "l"],
        )
        C.spacer(doc, 1)

    # ---- LUMBEE TRIBE ------------------------------------------------------------
    plain_bold(doc, f"PART A — REQUESTED FROM {C.TRIBE_FULL.upper()}",
               size=13, color=C.NAVY)
    C.spacer(doc, 1)
    category(doc, "A-1. Recognition & Eligibility")
    req_table(doc, [
        ("Documentation of federal-recognition status (Lumbee Act and related materials)",
         "Threshold TBCP eligibility " + C.FLAG_GRANT),
        ("Any prior NTIA / federal correspondence on Tribal eligibility",
         "Confirm eligibility basis"),
        ("Tribal governing documents (constitution, charter, ordinances)",
         "Authority to contract / waive immunity"),
        (C.PH("other eligibility documentation"), "As applicable"),
    ])
    category(doc, "A-2. Governance & Authorizations")
    req_table(doc, [
        ("Tribal Council membership, officers, and meeting/quorum rules",
         "Validity of resolution"),
        ("Draft/adopted Council resolution authorizing the transaction",
         "See 05.01"),
        ("Signature authority / delegation documentation",
         "Authorized representatives"),
    ])
    category(doc, "A-3. Land, Easements & Rights-of-Way")
    req_table(doc, [
        ("Land / parcel records for the service territory and asset sites",
         "Siting and easements"),
        ("Existing easements, rights-of-way, and land-status (trust/fee) records",
         "ROW durability; BIA triggers"),
        (C.PH("survey / GIS data if available"), "Network design"),
    ])
    category(doc, "A-4. Existing Grants & Funding")
    req_table(doc, [
        ("List of existing/pending broadband grants and awards",
         "Duplication-of-benefits check " + C.FLAG_GRANT),
        ("Related funding agreements and service commitments",
         "Overlap / conflict analysis"),
    ])
    category(doc, "A-5. Financial & Organizational")
    req_table(doc, [
        (C.PH("Tribal / Tribal-entity financial statements (as applicable)"),
         "Capacity as Recipient"),
        ("Designated Tribal entity formation documents (if used)",
         "Eligible-entity structure"),
    ])

    doc.add_page_break()
    C.status_banner(doc)
    C.spacer(doc, 1)

    # ---- RIVR TECH --------------------------------------------------------
    plain_bold(doc, f"PART B — REQUESTED FROM {C.OPERATOR_FULL.upper()}",
               size=13, color=C.NAVY)
    C.spacer(doc, 1)
    category(doc, "B-1. Formation & Good Standing")
    req_table(doc, [
        ("Articles of organization / operating agreement (and amendments)",
         "Authority to contract"),
        ("Certificate of good standing / existence (state of formation + N.C.)",
         "Good standing"),
        ("Parent/affiliate structure (LREMC) and required consents",
         "Change-of-control / authority"),
    ])
    category(doc, "B-2. Financial Condition")
    req_table(doc, [
        (C.PH("audited/reviewed financial statements (2-3 years)"),
         "Operator capacity"),
        (C.PH("evidence of financing / capital for construction"),
         "Funding of RIVR contribution"),
    ])
    category(doc, "B-3. Network & Technical")
    req_table(doc, [
        ("Existing-network maps (middle-mile / core / RIVR Existing Network)",
         "Hybrid design; demarcation"),
        (C.PH("proposed network design / homes-passed model"),
         "Footprint; economics"),
        ("Fiber, OSP, and equipment standards (BABA/§889 posture)",
         "Compliance " + C.FLAG_GRANT),
    ])
    category(doc, "B-4. Licenses, Authorizations & Regulatory")
    req_table(doc, [
        ("FCC / state carrier authorizations and filings",
         "Regulatory authority"),
        ("ETC designation status (if any) and USAC/USF participation",
         "ETC / USF obligations"),
        ("CPNI / privacy policies and procedures",
         "Privacy agreement"),
    ])
    category(doc, "B-5. Insurance, Bonding & Contracts")
    req_table(doc, [
        ("Certificates of insurance / current coverage schedule",
         "Exhibit I compliance"),
        (C.PH("surety / bonding capacity letter"), "Construction bonds"),
        ("Material contracts, subcontractor agreements, and supplier terms",
         "Flow-down / dependencies"),
    ])
    category(doc, "B-6. Litigation & Compliance")
    req_table(doc, [
        ("Pending / threatened litigation and material disputes",
         "Risk assessment"),
        ("SAM.gov / debarment-suspension status",
         "Eligibility " + C.FLAG_GRANT),
        (C.PH("compliance / debarment certifications"), "Grant flow-down"),
    ])

    C.spacer(doc, 1)
    C.flag_para(
        doc, C.FLAG_ATTORNEY,
        "Counsel to tailor this request list to the final deal structure and to "
        "the subrecipient-vs-contractor determination, which affects the scope of "
        "RIVR Tech diligence and flow-down obligations.",
    )

    path = C.save(doc, "06_Review_Materials", "07_Due_Diligence_Request_List.docx")
    GENERATED.append(path)


# ===========================================================================
#  MAIN
# ===========================================================================
def main():
    build_tribal_resolution()
    build_rivr_corporate_authorization()
    build_certificate_of_authority()
    build_attorney_issue_list()
    build_open_business_decisions()
    build_negotiation_risk_matrix()
    build_cross_reference_matrix()
    build_grant_compliance_checklist()
    build_closing_checklist()
    build_due_diligence_request_list()

    print("Generated {} files:".format(len(GENERATED)))
    for p in GENERATED:
        exists = "OK" if os.path.exists(p) else "MISSING"
        print(f"  [{exists}] {p}")


if __name__ == "__main__":
    main()
