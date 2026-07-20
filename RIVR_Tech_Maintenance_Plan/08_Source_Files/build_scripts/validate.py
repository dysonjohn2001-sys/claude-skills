# -*- coding: utf-8 -*-
"""Structural validation of all deliverables (LibreOffice render unavailable in sandbox)."""
import os, glob
from docx import Document
from docx.shared import Inches, Emu
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Emu as PEmu

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
ROOT = os.path.join(BASE, "RIVR_Tech_Maintenance_Plan")

USABLE_DOCX_IN = 6.7   # 8.5 - 2*0.9 margins
SLIDE_W_IN = 13.333
issues = []
ok = []


def check_docx(path):
    d = Document(path)
    name = os.path.basename(path)
    nparas = len([p for p in d.paragraphs])
    ntables = len(d.tables)
    # empty content?
    text_len = sum(len(p.text) for p in d.paragraphs)
    if text_len < 200:
        issues.append(f"[docx] {name}: suspiciously little text ({text_len} chars)")
    # table width overflow
    for ti, t in enumerate(d.tables):
        widths = []
        for cell in t.rows[0].cells:
            w = cell.width
            widths.append(Emu(w).inches if w else None)
        if all(w is not None for w in widths):
            total = sum(widths)
            if total > USABLE_DOCX_IN + 0.15:
                issues.append(f"[docx] {name}: table {ti} width {total:.2f}in > {USABLE_DOCX_IN}in usable")
    ok.append(f"[docx] {name}: {nparas} paras, {ntables} tables, {text_len} chars OK")


def check_xlsx(path):
    wb = load_workbook(path)
    name = os.path.basename(path)
    sheets = wb.sheetnames
    total_charts = 0
    for ws in wb.worksheets:
        total_charts += len(getattr(ws, "_charts", []))
    ok.append(f"[xlsx] {name}: sheets={len(sheets)} {sheets} charts={total_charts}")
    # check no sheet is empty
    for ws in wb.worksheets:
        if ws.max_row < 2:
            issues.append(f"[xlsx] {name}: sheet '{ws.title}' looks empty")


def check_pptx(path):
    prs = Presentation(path)
    name = os.path.basename(path)
    n = len(prs.slides._sldIdLst)
    sw = prs.slide_width
    sh = prs.slide_height
    overflow = 0
    empty = 0
    for si, slide in enumerate(prs.slides, 1):
        has_text = False
        for shp in slide.shapes:
            # off-slide / overflow check
            try:
                if shp.left is not None and shp.top is not None and shp.width and shp.height:
                    right = shp.left + shp.width
                    bottom = shp.top + shp.height
                    if right > sw + PEmu(Inches(0.05)) or bottom > sh + PEmu(Inches(0.05)):
                        # allow footer bar exactly at edges
                        if bottom > sh + PEmu(Inches(0.1)) or right > sw + PEmu(Inches(0.1)):
                            overflow += 1
            except Exception:
                pass
            if shp.has_text_frame and shp.text_frame.text.strip():
                has_text = True
            if shp.has_table:
                has_text = True
        if not has_text:
            empty += 1
    if overflow:
        issues.append(f"[pptx] {name}: {overflow} shapes may overflow slide bounds")
    if empty:
        issues.append(f"[pptx] {name}: {empty} slides with no text/table")
    ok.append(f"[pptx] {name}: {n} slides, overflow={overflow}, empty={empty}")


REQUIRED = [
    "01_Executive/RIVR_Tech_Maintenance_Business_Case.docx",
    "01_Executive/RIVR_Tech_Executive_Decision_Summary.docx",
    "02_Policy_and_SOP/Customer_Premises_Maintenance_Policy.docx",
    "02_Policy_and_SOP/Support_and_Field_Operations_SOP.docx",
    "02_Policy_and_SOP/Trouble_Visit_Responsibility_Matrix.xlsx",
    "03_Financial_Model/Maintenance_Program_Financial_Model.xlsx",
    "04_Training/Maintenance_Program_Employee_Training.pptx",
    "04_Training/Employee_Quick_Reference_Guide.docx",
    "04_Training/Facilitator_and_Manager_Coaching_Guide.docx",
    "04_Training/Employee_Knowledge_Assessment.docx",
    "05_Customer_Communications/Customer_Communication_Package.docx",
    "05_Customer_Communications/Customer_Maintenance_Handout.docx",
    "06_Implementation/90_Day_Implementation_Tracker.xlsx",
    "06_Implementation/Maintenance_Program_KPI_Dashboard.xlsx",
    "07_Final_Presentation/Maintenance_Program_Executive_Presentation.pptx",
    "RIVR_Tech_Maintenance_Plan_README.md",
]

print("=== EXISTENCE CHECK ===")
missing = []
for rel in REQUIRED:
    fp = os.path.join(ROOT, rel)
    exists = os.path.exists(fp)
    size = os.path.getsize(fp) if exists else 0
    print(f"  {'OK ' if exists else 'MISSING'} {rel} ({size} bytes)")
    if not exists:
        missing.append(rel)

print("\n=== STRUCTURAL VALIDATION ===")
for rel in REQUIRED:
    fp = os.path.join(ROOT, rel)
    if not os.path.exists(fp):
        continue
    if fp.endswith(".docx"):
        check_docx(fp)
    elif fp.endswith(".xlsx"):
        check_xlsx(fp)
    elif fp.endswith(".pptx"):
        check_pptx(fp)

for line in ok:
    print("  " + line)

print("\n=== ISSUES ===")
if missing:
    for m in missing:
        print("  MISSING:", m)
if issues:
    for i in issues:
        print("  " + i)
if not issues and not missing:
    print("  none")
