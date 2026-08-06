import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx_helpers import *
from docx.shared import Pt

doc = new_doc()
cover(doc,
    "Executive Implementation Plan",
    "Lifeline & Lumbee Broadband Affordability Initiative",
    "01 of 13",
    "Program Owner (with Executive Sponsor)",
    version="v0.1 — DRAFT (Phase 0: ETC application pending)",
    extra=("Do NOT launch, market RIVR Tech as an approved Lifeline provider, or claim reimbursement "
        "until the final NCUC ETC designation order is issued and USAC onboarding is complete. All "
        "prices and volumes in this plan are illustrative placeholders pending confirmation."))
toc(doc)

# ---------------- Executive summary ----------------
h1(doc, "1. Executive summary")
para(doc, "RIVR Tech (LREMC Technologies, LLC) has applied to the North Carolina Utilities Commission "
    "(NCUC) for designation as an Eligible Telecommunications Carrier (ETC). Upon approval, RIVR Tech "
    "intends to (a) apply the federal Lifeline discount to eligible customers on its existing fiber "
    "internet plans, (b) launch a new symmetrical 100/100 Mbps Senior Internet Plan targeting a $10 "
    "monthly customer payment, and (c) coordinate a broadband affordability initiative with the newly "
    "federally-recognized Lumbee Tribe of North Carolina.")
callout(doc, "THE SINGLE MOST IMPORTANT COMPLIANCE POINT: Lumbee tribal membership and residence in "
    "Robeson, Hoke, Scotland, or Cumberland County do NOT, by themselves, qualify a customer for the "
    "enhanced $34.25 Tribal Lifeline benefit. Enhanced support requires the customer's PRINCIPAL "
    "RESIDENCE to be on qualifying 'Tribal lands' as defined in 47 C.F.R. §54.400(e), verified "
    "per-address via USAC's Tribal Lands Verification Tool. The Lumbee currently have no reservation, "
    "no land in trust, and no FCC §54.412 designation — so the working assumption is ZERO enhanced-"
    "Tribal-eligible addresses until a written USAC/legal determination confirms otherwise (SR-4, SR-5).",
    kind="stop")
para(doc, "Key recommendations (each an executive decision in §11):", bold=True)
for t in [
    "Adopt the blended 'Alternative D' subsidy design: apply Lifeline first, then a defined RIVR/Tribe "
    "top-up so every approved participant reaches the $10 net target. $10 is the NET price after "
    "subsidies, not the retail list price.",
    "Set the senior age threshold at 62 (balances reach against company-funded exposure).",
    "Keep the enhanced Tribal benefit OFF (0 subscribers) until the written Tribal-lands determination "
    "is obtained; never condition it on Lumbee membership.",
    "Fold or waive equipment on the Senior plan to protect the $10 target; keep voice an optional, "
    "separately-disclosed add-on (broadband is the Lifeline-supported service).",
    "Obtain a tax opinion before advertising any 'plus tax'; broadband access itself is non-taxable "
    "under the Internet Tax Freedom Act.",
    "Keep four funding streams — federal Lifeline, any grant/TBCP funds, RIVR-funded discounts, and "
    "Lumbee Tribe contributions — in strictly separate accounting categories.",
]:
    bullet(doc, t)
para(doc, "This plan is one of 13 deliverables (index in §12). Figures are illustrative placeholders; "
    "confirmed RIVR pricing, a tax opinion, and the Tribal-lands legal determination are prerequisites "
    "to launch and are tracked in the Open Issues & Legal Decisions Log (Deliverable 13).", italic=True, size=9, color=GREY)

# ---------------- Program objectives ----------------
h1(doc, "2. Program objectives & scope")
for t in [
    "Apply Lifeline discounts to eligible customers on existing RIVR Tech fiber plans.",
    "Serve eligible Lumbee tribal members through a coordinated, privacy-respecting affordability initiative.",
    "Correctly determine standard ($9.25) vs. enhanced Tribal (up to $34.25) benefit eligibility on a "
    "per-household, per-address basis.",
    "Launch a symmetrical 100/100 Mbps Senior Internet Plan at a $10 net customer payment (plus only "
    "legally-applicable, tax-advisor-confirmed taxes/fees).",
    "Stand up all operational, compliance, billing, financial, technology, training, reporting, and "
    "communications capabilities required to run the program after ETC approval.",
]:
    numbered(doc, t)

# ---------------- Critical compliance distinctions ----------------
h1(doc, "3. Critical compliance distinctions — four separate tests")
para(doc, "The program is built on four independent tests. Confusing them is the primary compliance risk.")
table(doc,
    ["Test", "What it establishes", "How verified", "What it does NOT establish"],
    [
     ["1. Standard Lifeline eligibility", "Household qualifies for up to $9.25/mo (income ≤135% FPG or a qualifying program)",
      "USAC National Verifier approval BEFORE enrollment (SR-1, SR-3.1)", "Does not establish Tribal-lands or senior status"],
     ["2. Enhanced Tribal eligibility", "Principal residence is ON qualifying Tribal lands → additional up to $25 ($34.25 total)",
      "USAC Tribal Lands Verification Tool on the address (SR-4)", "Is NOT established by tribal membership or county of residence"],
     ["3. Lumbee tribal membership", "May support outreach or a Tribe-funded benefit only",
      "Tribe-run, privacy-preserving verification (no roll transfer)", "Is NOT a substitute for National Verifier approval or Tribal-lands verification (SR-5)"],
     ["4. Senior status", "Eligibility for the $10 Senior plan design",
      "Age ≥ threshold (recommend 62)", "Age ALONE does not establish Lifeline eligibility"],
    ],
    widths=[1.5, 2.2, 2.0, 2.1], small=True)
h2(doc, "3.1 Why the Lumbee case requires special care")
para(doc, "The Lumbee Tribe received full federal recognition via Public Law 119-60 (§8803 of the NDAA "
    "for FY2026, signed December 18, 2025) and was added to the DOI list in January 2026 (SR-5.1–5.2). "
    "That law extends recognition and federal services and designates a four-county service delivery "
    "area, and it AUTHORIZES a future land-into-trust process — but it does not itself create a "
    "reservation or place any land in trust (SR-5.3–5.4). Under 47 C.F.R. §54.400(e), off-reservation "
    "land becomes Lifeline 'Tribal lands' only through an FCC §54.412 designation made on the Tribe's "
    "request (SR-4.3). None of a reservation, trust land, or §54.412 designation exists today, so no "
    "Lumbee address is presumptively Tribal lands (SR-5.5).")
callout(doc, "Action: obtain a written USAC and/or outside-counsel determination on whether any Lumbee "
    "service-area, trust, or dependent-community land satisfies §54.400(e) or has a §54.412 designation, "
    "and confirm none is pending. Until then, model and operate with zero enhanced-Tribal subscribers.",
    kind="warn")

# ---------------- Current state ----------------
h1(doc, "4. Current-state snapshot (RIVR Tech)")
para(doc, "Confirmed vs. to-obtain (full inventory in Deliverable 02).")
table(doc, ["Item", "Status"], [
    ["Fiber lineup (symmetrical)", "CONFIRMED: 250/250 (entry) → 500/500 → 1 Gig; unlimited data (SR-8.2)"],
    ["Proposed Senior plan", "NEW 100/100 tier BELOW the 250 entry tier — exceeds Lifeline's 25/3 & 1,280 GB minimums (SR-2, SR-8.3)"],
    ["Monthly plan prices", "NOT PUBLISHED — obtain confirmed price sheet from RIVR (SR-8.6)"],
    ["Equipment", "Indoor Wi-Fi 6 router $5/mo; outdoor mesh $10/mo; install advertised free — CONFIRM & resolve 'free vs $5' conflict (SR-8.4)"],
    ["Voice", "Unlimited calling; price not published — obtain (SR-8.5)"],
    ["NCUC ETC docket", "Not found via public search — CONFIRM docket/status with NCUC (SR-7.3)"],
    ["ETC order", "PENDING — prerequisite to launch (SR-7.4)"],
], widths=[2.3, 5.0], small=True)

# ---------------- Senior plan & alternatives ----------------
h1(doc, "5. Senior 100/100 plan & subsidy alternatives")
para(doc, "The plan is a new symmetrical 100/100 Mbps tier (download and upload), unlimited data (meets "
    "the 1,280 GB minimum), with premium services (static IP, etc.) excluded. Equipment/managed-Wi-Fi, "
    "installation, and optional voice are decided in §11. Four subsidy designs were evaluated:")
table(doc, ["Alt", "Design", "Compliance risk", "Assessment"], [
    ["A", "Universal $10 retail for all qualifying seniors; RIVR absorbs the gap for non-Lifeline seniors; Lifeline applied where permissible",
     "Low–Med", "Simple to message; highest company subsidy on non-Lifeline seniors"],
    ["B", "Set retail so the bill nets to $10 after the $9.25 standard credit",
     "Med", "A tribal-lands-qualified sub would net BELOW $10 (over-recovery risk); non-Lifeline seniors pay more"],
    ["C", "Set retail so the bill nets to $10 after the $34.25 enhanced credit",
     "HIGH", "Only works for the (currently zero) Tribal-lands population; mis-signals eligibility; standard-LL seniors stay above $10"],
    ["D", "Apply Lifeline first, then a defined RIVR/Tribe top-up so every approved participant reaches $10",
     "Low", "RECOMMENDED — compliance-clean, protects all cohorts, simplest to administer in NISC"],
], widths=[0.5, 3.0, 1.1, 2.6], small=True)
callout(doc, "Recommendation: Alternative D. $10 is the NET target after subsidies; the retail list rate "
    "is unchanged. Per-subscriber and portfolio economics for all four designs are in Deliverable 06 "
    "(Financial Model).", kind="ok")
h2(doc, "5.1 Plan policy elements to finalize")
for t in ["Equipment & managed-Wi-Fi policy (recommend fold/waive to hold $10; confirm tax).",
    "Installation policy (recommend waive for program).",
    "Static IP & premium-service exclusions (excluded).",
    "Upgrade options (senior may upgrade to 250/500/1 Gig at standard rates, Lifeline credit follows).",
    "Service-call & customer-caused-trouble policy; nonpayment rules; equipment-return requirements.",
    "Transfer/move rules (re-run Tribal-lands check on the new address; update NLAD within 10 business days).",
    "Renewal/recertification; loss-of-Lifeline handling; loss of senior/tribal program status.",
    "Reasonable accommodations & authorized representatives."]:
    bullet(doc, t)

# ---------------- Pricing & subsidy summary ----------------
h1(doc, "6. Pricing & subsidy summary")
para(doc, "Illustrative per-subscriber monthly economics (placeholder retail $55.00 for the new Senior "
    "tier; confirm from RIVR). Live formulas in Deliverable 06; full matrix in Deliverable 02. All "
    "pre-tax — no 'plus tax' until the tax advisor confirms (SR-9.9).")
table(doc, ["Scenario", "Customer pays (pre-tax)", "USAC reimb.", "Notes"], [
    ["Regular retail", "$55.00", "$0.00", "New Senior tier list rate (placeholder)"],
    ["+ Standard Lifeline", "$45.75", "$9.25", "After $9.25 credit"],
    ["+ Enhanced Tribal", "$20.75", "$34.25", "ONLY for USAC-verified Tribal-lands address (≈0 today)"],
    ["Alt D → $10 target (std LL)", "$10.00", "$9.25", "RIVR/Tribe top-up covers the remainder"],
    ["Alt D → $10 target (no Lifeline)", "$10.00", "$0.00", "Fully company/Tribe-funded gap"],
], widths=[2.4, 1.6, 1.1, 2.1], small=True)
para(doc, "Federal Lifeline support, any TBCP/grant funds, RIVR-funded discounts, and Lumbee Tribe "
    "contributions are kept in separate accounting categories; TBCP funds are not treated as a "
    "recurring consumer subsidy unless the award expressly authorizes it (SR-10).", size=9, italic=True, color=GREY)

# ---------------- Operational workstreams ----------------
h1(doc, "7. Operational workstreams (overview)")
para(doc, "Detailed procedures live in Deliverables 03–05, 07–08. Summary:")
table(doc, ["Workstream", "What it covers", "Deliverable"], [
    ["Governance & approval", "Sponsor, Program Owner, Compliance, Counsel, Finance, Billing, IT, CS, Sales, Marketing, Network, Tribe, Audit; RACI", "01 §8, 08"],
    ["USAC & regulatory readiness", "ETC order, FRN, SAM.gov/UEI, SAC, 498, RAD, NV/NLAD/LCS, Forms 481/555, records, advertising, WFA", "03, 07"],
    ["Enrollment", "10 workflows + universal control points; NV-before-NLAD; QC; retention", "04"],
    ["Billing / NISC", "Rate & credit codes, tax config, GL mapping, bill text, effective dates, reversals, exceptions", "05"],
    ["Reimbursement & reconciliation", "Monthly close; three-way NISC↔NLAD↔LCS reconciliation; officer certification", "05"],
    ["Recertification & de-enrollment", "Annual recert (60-day), non-usage [confirm], de-enrollment, notices, no retroactive billing", "03, 04"],
    ["Privacy & information security", "Consent, data minimization, role-based access, no SSN/roll in unsecured fields, audit logs", "03, 05"],
], widths=[1.7, 4.2, 1.0], small=True)

# ---------------- Roadmap ----------------
h1(doc, "8. Implementation roadmap & governance")
para(doc, "Five phases (full task/owner/date/evidence detail + RACI in Deliverable 08).")
table(doc, ["Phase", "Focus", "Gate to exit"], [
    ["0 — ETC pending (now)", "Research, design, financial model, policy drafting, Tribe discussions, tax & legal determinations", "Signed source register + determinations obtained"],
    ["1 — ETC approval & USAC onboarding", "FRN/SAM/SAC/498/RAD; NV/NLAD/LCS access; NISC config & test; training", "Final ETC order + onboarding complete"],
    ["2 — Controlled pilot", "≤25–50 existing customers through enrollment→NLAD→bill→LCS→reconcile", "Clean three-way reconciliation"],
    ["3 — Public launch", "Coordinated RIVR + Lumbee comms; open enrollment across ETC area", "Launch readiness scorecard GO (Deliverable 12)"],
    ["4 — Stabilization", "30/60/90-day reviews; correct enrollment/billing/claims/training", "Issues closed; steady-state metrics"],
], widths=[1.9, 3.7, 1.7], small=True)
callout(doc, "Do not market RIVR Tech as an approved Lifeline provider during Phase 0. Public launch "
    "(Phase 3) follows a passing pilot, not before.", kind="warn")
h2(doc, "8.1 Governance roles")
para(doc, "Executive Sponsor; Program Owner; Compliance Officer; Regulatory Counsel; Finance; Accounting; "
    "Billing; IT; Customer Service; Sales; Marketing; Network Operations; Lumbee Tribe representatives; "
    "Internal Audit. Full RACI in Deliverable 08.")

# ---------------- Risks ----------------
h1(doc, "9. Key risks & mitigations")
table(doc, ["Risk", "Mitigation"], [
    ["Improperly claiming enhanced Tribal benefit based on membership/county", "Address-only verification via USAC tool; 0-subscriber default; written determination; training & prohibited-statements list (03, 10)"],
    ["Launching before ETC/USAC readiness", "Readiness scorecard gate (12); no marketing as approved provider until Phase 1 done"],
    ["Advertising taxes that don't apply", "No 'plus tax' until tax opinion; ITFA bars internet-access tax (09, SR-9)"],
    ["Commingling funding streams", "Separate GL categories; monthly reconciliation; Tribe settlement invoicing (05, SR-10)"],
    ["Duplicate/one-per-household violations", "National Verifier + NLAD de-dup; household worksheet (03, 04)"],
    ["Privacy exposure of eligibility/roll data", "Data minimization; role-based access; no roll transfer; audit logs (03, 05, 11)"],
    ["Unconfirmed RIVR pricing drives bad decisions", "Obtain confirmed price sheet before finalizing model & rates (02, 13)"],
], widths=[3.0, 4.3], small=True)

# ---------------- Concise financial view ----------------
h1(doc, "10. Financial summary")
para(doc, "The Financial Model (Deliverable 06) is fully formula-driven with editable assumptions. It "
    "computes per-subscriber contribution margin across all six pricing scenarios and the four senior "
    "alternatives, a 12-month subscriber build (conservative 0.6x / expected 1.0x / aggressive 1.5x), a "
    "contribution P&L, break-even, USAC cash-timing, and a 36-month scenario summary. The single largest "
    "lever on sustainability is the National Verifier approval rate — moving seniors from the fully "
    "company-funded cohort into the Lifeline-supported cohort (which carries the $9.25 USAC offset). "
    "Because retail and cost inputs are still placeholders, the sign of contribution margin must be "
    "re-checked once RIVR confirms actual figures.")

# ---------------- REQUIRED EXECUTIVE DECISIONS ----------------
h1(doc, "11. Required executive decisions")
para(doc, "Each item below has a recommended decision, rationale, responsible decision-maker, and "
    "required-by date. These mirror the Open Issues & Legal Decisions Log (Deliverable 13).")
table(doc,
    ["#", "Decision", "Recommended", "Rationale", "Decision-maker", "By"],
    [
     ["1", "Senior eligibility age", "62", "SSA early-retirement anchor; balances reach vs. company subsidy (65=narrow/cheaper, 60=broad/costly)", "Exec Sponsor", "Phase 0 end"],
     ["2", "$10 = retail or net?", "NET after subsidies (Alt D)", "Avoids over-recovery on tribal; protects non-Lifeline seniors; simplest NISC admin", "Exec Sponsor + CFO", "Phase 0 end"],
     ["3", "Company-funded discount for non-Lifeline seniors?", "Yes, capped (Alt D)", "Equity of a senior plan; but full cost w/o USAC offset — cap exposure", "CFO", "Phase 0 end"],
     ["4", "Lumbee Tribe additional funding?", "Pursue via MOU; base case $0", "No commitment yet; keep separate accounting (SR-10)", "Exec Sponsor + Tribe", "Phase 1"],
     ["5", "Equipment & installation charges", "Waive/fold equipment; free install", "$5/mo router would break the $10 target; confirm tax (SR-8.4, SR-9)", "Product + CFO", "Phase 0 end"],
     ["6", "Bundle voice?", "Optional add-on only", "Broadband is the Lifeline-supported service; voice is taxable/assessed (SR-9)", "Product + Tax", "Phase 1"],
     ["7", "Geographic scope", "Approved ETC service area only", "Cannot serve/claim outside the designated area (SR-6.1, SR-7)", "Reg. Counsel", "At ETC order"],
     ["8", "Pilot size", "≤25–50 existing customers", "De-risk enrollment/billing/claims before scale", "Program Owner", "Phase 2"],
     ["9", "Launch date", "After pilot 3-way reconciliation passes", "Green readiness scorecard gate (12)", "Exec Sponsor", "Phase 3"],
     ["10", "Branding", "Co-branded pending written Tribe approval", "Respect sovereignty; Tribe controls name/logo (11)", "Marketing + Tribe", "Phase 3"],
     ["11", "Customer migration rules", "Opt-in; existing customers keep plan + Lifeline credit or move to Senior plan", "Consent-based; avoids disruption", "Program Owner", "Phase 1"],
     ["12", "Treatment of customers who lose eligibility", "Prospective conversion to retail; notice; NO retroactive billing", "Compliance rule; future-dated reversals only (SR, 05)", "Compliance", "Phase 1"],
     ["13", "Tax treatment", "Obtain tax opinion; no 'plus tax' until confirmed", "ITFA bars internet-access tax; other components pending (SR-9)", "Tax Advisor + CFO", "Before launch"],
     ["14", "Enhanced Tribal legal determination", "Obtain written USAC/legal determination; assume 0 until confirmed", "Membership/county ≠ Tribal lands; §54.400(e)/§54.412 (SR-4, SR-5)", "Reg. Counsel + USAC", "Before enroll"],
    ],
    widths=[0.25, 1.35, 1.45, 2.1, 1.05, 0.75], small=True)

# ---------------- Deliverables index ----------------
h1(doc, "12. Deliverables index")
table(doc, ["#", "Deliverable", "Format"], [
    ["00", "Regulatory Source Register & Verified-Facts Base", "Markdown"],
    ["01", "Executive Implementation Plan (this document)", "Word"],
    ["02", "Current Plan & Pricing Audit", "Excel"],
    ["03", "Lifeline Operations Manual", "Word"],
    ["04", "Enrollment & De-enrollment SOPs", "Word"],
    ["05", "NISC Configuration & Billing Specification", "Word"],
    ["06", "Lifeline Financial Model", "Excel"],
    ["07", "Compliance & Audit Checklist", "Excel"],
    ["08", "Implementation Roadmap & RACI", "Excel"],
    ["09", "Customer Communications Toolkit", "Word"],
    ["10", "Employee Training Guide", "Word"],
    ["11", "Lumbee Tribe Partnership MOU Exhibit", "Word"],
    ["12", "Executive Readiness Scorecard", "Excel"],
    ["13", "Open Issues & Legal Decisions Log", "Excel"],
], widths=[0.5, 5.3, 1.2], small=True)
para(doc, "Note: Deliverables are provided in editable Office formats (.docx/.xlsx) rather than the "
    "originally-listed static formats, and a Regulatory Source Register (00) was added as the shared "
    "citation base. All Excel models use visible formulas and editable, color-coded assumption cells.",
    italic=True, size=9, color=GREY)

footer_revhist(doc)
out = "/home/user/claude-skills/deliverables/RIVR-Tech-Lifeline-Lumbee/01_Executive_Implementation_Plan.docx"
doc.save(out)
from docx import Document
Document(out)
print("OK 01 ->", out)
