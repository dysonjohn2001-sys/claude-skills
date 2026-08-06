import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx_helpers import *

PATH = "/home/user/claude-skills/deliverables/RIVR-Tech-Lifeline-Lumbee/11_Lumbee_Tribe_Partnership_MOU_Exhibit.docx"
PH = "[PLACEHOLDER — LEGAL REVIEW REQUIRED]"

doc = new_doc()

cover(
    doc,
    title="Memorandum of Understanding — Exhibit",
    subtitle="Broadband Affordability Initiative — RIVR Tech & the Lumbee Tribe of North Carolina",
    deliverable_no="11 — Lumbee Tribe Partnership MOU (Draft Exhibit)",
    owner="Executive Sponsor + Legal Counsel (both parties)",
    approvers="________________________  (authorized official / title / date) — each party",
    version="v0.1 — DRAFT EXHIBIT (for legal review)",
    extra=("DRAFT FOR LEGAL REVIEW BY BOTH PARTIES' COUNSEL. This exhibit is non-binding until "
           "executed by duly authorized officials of both parties. Nothing herein waives, limits, "
           "or diminishes the sovereign immunity of the Lumbee Tribe of North Carolina. All "
           "financial terms are placeholders pending confirmation."),
)

toc(doc)

# ------------------------------------------------------------------ 1. DRAFT / status
h1(doc, "1. Draft Status & Reservation of Rights")
callout(doc,
    "DRAFT FOR LEGAL REVIEW. This document is a working draft prepared for review by counsel for "
    "both the Lumbee Tribe of North Carolina (the \"Tribe\") and LREMC Technologies, LLC d/b/a RIVR "
    "Tech (\"RIVR Tech\"). It is NOT binding on either party and creates no obligations until it is "
    "reviewed, revised, approved, and executed by duly authorized officials of both parties. "
    "NOTHING IN THIS MEMORANDUM OF UNDERSTANDING WAIVES, LIMITS, OR OTHERWISE DIMINISHES THE "
    "SOVEREIGN IMMUNITY OF THE LUMBEE TRIBE OF NORTH CAROLINA (see Section 12). All names, titles, "
    "dates, dollar amounts, funding commitments, and settlement mechanics shown are placeholders.",
    kind="warn")
para(doc, "Placeholder convention: every term marked \"" + PH + "\" must be resolved by the "
     "parties' counsel before execution. Financial commitments are intentionally left as "
     "placeholders; no binding financial obligation is created by this draft.", italic=True,
     size=9, color=GREY)

# ------------------------------------------------------------------ 2. Background / recitals
h1(doc, "2. Background & Recitals")
para(doc, "This Memorandum of Understanding (\"MOU\") is entered into as of "
     + PH + " (the \"Effective Date\"), by and between:")
bullet(doc, "The Lumbee Tribe of North Carolina, a federally recognized Indian tribe (the \"Tribe\"); and")
bullet(doc, "LREMC Technologies, LLC d/b/a RIVR Tech, a North Carolina limited liability company "
       "with its principal office at 6090 NC Hwy 711, Pembroke, NC 28372 (\"RIVR Tech\").")
para(doc, "The Tribe and RIVR Tech are referred to individually as a \"Party\" and collectively as "
     "the \"Parties.\"")

h2(doc, "2.1 Recitals")
numbered(doc, "WHEREAS, the Lumbee Tribe of North Carolina received full federal recognition "
         "through Public Law 119-60, enacted as Section 8803 of the National Defense Authorization "
         "Act for Fiscal Year 2026, 139 Stat. 1973, signed December 18, 2025 (SR-5.1), and was "
         "added by the U.S. Department of the Interior to the official list of federally recognized "
         "tribes (SR-5.2);")
numbered(doc, "WHEREAS, the Tribe's federal service delivery area for the delivery of federal "
         "services to members comprises Robeson, Cumberland, Hoke, and Scotland Counties, North "
         "Carolina (SR-5.4);")
numbered(doc, "WHEREAS, RIVR Tech operates a symmetrical fiber broadband network in the region and "
         "has an application for designation as an Eligible Telecommunications Carrier (ETC) "
         "pending before the North Carolina Utilities Commission (SR-7.4);")
numbered(doc, "WHEREAS, the Parties wish to collaborate to improve broadband affordability, "
         "adoption, and digital literacy among Tribal members and residents of the service delivery "
         "area; and")
numbered(doc, "WHEREAS, the Parties intend to set out a framework of shared, non-binding "
         "understandings that respects the Tribe's sovereignty, protects member data, and "
         "complies with applicable federal and state law;")
para(doc, "NOW, THEREFORE, the Parties record the following mutual understandings.", bold=True)

callout(doc,
    "SCOPE NOTE. This MOU concerns a broadband AFFORDABILITY initiative and a possible, "
    "Tribe-funded consumer benefit. It is SEPARATE and DISTINCT from federal Lifeline eligibility. "
    "Federal Lifeline support (including any enhanced Tribal benefit) is governed exclusively by "
    "FCC rules and USAC processes and is not created, expanded, or guaranteed by this MOU. See "
    "Section 5.", kind="note")

# ------------------------------------------------------------------ 3. Purpose & principles
h1(doc, "3. Purpose & Guiding Principles")
h2(doc, "3.1 Purpose")
para(doc, "The purpose of this MOU is to establish a cooperative, non-binding framework under which "
     "the Parties may: (a) expand affordable broadband access for Tribal members and area residents; "
     "(b) coordinate outreach, enrollment, and digital-literacy education; (c) protect the privacy "
     "of member and applicant data; and (d) define a possible, separately funded Tribal "
     "affordability contribution, if the Tribe elects to fund one.")
h2(doc, "3.2 Guiding principles")
bullet(doc, "Tribal sovereignty is respected in all activities; the Tribe retains full control over "
       "its name, marks, facilities, membership information, and internal processes.")
bullet(doc, "Regulatory integrity: federal Lifeline eligibility is determined solely under FCC/USAC "
       "rules; nothing in this collaboration is represented to members as changing that.")
bullet(doc, "Data minimization and consent: only the minimum data necessary is shared, only with "
       "consent, and only for the stated purpose.")
bullet(doc, "Separate accounting: any Tribe-funded benefit is tracked separately from federal funds "
       "and grant funds (SR-10.2).")
bullet(doc, "Transparency to members: communications are accurate, non-deceptive, and jointly "
       "reviewed.")

# ------------------------------------------------------------------ 4. Division of responsibilities
h1(doc, "4. Division of Responsibilities")
para(doc, "The following table allocates anticipated responsibilities. \"R\" denotes the responsible "
     "party; \"S\" denotes a supporting role; \"Joint\" denotes shared responsibility. Allocations "
     "marked \"" + PH + "\" require confirmation by the Parties.", size=9.5)
table(doc,
    ["Responsibility", "RIVR Tech", "Lumbee Tribe", "Joint"],
    [
        ["Tribal-member outreach", "S", "R", "Joint messaging"],
        ["Membership verification", "S (receives confirmation only)", "R (Tribe-operated)", "—"],
        ["Community enrollment events", "S (staffing, systems)", "R (venue, member notice)", "Joint calendar"],
        ["Digital-literacy education", "S (materials, trainers)", "R (community delivery)", "Joint curriculum"],
        ["Use of Tribal facilities for enrollment assistance", "Requests access", "R (grants/schedules)", "Joint scheduling"],
        ["Data-sharing restrictions", "R (safeguards, minimization)", "R (roster control)", "Joint DPA " + PH],
        ["Funding of any additional Tribal discount", "S (billing/settlement)", "R (funds, if elected) " + PH, "—"],
        ["Approval of Tribe name/logo use", "Requests approval", "R (written approval)", "—"],
        ["Joint communications", "Drafts", "Reviews/approves", "Joint sign-off"],
        ["Complaint escalation", "R (Tier-1 intake)", "S (member liaison)", "Joint escalation path"],
        ["Monthly reporting", "R (prepares metrics)", "Reviews", "Joint review meeting"],
        ["Audit rights", "Provides records", "R (audit of shared program)", "Joint scope " + PH],
        ["Program evaluation", "S (data)", "S (member input)", "Joint annual review"],
    ],
    widths=[2.5, 1.5, 1.5, 1.4], small=True)
para(doc, "Role letters and any funding role above are non-binding drafting placeholders subject to "
     "counsel review and to the Tribe's election whether to fund a benefit at all.", italic=True,
     size=9, color=GREY)

# ------------------------------------------------------------------ 5. CRITICAL compliance clause
h1(doc, "5. Critical Compliance Clause — Lifeline Eligibility Is Address-Based")
callout(doc,
    "TRIBAL MEMBERSHIP DOES NOT ESTABLISH ENHANCED TRIBAL LIFELINE ELIGIBILITY. The Parties "
    "expressly acknowledge and agree that neither enrollment in the Tribe nor residence within the "
    "four-county service delivery area, by itself, qualifies any person for the enhanced Tribal "
    "Lifeline benefit.", kind="stop")
para(doc, "The Parties acknowledge and agree that:")
bullet(doc, "The enhanced federal Tribal Lifeline benefit requires the subscriber's principal "
       "residence to be physically located on qualifying \"Tribal lands\" as defined at 47 C.F.R. "
       "§54.400(e), and is not conferred by Tribal membership (SR-4.2, SR-5.6).")
bullet(doc, "Qualifying-address status must be verified on a per-applicant basis using USAC's Tribal "
       "Lands Verification Tool; the NLAD Tribal benefit flag is set only for verified addresses "
       "(SR-4.4, SR-5).")
bullet(doc, "As of the date of this draft, the Tribe has no reservation, no land held in trust, and "
       "no FCC §54.412 off-reservation designation; therefore addresses in the four counties are "
       "not presumptively Tribal lands and do not qualify for the enhanced benefit by default "
       "(SR-5.5). " + PH + " (confirm current §54.412 / trust status with USAC/FCC ONAP and BIA).")
bullet(doc, "Any Tribe-funded benefit contemplated by this MOU is a SEPARATE, Tribe-funded "
       "affordability contribution. It is not a federal benefit, is not represented to members as a "
       "federal benefit, and is accounted for separately from federal Lifeline support and grant "
       "funds (SR-10.2).")
para(doc, "Neither Party will represent to any member, applicant, or the public that Tribal "
     "membership or four-county residence establishes enhanced Tribal Lifeline eligibility.",
     bold=True)

# ------------------------------------------------------------------ 6. Privacy-preserving membership verification
h1(doc, "6. Privacy-Preserving Membership Verification")
para(doc, "The Parties intend to confirm Tribal membership for program purposes WITHOUT the Tribe "
     "transferring its full membership roll to RIVR Tech. The specific method will be finalized by "
     "counsel and the Tribe's enrollment office (" + PH + "), and is expected to use one or more of "
     "the following privacy-preserving approaches:")
bullet(doc, "Tribe-operated verification: an applicant presents to the Tribe (or a Tribe-designated "
       "portal), and the Tribe returns to RIVR Tech only a per-applicant yes/no confirmation of "
       "membership.")
bullet(doc, "Applicant attestation with Tribal confirmation: the applicant attests to membership and "
       "the Tribe confirms individual records on a case-by-case basis.")
bullet(doc, "Hashed / tokenized match: matching on a one-way hashed identifier so that no readable "
       "roster is transferred to RIVR Tech.")
h2(doc, "6.1 Data-handling commitments")
bullet(doc, "Data minimization: only the minimum fields necessary for the confirmation are exchanged.")
bullet(doc, "Consent: no member/applicant information is processed without the individual's informed, "
       "documented consent.")
bullet(doc, "Purpose limitation: eligibility and membership data are used ONLY for this program and "
       "NOT for unrelated marketing, resale, or any secondary purpose.")
bullet(doc, "Security & retention: safeguards, breach notification, and retention limits to be set "
       "in a Data Processing Addendum (" + PH + ").")
bullet(doc, "Tribal control: the Tribe retains ownership and control of its membership records at all "
       "times; RIVR Tech receives confirmations, not the roll.")

# ------------------------------------------------------------------ 7. Funding & settlement
h1(doc, "7. Funding & Settlement (Placeholders — No Binding Commitment)")
callout(doc,
    "No binding financial commitment is created by this draft. All amounts, rates, and mechanics in "
    "this Section are placeholders subject to the Tribe's election and to counsel review.",
    kind="warn")
para(doc, "If, and only if, the Tribe elects to fund an additional consumer discount, the Parties "
     "anticipate the following framework:")
numbered(doc, "Discount amount: " + PH + " per eligible household per month, funded solely by the "
         "Tribe.")
numbered(doc, "Invoicing: RIVR Tech invoices the Tribe monthly for the aggregate Tribe-funded "
         "discount actually applied to member bills, with supporting detail (" + PH + " format).")
numbered(doc, "Settlement & reconciliation: the Tribe remits within " + PH + " days; the Parties "
         "reconcile monthly and true-up any discrepancies.")
numbered(doc, "Separate accounting: the Tribe-funded contribution is booked in a separate GL "
         "category and is NOT commingled with (a) federal Lifeline support (USAC), (b) TBCP or other "
         "grant funds, or (c) RIVR-funded discounts (SR-10.2).")
numbered(doc, "Grant-funds restriction: Tribal Broadband Connectivity Program (TBCP) funds are NOT "
         "used as a recurring consumer subsidy unless the specific award and grant terms expressly "
         "authorize that use (SR-10.1).")
numbered(doc, "No offset against federal benefit: any Tribe-funded discount is applied "
         "transparently and does not reduce, replace, or offset the member's federal Lifeline "
         "benefit.")
para(doc, "This Section does not obligate the Tribe to provide any funding and does not obligate "
     "RIVR Tech to provide any service at a particular price.", italic=True, size=9, color=GREY)

# ------------------------------------------------------------------ 8. Name, logo & branding
h1(doc, "8. Name, Logo & Branding Approval")
para(doc, "RIVR Tech will not use the Tribe's name, seal, logo, marks, or any indicia of Tribal "
     "affiliation or endorsement in any material, communication, advertisement, or website without "
     "the Tribe's prior WRITTEN approval, which the Tribe may grant, condition, or withhold in its "
     "sole discretion.")
bullet(doc, "Each proposed use is submitted to the Tribe's designated contact for review (" + PH + " "
       "review turnaround).")
bullet(doc, "The Tribe may revoke approval prospectively; RIVR Tech will cease the affected use "
       "within a reasonable, agreed period.")
bullet(doc, "This Section respects the Tribe's sovereignty and control over its identity and does "
       "not grant RIVR Tech any license except as expressly approved in writing.")

# ------------------------------------------------------------------ 9. Joint comms & complaint escalation
h1(doc, "9. Joint Communications & Complaint Escalation")
h2(doc, "9.1 Joint communications")
bullet(doc, "Member-facing communications about the program are drafted by RIVR Tech and reviewed "
       "and approved by the Tribe before release.")
bullet(doc, "Communications are accurate and non-deceptive, and clearly distinguish federal Lifeline "
       "benefits from any Tribe-funded contribution (see Section 5).")
h2(doc, "9.2 Complaint escalation")
numbered(doc, "Tier 1 — RIVR Tech intake: RIVR Tech logs and acknowledges member complaints within "
         "" + PH + " and attempts first-contact resolution.")
numbered(doc, "Tier 2 — Tribal liaison: unresolved or Tribe-sensitive complaints are escalated to "
         "the Tribe's designated liaison for coordinated handling.")
numbered(doc, "Tier 3 — Joint escalation: material or unresolved matters are escalated to the "
         "Executive Sponsors of both Parties.")
numbered(doc, "Regulatory complaints: nothing in this process limits a member's right to complain "
         "to the FCC, USAC, or the NCUC.")

# ------------------------------------------------------------------ 10. Reporting, audit, evaluation
h1(doc, "10. Reporting, Audit Rights & Program Evaluation")
h2(doc, "10.1 Monthly reporting")
para(doc, "RIVR Tech provides the Tribe a monthly program report including (metrics to be finalized, "
     "" + PH + "):")
bullet(doc, "Members reached via outreach and enrollment events held.")
bullet(doc, "Applications received, approved, and pending (program-level counts, not federal "
       "eligibility determinations).")
bullet(doc, "Households receiving any Tribe-funded discount and aggregate Tribe-funded amount.")
bullet(doc, "Complaints received, resolved, and open; escalations.")
bullet(doc, "Digital-literacy sessions delivered and attendance.")
h2(doc, "10.2 Audit rights")
bullet(doc, "The Tribe may audit records relating to the shared program and to any Tribe-funded "
       "contribution, on reasonable notice, at reasonable times (scope " + PH + ").")
bullet(doc, "Audits are limited to program records and do not extend to RIVR Tech's unrelated "
       "business or to federal Lifeline claims records governed separately by USAC/FCC.")
h2(doc, "10.3 Program evaluation")
para(doc, "The Parties conduct an annual joint evaluation assessing reach, affordability impact, "
     "member satisfaction, and privacy compliance, and use the findings to adjust the program by "
     "mutual agreement.")

# ------------------------------------------------------------------ 11. Term, termination, amendment
h1(doc, "11. Term, Termination & Amendment")
bullet(doc, "Term: this MOU takes effect on the Effective Date and continues for " + PH + " unless "
       "earlier terminated.")
bullet(doc, "Termination for convenience: either Party may terminate on " + PH + " days' written "
       "notice.")
bullet(doc, "Termination for cause: either Party may terminate on material breach uncured within "
       "" + PH + " days of written notice.")
bullet(doc, "Wind-down: on termination, the Parties cooperate on an orderly transition, cease "
       "branding use, and reconcile any outstanding Tribe-funded amounts; data-handling and "
       "confidentiality obligations survive.")
bullet(doc, "Amendment: this MOU may be amended only by a written instrument signed by authorized "
       "officials of both Parties.")

# ------------------------------------------------------------------ 12. Non-waiver + dispute + governing law
h1(doc, "12. Non-Waiver of Sovereign Immunity; Dispute Resolution; Governing Law")
callout(doc,
    "EXPRESS NON-WAIVER OF SOVEREIGN IMMUNITY. Nothing in this MOU constitutes, and this MOU shall "
    "not be construed as, a waiver — express, implied, or by operation of law — of the sovereign "
    "immunity of the Lumbee Tribe of North Carolina, its officials, employees, agents, "
    "instrumentalities, or enterprises. No officer, employee, or agent of the Tribe has authority "
    "to waive the Tribe's sovereign immunity except by an express, written waiver duly authorized "
    "and adopted by the governing body of the Tribe in accordance with Tribal law. No such waiver "
    "is granted, intended, or implied by this MOU. Any limited waiver, if ever granted, would be "
    "set forth only in a separate, express written instrument and would be strictly construed.",
    kind="stop")
h2(doc, "12.1 Dispute resolution")
para(doc, "The Parties will first attempt to resolve any dispute through good-faith negotiation "
     "between their Executive Sponsors, followed by " + PH + " (e.g., non-binding mediation), "
     "without prejudice to the non-waiver provision above. The dispute-resolution mechanism, forum, "
     "and any consent (or absence of consent) to any forum are " + PH + " and subject to the Tribe's "
     "sovereign immunity, which is expressly not waived.")
h2(doc, "12.2 Governing law")
para(doc, "Governing law is " + PH + ". The Parties acknowledge that application of any body of law "
     "does not, by itself, waive the Tribe's sovereign immunity, and any question of Tribal law and "
     "jurisdiction is reserved to the Tribe.")
h2(doc, "12.3 No third-party beneficiaries; independent parties")
bullet(doc, "This MOU creates no rights in any third party, including members or applicants.")
bullet(doc, "The Parties are independent; nothing herein creates a partnership, joint venture, "
       "agency, or employment relationship.")
bullet(doc, "This MOU is non-binding except for any provisions the Parties expressly designate as "
       "binding in the executed version (" + PH + ").")

# ------------------------------------------------------------------ 13. Signature block
h1(doc, "13. Signatures")
callout(doc,
    "This MOU is not effective until executed by DULY AUTHORIZED officials of both Parties. For the "
    "Tribe, execution must be by an authorized Tribal official acting pursuant to Tribal law and "
    "any required action of the Tribe's governing body. Names, titles, and dates below are "
    "placeholders.", kind="note")

para(doc, "LREMC TECHNOLOGIES, LLC d/b/a RIVR TECH", bold=True)
table(doc,
    ["Field", "Entry"],
    [
        ["By (signature)", "________________________________"],
        ["Name", PH],
        ["Title", PH],
        ["Date", PH],
    ],
    widths=[2.0, 4.5])
doc.add_paragraph()
para(doc, "LUMBEE TRIBE OF NORTH CAROLINA", bold=True)
table(doc,
    ["Field", "Entry"],
    [
        ["By (signature)", "________________________________"],
        ["Name", PH],
        ["Title (authorized Tribal official)", PH],
        ["Date", PH],
        ["Governing-body authorization ref.", PH],
    ],
    widths=[2.0, 4.5])
para(doc, "Each signatory represents that they are duly authorized to execute this MOU on behalf of "
     "their respective Party. Tribal execution is subject to Section 12 (non-waiver of sovereign "
     "immunity).", italic=True, size=9, color=GREY)

footer_revhist(doc)

doc.save(PATH)
print("saved", PATH)

# reload
from docx import Document
d2 = Document(PATH)
print("reloaded paragraphs:", len(d2.paragraphs), "tables:", len(d2.tables))
print("OK 11")
