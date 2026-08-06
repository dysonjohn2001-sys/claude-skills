"""Shared python-docx helpers for RIVR Tech Lifeline Word deliverables.

Gives every document the same executive look: navy headings, a cover block with
owner/version/approval, a table-of-contents field, callout boxes, and clean tables.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def new_doc():
    doc = Document()
    # base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for h, sz, col in [("Heading 1", 16, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, BLUE)]:
        st = doc.styles[h]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.color.rgb = col
        st.font.bold = True
    return doc

def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def cover(doc, title, subtitle, deliverable_no, owner, approvers=None, version="v0.1 — DRAFT",
          extra=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("LREMC Technologies, LLC d/b/a RIVR Tech")
    r.font.size = Pt(11); r.font.color.rgb = GREY; r.bold = True
    t = doc.add_paragraph()
    tr = t.add_run(title); tr.font.size = Pt(24); tr.font.bold = True; tr.font.color.rgb = NAVY
    if subtitle:
        s = doc.add_paragraph(); sr = s.add_run(subtitle)
        sr.font.size = Pt(13); sr.font.color.rgb = BLUE
    doc.add_paragraph()
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light List Accent 1"
    rows = [
        ("Deliverable", deliverable_no),
        ("Document owner", owner),
        ("Approval / sign-off", approvers or "________________________  (name / title / date)"),
        ("Version", version),
        ("Status", "PRE-DECISION build artifact. Figures are illustrative/placeholder pending confirmation."),
        ("Source basis", "See 00_Regulatory_Source_Register.md (item refs 'SR-x.y')."),
    ]
    for k, v in rows:
        c = meta.add_row().cells
        c[0].text = k; c[1].text = v
        for run in c[0].paragraphs[0].runs: run.bold = True
    if extra:
        doc.add_paragraph()
        callout(doc, extra, kind="warn")
    doc.add_page_break()

def toc(doc, heading="Table of Contents"):
    doc.add_heading(heading, level=1)
    par = doc.add_paragraph()
    run = par.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    r2 = OxmlElement("w:r"); t = OxmlElement("w:t")
    t.text = "Right-click and 'Update Field' to build the table of contents."
    r2.append(t); fld.append(r2)
    par._p.append(fld)
    doc.add_page_break()

def callout(doc, text, kind="note"):
    colors = {"warn": ("FCE4D6", RED, "⚠ "), "note": ("D9E1F2", BLUE, "ⓘ "),
              "ok": ("E2EFDA", RGBColor(0x37,0x6E,0x37), "✔ "),
              "stop": ("F8CBAD", RED, "⛔ ")}
    fill, col, icon = colors.get(kind, colors["note"])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    _shade(cell, fill)
    p = cell.paragraphs[0]
    run = p.add_run(icon + text)
    run.font.size = Pt(10); run.font.color.rgb = col
    if kind in ("warn", "stop"): run.bold = True
    return tbl

def table(doc, headers, rows, widths=None, header_fill="2E5496", small=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]; run = p.add_run(str(h))
        run.bold = True; run.font.color.rgb = WHITE; run.font.size = Pt(9 if small else 9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]; run = p.add_run("" if val is None else str(val))
            run.font.size = Pt(8.5 if small else 9.5)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Inches(w)
    return t

def h1(doc, text): return doc.add_heading(text, level=1)
def h2(doc, text): return doc.add_heading(text, level=2)
def h3(doc, text): return doc.add_heading(text, level=3)
def para(doc, text, bold=False, italic=False, size=10.5, color=None):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p
def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text); return p
def numbered(doc, text):
    p = doc.add_paragraph(style="List Number"); p.add_run(text); return p

def footer_revhist(doc):
    doc.add_page_break()
    h2(doc, "Revision history")
    table(doc, ["Version", "Date", "Author", "Summary of change", "Approved by"],
          [["v0.1", "(build)", "Program build", "Initial draft for internal review", "(pending)"],
           ["", "", "", "", ""], ["", "", "", "", ""]])
    doc.add_paragraph()
    para(doc, "This is a build-phase draft. All figures, prices, and regulatory determinations are "
        "subject to confirmation per the Open Issues & Legal Decisions Log (Deliverable 13) and the "
        "Regulatory Source Register (Deliverable 00). Do not launch, market, or bill against this "
        "document until the final NCUC ETC order is issued and USAC onboarding is complete.",
        italic=True, size=9, color=GREY)
