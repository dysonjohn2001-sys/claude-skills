"""Shared openpyxl styling helpers for the RIVR Tech Lifeline deliverables.

Color / cell conventions used across every workbook:
  * INPUT (yellow)          -> user-editable assumption cell
  * PLACEHOLDER (orange)    -> illustrative value that MUST be replaced with a
                               confirmed RIVR Tech / USAC figure before use
  * CONSTANT (blue)         -> federal program constant to be re-verified each
                               program year against the cited source
  * CALC (no fill)          -> formula-driven, do not hand-edit
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

# --- palette ---------------------------------------------------------------
NAVY = "1F3864"
BLUE = "2E5496"
LIGHT_BLUE = "D9E1F2"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
GREEN = "E2EFDA"
GREY = "F2F2F2"
RED = "C00000"
WHITE = "FFFFFF"

THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX = Border(left=MED, right=MED, top=MED, bottom=MED)

def title(ws, cell, text, size=16):
    ws[cell] = text
    ws[cell].font = Font(bold=True, size=size, color=NAVY)

def subtitle(ws, cell, text, size=11):
    ws[cell] = text
    ws[cell].font = Font(italic=True, size=size, color="595959")

def header_row(ws, row, headers, start_col=1, fill=BLUE):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

def cell(ws, coord, value=None, kind="calc", num=None, bold=False, wrap=False,
         align=None, color=None, size=10):
    c = ws[coord]
    if value is not None:
        c.value = value
    fills = {"input": YELLOW, "placeholder": ORANGE, "constant": LIGHT_BLUE,
             "calc": None, "good": GREEN, "grey": GREY}
    fill = fills.get(kind)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.font = Font(bold=bold, size=size, color=(RED if kind == "placeholder" else "000000"))
    if num:
        c.number_format = num
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    c.border = BORDER
    return c

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

USD = '"$"#,##0.00'
USD0 = '"$"#,##0'
PCT = '0.0%'
NUM = '#,##0'

def legend(ws, row, col=1):
    items = [("INPUT — edit me", YELLOW), ("PLACEHOLDER — replace w/ confirmed figure", ORANGE),
             ("PROGRAM CONSTANT — re-verify each year", LIGHT_BLUE), ("CALCULATED — do not edit", WHITE)]
    ws.cell(row=row, column=col, value="Legend:").font = Font(bold=True, size=9)
    for i, (txt, clr) in enumerate(items):
        c = ws.cell(row=row + 1 + i, column=col, value=txt)
        c.fill = PatternFill("solid", fgColor=clr)
        c.font = Font(size=9)
        c.border = BORDER
