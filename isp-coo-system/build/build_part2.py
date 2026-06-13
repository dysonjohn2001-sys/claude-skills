"""
Board Report + Data Dictionary + User Guide + Assumptions tabs.
Called from build_workbook.py with shared context.
"""
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment


def add_board_and_reference(wb, S, D, ctx):
    PM = ctx["PM"]; PIDX = ctx["PIDX"]; rng = ctx["rng"]
    sumifs = ctx["sumifs"]; countifs = ctx["countifs"]
    apply_ryg = ctx["apply_ryg_word"]
    DS, DE = ctx["DATA_START"], ctx["DATA_END"]
    CUR, NUM, PCT, DEC1 = ctx["CUR"], ctx["NUM"], ctx["PCT"], ctx["DEC1"]
    FD, CD, OD, ST = ctx["FD"], ctx["CD"], ctx["OD"], ctx["ST"]
    SD, PL, RD = ctx["SD"], ctx["PL"], ctx["RD"]
    PRJ, GR, RR, BA = ctx["PRJ"], ctx["GR"], ctx["RR"], ctx["BA"]
    prj_last = ctx["prj_last"]

    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ===================================================================
    # BOARD REPORT
    # ===================================================================
    br = wb.create_sheet("Board_Report")
    S.page_setup(br, tab_color="C00000")
    for c in range(1, 9):
        br.column_dimensions[get_column_letter(c)].width = 17
    br.column_dimensions["A"].width = 3

    br.merge_cells("A1:H2")
    t = br["A1"]; t.value = "QUARTERLY BOARD REPORT  /  MONTHLY COO REPORT"
    t.font = S.F_TITLE; t.fill = S.fill(S.NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    br.merge_cells("A3:H3")
    s = br["A3"]
    s.value = ('="Reporting Period: "&Exec_Dashboard!C5&"   |   Regional Fiber Broadband   |   '
               'Prepared by the COO   |   CONFIDENTIAL — Board of Directors"')
    s.font = S.F_SUBTITLE; s.fill = S.fill(S.SLATE)
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    def sec(row, text):
        br.merge_cells(f"A{row}:H{row}")
        c = br.cell(row=row, column=1, value=text)
        c.font = S.F_SECTION; c.fill = S.fill(S.SLATE)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        br.row_dimensions[row].height = 20
        return row + 1

    # shared metric fragments
    tot = sumifs(CD, "C", ("A", PM)); new = sumifs(CD, "D", ("A", PM))
    dis = sumifs(CD, "E", ("A", PM)); homes = sumifs(CD, "F", ("A", PM))
    prior = sumifs(CD, "C", ("B", f"{PIDX}-1"))
    rev = sumifs(FD, "C", ("A", PM)); mrr = sumifs(FD, "D", ("A", PM))
    budrev = sumifs(FD, "J", ("A", PM)); cash = sumifs(FD, "N", ("A", PM))
    capex = sumifs(FD, "E", ("A", PM)); opex = sumifs(FD, "F", ("A", PM))
    ytd_rev = f'SUMIFS({rng(FD,"C")},{rng(FD,"B")},"<="&{PIDX})'
    ytd_grant = f'SUMIFS({rng(FD,"I")},{rng(FD,"B")},"<="&{PIDX})'
    backlog = f'SUMIFS({rng(OD,"F")},{rng(OD,"B")},"<="&{PIDX})-SUMIFS({rng(OD,"G")},{rng(OD,"B")},"<="&{PIDX})'
    adays = f'AVERAGEIFS({rng(OD,"H")},{rng(OD,"A")},{PM})'
    behind = countifs(PRJ, ("F", '"Behind"')); atrisk = countifs(PRJ, ("F", '"At Risk"'))
    active = ("(" + countifs(PRJ, ("F", '"On Track"')) + "+" +
              countifs(PRJ, ("F", '"At Risk"')) + "+" +
              countifs(PRJ, ("F", '"Behind"')) + ")")
    redrisk = f'COUNTIFS({rng(RR,"F")},">=15")'
    decisions = countifs(BA, ("I", '"Yes"'), ("H", '"<>Complete"'))

    row = 5
    # Executive summary narrative (auto-generated)
    row = sec(row, "1.  EXECUTIVE SUMMARY")
    br.merge_cells(f"A{row}:H{row+3}")
    es = br.cell(row=row, column=1)
    es.value = ('="For the period ending "&Exec_Dashboard!C5&", the company served "&TEXT(' + tot + ',"#,##0")&'
                '" customers across the service area, a net change of "&TEXT(' + new + '-' + dis + ',"#,##0")&'
                '" connections ("&TEXT(' + new + ',"#,##0")&" new sales less "&TEXT(' + dis + ',"#,##0")&" disconnects, '
                'churn "&TEXT(IF(' + prior + '=0,0,' + dis + '/' + prior + '),"0.0%")&"). '
                'Monthly revenue was "&TEXT(' + rev + ',"$#,##0")&" against a budget of "&TEXT(' + budrev + ',"$#,##0")&'
                '" (variance "&TEXT(' + rev + '-' + budrev + ',"$#,##0")&"). Homes passed reached "&TEXT(' + homes + ',"#,##0")&". '
                'Of "&' + active + '&" active capital projects, "&' + behind + '&" are behind schedule and "&' + atrisk + '&" at risk. '
                'Installation backlog stands at "&TEXT(' + backlog + ',"#,##0")&" orders with an average install time of "&'
                'TEXT(' + adays + ',"0.0")&" days. End-of-month cash was "&TEXT(' + cash + ',"$#,##0")&". '
                '"&' + redrisk + '&" risk(s) are red and "&' + decisions + '&" decision(s) require board attention."')
    es.font = S.F_BODY; es.alignment = LEFT_TOP; es.fill = S.fill(S.LIGHT_BG)
    for rr in range(row, row + 4):
        for cc in range(1, 9):
            br.cell(row=rr, column=cc).border = S.box_thin
    row += 5

    # Going well / behind / needs attention (3 columns)
    row = sec(row, "2.  WHAT IS GOING WELL  /  WHAT IS BEHIND  /  WHAT NEEDS EXECUTIVE ATTENTION")
    heads = ["✅ Going Well", "⚠️ Behind", "🚩 Needs Executive Attention"]
    spans = [("A", "C"), ("D", "E"), ("F", "H")]
    for (h, (c1, c2)) in zip(heads, spans):
        br.merge_cells(f"{c1}{row}:{c2}{row}")
        cell = br[f"{c1}{row}"]; cell.value = h; cell.font = S.F_BODY_B
        cell.fill = S.fill("E8EDF5"); cell.alignment = CENTER; cell.border = S.box_thin
    going = ["Pembroke Core Ring & Lumberton MDU on/ahead of plan",
             "ARPU trending up; margin healthy on residential & business",
             "BEAD Zone 1 drawdowns paid in full; strong grant execution",
             "Avg days-to-install improving month over month"]
    behindl = ["Scotland RDOF build behind — grant drawdown at risk",
               "Columbus BEAD Zone 2 design rework slipping Milestone 1",
               "Install backlog rising vs crew capacity in Robeson",
               "Make-ready / pole-attachment delays from utility"]
    attn = ["Approve revolver draw to bridge BEAD reimbursement timing",
            "Authorize 2nd construction contractor (Scotland + Columbus)",
            "Approve install-crew headcount add to clear backlog",
            "Ratify HOA / bulk pricing policy"]
    for i in range(4):
        rr = row + 1 + i
        br.merge_cells(f"A{rr}:C{rr}"); br.merge_cells(f"D{rr}:E{rr}"); br.merge_cells(f"F{rr}:H{rr}")
        br.cell(row=rr, column=1, value="• " + going[i]).font = S.F_SMALL
        br.cell(row=rr, column=1).alignment = LEFT
        br.cell(row=rr, column=4, value="• " + behindl[i]).font = S.F_SMALL
        br.cell(row=rr, column=4).alignment = LEFT
        br.cell(row=rr, column=6, value="• " + attn[i]).font = S.F_SMALL
        br.cell(row=rr, column=6).alignment = LEFT
        for cc in range(1, 9):
            br.cell(row=rr, column=cc).border = S.box_thin
    row += 6

    # Financial summary table
    row = sec(row, "3.  FINANCIAL PERFORMANCE")
    fin_rows = [
        ("Revenue (MTD)", f"={rev}", CUR, f"={budrev}", f"={rev}-{budrev}"),
        ("Revenue (YTD)", f"={ytd_rev}", CUR, "", ""),
        ("MRR", f"={mrr}", CUR, "", ""),
        ("Capital Spend (MTD)", f"={capex}", CUR, f'=SUMIFS({rng(FD,"L")},{rng(FD,"A")},{PM})', f'={capex}-SUMIFS({rng(FD,"L")},{rng(FD,"A")},{PM})'),
        ("Operating Expense (MTD)", f"={opex}", CUR, f'=SUMIFS({rng(FD,"K")},{rng(FD,"A")},{PM})', f'={opex}-SUMIFS({rng(FD,"K")},{rng(FD,"A")},{PM})'),
        ("Grant Reimb (YTD)", f"={ytd_grant}", CUR, "", ""),
        ("Cash (End of Month)", f"={cash}", CUR, "", ""),
    ]
    for j, h in enumerate(["Metric", "Actual", "", "Budget", "Variance"]):
        cc = br.cell(row=row, column=1 + j + (1 if j >= 2 else 0) if False else 1 + j, value=h)
    # simpler header
    for j, h in enumerate(["Metric", "Actual", "Budget", "Variance"]):
        cell = br.cell(row=row, column=1 + (j if j == 0 else j + 1), value=h)
        cell.font = S.F_HDR; cell.fill = S.fill(S.TABLE_HDR); cell.alignment = CENTER
        cell.border = S.box_thin
    br.merge_cells(f"A{row}:B{row}")
    row += 1
    for name, actf, fmt, budf, varf in fin_rows:
        br.merge_cells(f"A{row}:B{row}")
        br.cell(row=row, column=1, value=name).font = S.F_BODY_B
        a = br.cell(row=row, column=3, value=actf); a.number_format = fmt
        if budf:
            b = br.cell(row=row, column=4, value=budf); b.number_format = fmt
            v = br.cell(row=row, column=5, value=varf); v.number_format = fmt
        for cc in range(1, 6):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1
    row += 1

    # Customer growth + operations summary side by side tables
    row = sec(row, "4.  CUSTOMER GROWTH  &  OPERATIONAL PERFORMANCE")
    cust_metrics = [
        ("Total Customers", f"={tot}", NUM),
        ("New Sales (MTD)", f"={new}", NUM),
        ("Disconnects (MTD)", f"={dis}", NUM),
        ("Net Adds (MTD)", f"={new}-{dis}", NUM),
        ("Churn Rate", f"=IF({prior}=0,0,{dis}/{prior})", PCT),
        ("Homes Passed", f"={homes}", NUM),
        ("Take Rate", f"=IF({homes}=0,0,{tot}/{homes})", PCT),
    ]
    ops_metrics = [
        ("Installs Completed (MTD)", f'=SUMIFS({rng(OD,"G")},{rng(OD,"A")},{PM})', NUM),
        ("Install Backlog", f"={backlog}", NUM),
        ("Avg Days to Install", f"={adays}", DEC1),
        ("Trouble Tickets (MTD)", f'=SUMIFS({rng(OD,"I")},{rng(OD,"A")},{PM})', NUM),
        ("Outages (MTD)", f'=SUMIFS({rng(OD,"K")},{rng(OD,"A")},{PM})', "0"),
        ("Active Projects", f"={active}", "0"),
        ("Projects Behind / At Risk", f"={behind}+{atrisk}", "0"),
    ]
    br.cell(row=row, column=1, value="Customer").font = S.F_HDR
    br.cell(row=row, column=1).fill = S.fill(S.TABLE_HDR); br.merge_cells(f"A{row}:B{row}")
    br.cell(row=row, column=1).alignment = CENTER
    br.cell(row=row, column=3, value="Value").font = S.F_HDR
    br.cell(row=row, column=3).fill = S.fill(S.TABLE_HDR); br.cell(row=row, column=3).alignment = CENTER
    br.cell(row=row, column=5, value="Operations").font = S.F_HDR
    br.cell(row=row, column=5).fill = S.fill(S.TABLE_HDR); br.merge_cells(f"E{row}:G{row}")
    br.cell(row=row, column=5).alignment = CENTER
    br.cell(row=row, column=8, value="Value").font = S.F_HDR
    br.cell(row=row, column=8).fill = S.fill(S.TABLE_HDR); br.cell(row=row, column=8).alignment = CENTER
    for cc in range(1, 9):
        br.cell(row=row, column=cc).border = S.box_thin
    row += 1
    for i in range(7):
        br.merge_cells(f"A{row}:B{row}")
        br.cell(row=row, column=1, value=cust_metrics[i][0]).font = S.F_BODY
        cm = br.cell(row=row, column=3, value=cust_metrics[i][1]); cm.number_format = cust_metrics[i][2]
        br.merge_cells(f"E{row}:G{row}")
        br.cell(row=row, column=5, value=ops_metrics[i][0]).font = S.F_BODY
        om = br.cell(row=row, column=8, value=ops_metrics[i][1]); om.number_format = ops_metrics[i][2]
        for cc in range(1, 9):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1
    row += 1

    # Major project updates
    row = sec(row, "5.  MAJOR CAPITAL PROJECT UPDATES")
    for j, h in enumerate(["Project", "County", "Status", "% Complete", "Capital Spent", "Blocker / Note"]):
        spanmap = {0: ("A", "B"), 5: ("F", "H")}
        if j == 0:
            br.merge_cells(f"A{row}:B{row}"); col = 1
        elif j == 1:
            col = 3
        elif j == 2:
            col = 4
        elif j == 3:
            col = 5
        elif j == 4:
            col = 6 if False else 6
        cmap = {0: 1, 1: 3, 2: 4, 3: 5, 4: 6, 5: 7}
        cell = br.cell(row=row, column=cmap[j], value=h)
        cell.font = S.F_HDR; cell.fill = S.fill(S.TABLE_HDR); cell.alignment = CENTER
    br.merge_cells(f"G{row}:H{row}")
    for cc in range(1, 9):
        br.cell(row=row, column=cc).border = S.box_thin
    row += 1
    n_proj = prj_last - DS + 1
    for i in range(n_proj):
        dr = DS + i
        br.merge_cells(f"A{row}:B{row}")
        br.cell(row=row, column=1, value=f"=Project_Data!A{dr}").font = S.F_SMALL
        br.cell(row=row, column=3, value=f"=Project_Data!B{dr}").font = S.F_SMALL
        sc = br.cell(row=row, column=4, value=f"=Project_Data!F{dr}"); sc.font = S.F_SMALL; sc.alignment = CENTER
        pcc = br.cell(row=row, column=5, value=f"=Project_Data!E{dr}"); pcc.number_format = "0%"; pcc.alignment = CENTER
        spc = br.cell(row=row, column=6, value=f"=Project_Data!L{dr}"); spc.number_format = CUR
        br.merge_cells(f"G{row}:H{row}")
        br.cell(row=row, column=7, value=f"=Project_Data!R{dr}").font = S.F_SMALL
        br.cell(row=row, column=7).alignment = LEFT
        for cc in range(1, 9):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1
    apply_ryg(br, f"D{row-n_proj}:D{row-1}")
    # colour status words in board project table
    from openpyxl.formatting.rule import CellIsRule
    for word, (bg, tx) in {"On Track": (S.GREEN_BG, S.GREEN_TX),
                            "At Risk": (S.YELLOW_BG, S.YELLOW_TX),
                            "Behind": (S.RED_BG, S.RED_TX),
                            "Complete": ("D9E1F2", "1F3864")}.items():
        br.conditional_formatting.add(f"D{row-n_proj}:D{row-1}", CellIsRule(
            operator="equal", formula=[f'"{word}"'], fill=S.fill(bg),
            font=S.Font(bold=True, color=tx)))
    row += 1

    # Grant / capital project summary
    row = sec(row, "6.  GRANT / CAPITAL PROJECT SUMMARY")
    grant_rows = [
        ("Total Approved (all programs)", f'=SUM({rng(GR,"E")})'),
        ("Total Received to Date", f'=SUM({rng(GR,"F")})'),
        ("Outstanding (Approved − Received)", f'=SUM({rng(GR,"E")})-SUM({rng(GR,"F")})'),
        ("In Review / Submitted (Requested)", f'=SUMIFS({rng(GR,"D")},{rng(GR,"G")},"Under Review")+SUMIFS({rng(GR,"D")},{rng(GR,"G")},"Submitted")'),
        ("Total Capital Budget (projects)", f'=SUM({rng(PRJ,"K")})'),
        ("Total Capital Spent (projects)", f'=SUM({rng(PRJ,"L")})'),
    ]
    for name, f in grant_rows:
        br.merge_cells(f"A{row}:E{row}")
        br.cell(row=row, column=1, value=name).font = S.F_BODY_B
        v = br.cell(row=row, column=6, value=f); v.number_format = CUR
        br.merge_cells(f"F{row}:H{row}")
        for cc in range(1, 9):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1
    row += 1

    # Key risks
    row = sec(row, "7.  KEY RISKS  (top by score)")
    for j, h in enumerate(["Risk", "Cat.", "Score", "Owner", "Mitigation"]):
        cmap = {0: 1, 1: 4, 2: 5, 3: 6, 4: 7}
        cell = br.cell(row=row, column=cmap[j], value=h)
        cell.font = S.F_HDR; cell.fill = S.fill(S.TABLE_HDR); cell.alignment = CENTER
    br.merge_cells(f"A{row}:C{row}"); br.merge_cells(f"G{row}:H{row}")
    for cc in range(1, 9):
        br.cell(row=row, column=cc).border = S.box_thin
    row += 1
    for k in range(1, 6):
        large = f'LARGE(Risk_Register!$F$4:$F$1003,{k})'
        mr = f'MATCH({large},Risk_Register!$F$4:$F$1003,0)'
        br.merge_cells(f"A{row}:C{row}")
        br.cell(row=row, column=1, value=f'=IFERROR(INDEX(Risk_Register!$B$4:$B$1003,{mr}),"")').font = S.F_SMALL
        br.cell(row=row, column=1).alignment = LEFT
        br.cell(row=row, column=4, value=f'=IFERROR(INDEX(Risk_Register!$C$4:$C$1003,{mr}),"")').font = S.F_SMALL
        sc = br.cell(row=row, column=5, value=f'=IFERROR({large},"")'); sc.alignment = CENTER
        br.cell(row=row, column=6, value=f'=IFERROR(INDEX(Risk_Register!$H$4:$H$1003,{mr}),"")').font = S.F_SMALL
        br.merge_cells(f"G{row}:H{row}")
        br.cell(row=row, column=7, value=f'=IFERROR(INDEX(Risk_Register!$I$4:$I$1003,{mr}),"")').font = S.F_SMALL
        br.cell(row=row, column=7).alignment = LEFT
        for cc in range(1, 9):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1
    row += 1

    # Decisions needed from board
    row = sec(row, "8.  DECISIONS NEEDED FROM THE BOARD")
    for j, h in enumerate(["Decision / Action", "Owner", "Priority", "Due"]):
        cmap = {0: 1, 1: 5, 2: 6, 3: 7}
        cell = br.cell(row=row, column=cmap[j], value=h)
        cell.font = S.F_HDR; cell.fill = S.fill(S.TABLE_HDR); cell.alignment = CENTER
    br.merge_cells(f"A{row}:D{row}"); br.merge_cells(f"G{row}:H{row}")
    for cc in range(1, 9):
        br.cell(row=row, column=cc).border = S.box_thin
    row += 1
    for k in range(1, 7):
        idx = (f'AGGREGATE(15,6,(ROW(Board_Actions!$I$4:$I$1003)-3)/'
               f'((Board_Actions!$I$4:$I$1003="Yes")*(Board_Actions!$H$4:$H$1003<>"Complete")),{k})')
        br.merge_cells(f"A{row}:D{row}")
        br.cell(row=row, column=1, value=f'=IFERROR(INDEX(Board_Actions!$B$4:$B$1003,{idx}),"")').font = S.F_SMALL
        br.cell(row=row, column=1).alignment = LEFT
        br.cell(row=row, column=5, value=f'=IFERROR(INDEX(Board_Actions!$D$4:$D$1003,{idx}),"")').font = S.F_SMALL
        br.cell(row=row, column=6, value=f'=IFERROR(INDEX(Board_Actions!$G$4:$G$1003,{idx}),"")').alignment = CENTER
        br.merge_cells(f"G{row}:H{row}")
        br.cell(row=row, column=7, value=f'=IFERROR(INDEX(Board_Actions!$F$4:$F$1003,{idx}),"")').alignment = CENTER
        for cc in range(1, 9):
            br.cell(row=row, column=cc).border = S.box_thin
        row += 1

    br.freeze_panes = "A4"

    # ===================================================================
    # DATA DICTIONARY
    # ===================================================================
    dd = wb.create_sheet("Data_Dictionary")
    S.page_setup(dd, tab_color="595959")
    dd.column_dimensions["A"].width = 3
    for col, w in {"B": 26, "C": 34, "D": 48, "E": 26}.items():
        dd.column_dimensions[col].width = w
    dd.merge_cells("B1:E1")
    t = dd["B1"]; t.value = "DATA DICTIONARY  —  KPI Definitions & Formulas"
    t.font = S.F_SECTION; t.fill = S.fill(S.NAVY); t.alignment = LEFT
    dd.row_dimensions[1].height = 24
    headers = ["KPI / Term", "Definition", "How it is calculated", "Source tab"]
    for j, h in enumerate(headers):
        c = dd.cell(row=3, column=2 + j, value=h)
        c.font = S.F_HDR; c.fill = S.fill(S.TABLE_HDR); c.alignment = CENTER; c.border = S.box_thin
    entries = [
        ("Total Customers", "End-of-month active subscriber count", "Latest Customer_Data row for the month", "Customer_Data"),
        ("New Sales", "Gross new connections in the month", "Sum of New Sales for the month", "Customer_Data"),
        ("Disconnects", "Subscribers lost in the month", "Sum of Disconnects for the month", "Customer_Data"),
        ("Net Adds", "Net subscriber change", "New Sales − Disconnects", "Customer_Data"),
        ("Churn Rate", "Monthly subscriber churn", "Disconnects ÷ prior-month Total Customers", "Customer_Data"),
        ("Homes Passed", "Cumulative premises where fiber is available", "Latest cumulative value", "Customer_Data"),
        ("Take Rate", "Penetration of homes passed", "Total Customers ÷ Homes Passed", "Customer + Customer"),
        ("ARPU", "Average revenue per user", "MRR ÷ Total Customers", "Financial + Customer"),
        ("MRR", "Monthly recurring revenue", "Entered monthly", "Financial_Data"),
        ("Revenue", "Total billed revenue in month", "Entered monthly (recurring + non-recurring)", "Financial_Data"),
        ("CapEx", "Capitalized construction spend", "Entered monthly", "Financial_Data"),
        ("OpEx", "Operating expense (excl. capex)", "Entered monthly", "Financial_Data"),
        ("Budget vs Actual", "Variance to plan", "Actual − Budget (Revenue/OpEx/CapEx)", "Financial_Data"),
        ("Cash Forecast (90-day)", "Projected cash 3 months out", "Cash + 3×Forecast Rev + Grant Outstanding − 3×(OpEx+CapEx+Contractor)", "Financial + Grant"),
        ("Cost per Install", "Direct cost to install a subscriber", "Contractor cost ÷ Installs Completed", "Financial + Operations"),
        ("Cost per Home Passed", "Build cost efficiency", "CapEx ÷ Homes Passed Added", "Financial + Operations"),
        ("Gross Margin %", "Service-line profitability", "(Revenue − COGS) ÷ Revenue", "Revenue_Detail / Financial"),
        ("Homes Passed", "Premises fiber is available to (ISP term)", "Cumulative construction output", "Operations / Customer"),
        ("Footage Built", "Linear feet of fiber constructed", "Entered monthly", "Operations_Data"),
        ("Drops", "Fiber drop from street to premises", "Entered monthly (Drops Completed)", "Operations_Data"),
        ("Installs / Activations", "Premises connected & turned up (ONT live)", "Entered monthly", "Operations_Data"),
        ("Avg Days to Install", "Order-to-activation cycle time", "Average for the month", "Operations_Data"),
        ("Install Backlog", "Open install orders", "YTD Scheduled − YTD Completed", "Operations_Data"),
        ("Trouble Tickets", "Customer service issues opened", "Sum for the month", "Operations_Data"),
        ("Repeat Tickets", "Re-opened within 30 days (quality signal)", "Sum for the month; % of total", "Operations_Data"),
        ("Splicing", "Fiber splice work (make-ready to active)", "Splices Completed entered monthly", "Operations_Data"),
        ("Make-Ready", "Pole/conduit prep before attach", "Phase status per project", "Project_Data"),
        ("Permitting", "Regulatory permits to build", "Phase status per project", "Project_Data"),
        ("Contractor Performance", "Build-partner quality score 0–10", "Average score by contractor", "Project_Data"),
        ("Pipeline Value", "Open opportunity ARR", "Sum of open (non-closed) opp Value", "Pipeline"),
        ("Conversion Rate", "Lead-to-connect efficiency", "Sales ÷ Leads", "Sales_Data"),
        ("Enterprise / DIA", "Dedicated Internet Access deals", "Opp Type = Enterprise/DIA", "Pipeline"),
        ("HOA / Bulk", "Developer & multi-dwelling agreements", "Opp Type = HOA/Developer or Bulk", "Pipeline"),
        ("Grant Reimbursement", "BEAD/CAB/ARPA/State/RDOF drawdowns", "Approved & Received amounts", "Grant_Data"),
        ("BEAD", "Broadband Equity Access & Deployment (federal)", "Grant program tag", "Grant_Data"),
        ("CAB", "Completing Access to Broadband (NC state)", "Grant program tag", "Grant_Data"),
        ("Risk Score", "Likelihood × Impact (1–25)", "L × I; RED ≥15, YELLOW ≥8", "Risk_Register"),
        ("RYG Health", "Red/Yellow/Green status", "Threshold rules per KPI (see card formulas)", "All dashboards"),
    ]
    r = 4
    for e in entries:
        for j, v in enumerate(e):
            c = dd.cell(row=r, column=2 + j, value=v)
            c.font = S.F_SMALL if j != 0 else S.F_BODY_B
            c.alignment = LEFT_TOP; c.border = S.box_thin
        r += 1
    dd.freeze_panes = "B4"

    # ===================================================================
    # USER GUIDE
    # ===================================================================
    ug = wb.create_sheet("User_Guide")
    S.page_setup(ug, tab_color="375623")
    ug.column_dimensions["A"].width = 3
    ug.column_dimensions["B"].width = 110
    ug.merge_cells("B1:B1")
    ug["B1"] = "USER GUIDE & UPDATE PROCESS"
    ug["B1"].font = S.F_SECTION; ug["B1"].fill = S.fill(S.NAVY)
    ug["B1"].alignment = LEFT; ug.row_dimensions[1].height = 24

    guide_blocks = [
        ("HOW THIS WORKBOOK IS ORGANIZED", [
            "• DASHBOARDS (read-only views): Exec_Dashboard, Sales_Dashboard, Ops_Dashboard, Financial_Dashboard, Board_Report.",
            "• INPUT TABS (where staff type): Customer_Data, Sales_Data, Pipeline, Operations_Data, Project_Data,",
            "   Financial_Data, Revenue_Detail, Grant_Data, Sales_Targets, Risk_Register, Board_Action Items.",
            "• REFERENCE: Data_Dictionary (every KPI defined), Assumptions, User_Guide. Calc & Lists are hidden engine tabs.",
            "• Everything on the dashboards is a formula that reads the input tabs. Type data in, dashboards update automatically.",
        ]),
        ("THE ONE CONTROL YOU SET — REPORTING MONTH", [
            "• On Exec_Dashboard, cell C5 ('Reporting Month') has a dropdown. Pick the month you are reporting on.",
            "• Every dashboard and the Board Report follow that selection for This-Month / MTD / YTD / Rolling-12 math.",
        ]),
        ("WEEKLY UPDATE PROCESS  (≈20 minutes)", [
            "1. Customer_Data — confirm/adjust the current month row (Total Customers, New Sales, Disconnects, Homes Passed).",
            "2. Sales_Data — add rows for connects closed this week (Month, County, Segment, Channel, Campaign, Rep, units, New MRR).",
            "3. Pipeline — update opportunity Stages; mark Closed-Won / Closed-Lost (add a Lost Reason for losses).",
            "4. Operations_Data — update installs scheduled/completed, drops, splices, tickets, outages, avg days to install.",
            "5. Project_Data — update % Complete, Status, phase statuses, capital spent, and any new Blocker.",
            "6. Risk_Register — review scores, owners, mitigations; close resolved risks; add new ones.",
            "7. Board_Action Items — log any new decisions needed; update statuses.",
            "8. Open Exec_Dashboard — confirm RYG health and the auto-generated summary read correctly.",
        ]),
        ("MONTHLY COO REPORTING PROCESS  (≈45 minutes)", [
            "1. Complete the month in Financial_Data (Revenue, MRR, CapEx, OpEx, Labor, Contractor, Grant, Budget, Cash).",
            "2. Complete Revenue_Detail rows for the month (revenue & COGS by product) so margin-by-line is accurate.",
            "3. Update Grant_Data with drawdown status changes (Submitted / Approved / Partially Paid / Paid).",
            "4. Set Exec_Dashboard C5 to the closed month. Review all four dashboards.",
            "5. Open Board_Report — Section 1 (Executive Summary) and the financial/customer/ops tables auto-populate.",
            "6. Edit the qualitative cells in Section 2 (Going Well / Behind / Needs Attention) to reflect the month's narrative.",
            "7. Print/PDF the Board_Report tab for distribution (set print area to the Board_Report sheet).",
        ]),
        ("QUARTERLY BOARD REPORTING PROCESS", [
            "1. Run the monthly process for the quarter-end month (this gives YTD figures through the quarter).",
            "2. On Board_Report, review Sections 3–8 (Financial, Customer & Ops, Project Updates, Grant Summary, Risks, Decisions).",
            "3. Refresh the narrative bullets in Section 2 for the quarter; confirm Decisions Needed reflects board-level asks.",
            "4. Export Board_Report (and optionally each dashboard) to PDF: File ▸ Export ▸ PDF, or Print ▸ Save as PDF.",
            "5. Pair with the Strategic Initiative / Risk registers for the full board package.",
        ]),
        ("READING THE HEALTH COLORS", [
            "• GREEN = on/above target.  YELLOW = watch / within tolerance.  RED = off target, needs action.",
            "• Thresholds are documented in Data_Dictionary and encoded in each KPI card's status formula.",
            "• To change a threshold, edit the IF() formula in the small status cell beneath the KPI value.",
        ]),
        ("ADDING ROWS / NEW MONTHS", [
            "• Input tabs are Excel Tables — click the last row and press Tab, or type under the table; it expands automatically.",
            "• Rollup formulas read rows 4–1003 on each input tab, so newly added rows are picked up with no formula edits.",
            "• When you start a new fiscal year, copy the workbook, clear the input rows, and reset Sales_Targets & budgets.",
        ]),
    ]
    r = 3
    for title, lines in guide_blocks:
        c = ug.cell(row=r, column=2, value=title)
        c.font = S.F_BODY_B; c.fill = S.fill("E8EDF5"); c.alignment = LEFT; c.border = S.box_thin
        r += 1
        for ln in lines:
            cc = ug.cell(row=r, column=2, value=ln)
            cc.font = S.F_SMALL; cc.alignment = LEFT_TOP
            ug.row_dimensions[r].height = 14
            r += 1
        r += 1
    ug.freeze_panes = "B3"

    # ===================================================================
    # ASSUMPTIONS
    # ===================================================================
    asm = wb.create_sheet("Assumptions")
    S.page_setup(asm, tab_color="7F6000")
    asm.column_dimensions["A"].width = 3
    asm.column_dimensions["B"].width = 110
    asm["B1"] = "ASSUMPTIONS  &  RECOMMENDATIONS FOR FUTURE AUTOMATION"
    asm["B1"].font = S.F_SECTION; asm["B1"].fill = S.fill(S.NAVY)
    asm["B1"].alignment = LEFT; asm.row_dimensions[1].height = 24
    blocks = [
        ("ASSUMPTIONS USED IN THE BUILD", [
            "• Company profile: a regional fiber/FTTH ISP operating across southeastern NC counties (Robeson, Hoke,",
            "   Cumberland, Scotland, Bladen, Columbus), part-funded by BEAD / CAB / ARPA / State / RDOF capital programs.",
            "• Fiscal year shown is Jul-2025 → Jun-2026 (12 months). Current reporting month defaults to Jun-2026.",
            "• Sample data is illustrative and seeded for realism; replace with your actuals. All formulas remain valid.",
            "• 'Revenue' = total billed (recurring MRR + non-recurring install/equipment). 'OpEx' excludes capitalized build.",
            "• 'CapEx' = capitalized construction. Contractor cost is shown separately for cost-per-install/home math.",
            "• Cost per Install ≈ Contractor cost ÷ installs completed (a directional unit-economics proxy, not fully-loaded).",
            "• Cost per Home Passed ≈ CapEx ÷ homes passed added that month (directional; not all CapEx is plant).",
            "• Install backlog = YTD installs scheduled − YTD installs completed (a running open-order proxy).",
            "• Churn = monthly disconnects ÷ prior-month customers (simple logo churn, not revenue churn).",
            "• Risk score = Likelihood × Impact (1–5 each); RED ≥ 15, YELLOW ≥ 8, else GREEN.",
            "• RYG thresholds (churn, ARPU, days-to-install, margin, cash, etc.) are starting defaults — tune to your plan.",
            "• 90-day cash forecast is a simple run-rate projection (3× forecast revenue + outstanding grants − 3× cash outflows).",
            "• Trouble-ticket categories on the Ops dashboard are an editable helper block (no category field in the raw feed).",
        ]),
        ("WHAT WOULD MATERIALLY CHANGE THE DESIGN (confirm with COO)", [
            "• If revenue is recognized on a deferred/GAAP basis rather than billed, the Financial tab needs a recognition schedule.",
            "• If you track subscribers by service tier/speed, add a tier dimension to Customer_Data & Revenue_Detail.",
            "• If grant drawdowns are milestone-based with clawback risk, add a milestone schedule tab feeding the cash forecast.",
            "• If multiple legal entities/markets roll up, add an Entity column and a consolidation layer.",
        ]),
        ("RECOMMENDATIONS FOR FUTURE AUTOMATION", [
            "1. Connect input tabs to source systems: billing/CRM (customers, MRR, churn), OSS/field service (installs, tickets),",
            "   GIS/construction (homes passed, footage), and accounting (financials) via Power Query for scheduled refresh.",
            "2. Move to Power BI / Looker Studio for distribution: keep this workbook as the data model, publish dashboards online.",
            "3. Add a true subscriber ledger (one row per subscriber event) and derive cohorts, revenue churn, and LTV.",
            "4. Automate grant drawdown tracking against milestones with reimbursement-aging and cash-timing alerts.",
            "5. Add DORA-style operational SLAs (install SLA %, MTTR for outages) and contractor scorecards with auto-grading.",
            "6. Schedule a monthly snapshot (Power Automate) that saves a dated PDF of Board_Report to a board folder.",
            "7. Layer a rolling 13-week cash-flow model for treasury, fed by the same Financial & Grant inputs.",
            "8. Add data-validation guards / input forms to prevent staff entry errors at scale.",
        ]),
    ]
    r = 3
    for title, lines in blocks:
        c = asm.cell(row=r, column=2, value=title)
        c.font = S.F_BODY_B; c.fill = S.fill("FFF2CC"); c.alignment = LEFT; c.border = S.box_thin
        r += 1
        for ln in lines:
            cc = asm.cell(row=r, column=2, value=ln)
            cc.font = S.F_SMALL; cc.alignment = LEFT_TOP
            asm.row_dimensions[r].height = 14
            r += 1
        r += 1
    asm.freeze_panes = "B3"
