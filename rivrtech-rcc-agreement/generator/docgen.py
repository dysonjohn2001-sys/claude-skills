#!/usr/bin/env python3
"""Shared document-building helpers for the RIVRTECH / RCC package.

Font: Arial 11pt (renders as true Arial on Windows; Liberation Sans substitute
locally for PDF). Placeholders written as [ ... ] or $[ ... ] are auto-highlighted
yellow. Header shows the DRAFT banner; footer shows title + Page X of Y.
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
PLACEHOLDER_RE = re.compile(r'(\$?\[[^\]]+\])')

DRAFT_BANNER = "DRAFT FOR DISCUSSION AND LEGAL REVIEW"

# ---------------------------------------------------------------- base styles
def new_document():
    doc = Document()
    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.08
    # Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    return doc


def _set_run_font(run, size=11, bold=False, italic=False, color=None, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:cs"), font)


def add_rich(paragraph, text, size=11, bold=False, italic=False, color=None,
             font=FONT, highlight_placeholders=True):
    """Add text to a paragraph, auto-highlighting [placeholders] yellow."""
    if highlight_placeholders:
        parts = PLACEHOLDER_RE.split(text)
    else:
        parts = [text]
    for part in parts:
        if part == "":
            continue
        is_ph = highlight_placeholders and PLACEHOLDER_RE.fullmatch(part)
        # split on newlines so multi-line cells break correctly in Word
        lines = part.split("\n")
        for li, line in enumerate(lines):
            run = paragraph.add_run(line)
            _set_run_font(run, size=size, bold=bold, italic=italic,
                          color=color, font=font)
            if is_ph:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if li < len(lines) - 1:
                brk = paragraph.add_run()
                _set_run_font(brk, size=size, font=font)
                brk.add_break()
    return paragraph


def para(doc, text="", size=11, bold=False, italic=False, align=None,
         space_after=6, space_before=0, color=None, indent=None,
         highlight_placeholders=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic, color=color,
                 highlight_placeholders=highlight_placeholders)
    return p


def section_heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{number}. {title}")
    _set_run_font(run, size=12, bold=True)
    return p


def sub_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_rich(p, text, size=11, bold=True)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.3 + 0.3 * level)
    p.paragraph_format.space_after = Pt(2)
    add_rich(p, text)
    # force Arial on the bullet run(s)
    for r in p.runs:
        _set_run_font(r, size=11)
    return p


def numbered(doc, idx, text, indent=0.3):
    """Manual numbered item with hanging indent look."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent + 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    add_rich(p, f"{idx}\t{text}")
    return p


def page_break(doc):
    doc.add_page_break()


def quote_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_rich(p, text, italic=True)
    return p


# ------------------------------------------------------------- fields (X of Y)
def _add_field(paragraph, field_code, size=9):
    run = paragraph.add_run()
    _set_run_font(run, size=size)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_header_footer(doc, footer_title):
    for section in doc.sections:
        # Header: centered DRAFT banner
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = hp.add_run(DRAFT_BANNER)
        _set_run_font(hr, size=9, bold=True, color=RGBColor(0xB0, 0x00, 0x00))
        # bottom border on header
        _bottom_border(hp)

        # Footer: title (left) + Page X of Y (right) using a tab
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        # tab stops: left title, right page
        from docx.enum.text import WD_TAB_ALIGNMENT
        tab_stops = fp.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        fr = fp.add_run(footer_title + "\t")
        _set_run_font(fr, size=8, color=RGBColor(0x55, 0x55, 0x55))
        pr = fp.add_run("Page ")
        _set_run_font(pr, size=8, color=RGBColor(0x55, 0x55, 0x55))
        _add_field(fp, "PAGE", size=8)
        pr2 = fp.add_run(" of ")
        _set_run_font(pr2, size=8, color=RGBColor(0x55, 0x55, 0x55))
        _add_field(fp, "NUMPAGES", size=8)
        _top_border(fp)


def _bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _top_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), "999999")
    pBdr.append(top)
    pPr.append(pBdr)


# ------------------------------------------------------------- tables
def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=10, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    add_rich(p, text, size=size, bold=bold, color=color)
    return cell


def make_table(doc, headers, rows, col_widths=None, header_fill="1F3864",
               header_color=RGBColor(0xFF, 0xFF, 0xFF)):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, size=10, color=header_color)
        shade_cell(hdr[i], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table
